"""Full-document review is exhaustive, source-bound and separate from acceptance."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document

SCRIPTS = Path(__file__).resolve().parents[1] / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F
import review_plan as R


class ReviewPlanTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='fit-full-review-')
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.source = self.work / 'source.docx'
        doc = Document()
        for number in range(255):
            doc.add_paragraph(f'段落 {number + 1}：独立的完整原文。' + ('这是长段落，不能截成预览。' * 100 if number == 130 else ''))
        table = doc.add_table(rows=2, cols=2)
        for number, cell in enumerate(table._cells):
            cell.text = f'单元格 {number + 1}'
        p = doc.add_paragraph('图形宿主段落')
        drawing = F.E.SubElement(p._p, F.q('drawing'))
        textbox = F.E.SubElement(drawing, F.q('txbxContent'))
        nested = F.child(textbox, 'p')
        F.child(F.child(nested, 'r'), 't').text = '受保护图形内的文本也必须审核'
        doc.sections[0].header.paragraphs[0].text = '测试论文页眉'
        doc.sections[0].footer.paragraphs[0].text = '测试论文页脚'
        doc.save(self.source)
        # Comments are content to preserve even though existing apply has no comment map.
        parts = F.package(self.source)
        comments = F.E.Element(F.q('comments'), nsmap={'w': F.W})
        comment = F.child(comments, 'comment', id=0)
        F.child(F.child(F.child(comment, 'p'), 'r'), 't').text = '未开放自动排版的批注内容'
        parts['word/comments.xml'] = F.xml(comments)
        with zipfile.ZipFile(self.source, 'w') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        self.source_hash = F.sha(self.source.read_bytes())
        self.addCleanup(lambda: self.assertEqual(self.source_hash, F.sha(self.source.read_bytes())))

    def create(self, name='review'):
        folder = self.work / name
        result = R.create_review_plan(self.source, folder)
        self.assertEqual(result['counts']['reviewed'], 0)
        self.assertFalse(result['delivery_ready'])
        return folder

    def finish(self, folder, preserve_all=False):
        index = R._load(folder / 'index.json')
        for entry in index['chunks']:
            path = folder / entry['path']
            chunk = R._load(path)
            for item in chunk['items']:
                if preserve_all or not item['format_supported']:
                    item['decision'] = {'reviewed': True, 'action': 'preserve', 'reason': '已核对本项，保留原稿对象及现有排版'}
                else:
                    role = 'three-line' if item['kind'] == 'table' else 'table_text' if item['table_id'] else 'body'
                    item['decision'] = {'reviewed': True, 'action': 'format', 'role': role}
            path.write_text(json.dumps(chunk, ensure_ascii=False), encoding='utf-8')

    def assemble(self, folder, name='assembled'):
        mapping, coverage = self.work / (name + '.map.json'), self.work / (name + '.coverage.json')
        result = R.assemble_review_plan(self.source, folder, mapping, coverage)
        return result, mapping, coverage

    def test_complete_multichunk_apply_and_all_preserve_workflows(self):
        draft = self.work / 'incoming.draft.json'
        draft.write_text(json.dumps({'source_sha256': self.source_hash, 'reviewed': False, 'paragraphs': {},
                                     'suggestions_only': {'p000001': 'chapter'}}), encoding='utf-8')
        draft_bytes = draft.read_bytes()
        folder = self.work / 'review'
        result = R.create_review_plan(self.source, folder, draft_path=draft)
        self.assertGreaterEqual(result['counts']['chunks'], 3)
        first = R.next_chunk(self.source, folder)
        self.assertEqual(first['pending_chunks'], result['counts']['chunks'])
        self.assertLessEqual(len(first['chunk']['items']), 120)
        self.assertFalse(first['chunk']['items'][0]['decision']['reviewed'])
        self.assertTrue(first['chunk']['context']['after']['context_only'])
        status = R.review_status(self.source, folder)
        self.assertNotIn('chunk', status)
        self.assertEqual(status['counts']['unreviewed'], status['counts']['total'])
        self.assertEqual(status['next_chunk'], first['chunk_path'])
        index = R._load(folder / 'index.json')
        originals = [i for e in index['chunks'] for i in R._load(folder / e['path'])['items']]
        self.assertEqual(len({i['key'] for i in originals}), len(originals))
        self.assertGreater(len(originals[130]['text']), 1000)
        self.assertEqual(originals[130]['text_sha256'], F.sha(originals[130]['text'].encode('utf-8')))
        self.finish(folder)
        self.assertTrue(R.next_chunk(self.source, folder)['coverage_complete'])
        self.assertEqual(R.review_status(self.source, folder)['pending_chunks'], [])
        result, mapping_path, coverage = self.assemble(folder)
        self.assertTrue(result['apply_required'])
        verified = R.verify_coverage(self.source, coverage)
        self.assertTrue(all(verified['verifiable'].values()))
        self.assertEqual(verified['counts']['total'], verified['counts']['reviewed'])
        self.assertEqual(verified['counts']['total'], verified['mapped_item_count'] + verified['preserve_count'])
        self.assertEqual(draft.read_bytes(), draft_bytes)
        self.assertEqual((folder / 'style-map.draft.json').read_bytes(), draft_bytes)
        mapping = R._load(mapping_path)
        self.assertTrue(mapping['reviewed'])
        self.assertEqual(mapping['paragraphs']['p000001']['role'], 'body')
        self.assertNotIn('word/comments.xml', mapping['story_paragraphs'])
        output = self.work / 'formatted.docx'
        applied = F.apply_doc(self.source, output, mapping, self.work / 'format-report.json')
        self.assertTrue(applied['ok'], applied)
        self.assertTrue(applied['content_preserved'])
        with self.assertRaisesRegex(ValueError, 'new paths'):
            self.assemble(folder)
        all_preserve = self.create('all-preserve')
        self.finish(all_preserve, preserve_all=True)
        result, _, preserved_coverage = self.assemble(all_preserve, 'preserved')
        self.assertFalse(result['apply_required'])
        self.assertEqual(R.verify_coverage(self.source, preserved_coverage)['mapped_item_count'], 0)
        self.assertFalse(result['delivery_ready'])

    def test_missing_tail_duplicate_unknown_and_modified_original_fail(self):
        for case in ('missing_tail', 'duplicate', 'unknown', 'changed_text', 'missing_chunk'):
            with self.subTest(case=case):
                folder = self.create(case)
                self.finish(folder)
                index_path = folder / 'index.json'
                index = R._load(index_path)
                chunk_path = folder / index['chunks'][-1]['path']
                chunk = R._load(chunk_path)
                if case == 'missing_tail':
                    chunk['items'].pop()
                elif case == 'duplicate':
                    chunk['items'][-1] = copy.deepcopy(chunk['items'][0])
                elif case == 'unknown':
                    chunk['items'][-1]['key'] = 'word/document.xml#p999999'
                elif case == 'changed_text':
                    chunk['items'][-1]['text'] += '未经授权的修改'
                else:
                    index['chunks'].pop()
                    index_path.write_text(json.dumps(index), encoding='utf-8')
                chunk_path.write_text(json.dumps(chunk), encoding='utf-8')
                with self.assertRaises(ValueError):
                    self.assemble(folder, case)
                self.assertFalse((self.work / (case + '.map.json')).exists())
                self.assertFalse((self.work / (case + '.coverage.json')).exists())

    def test_unreviewed_tail_reason_and_protected_format_fail(self):
        for case in ('unreviewed_tail', 'empty_reason', 'unknown_role', 'protected_format'):
            with self.subTest(case=case):
                folder = self.create(case)
                self.finish(folder)
                index = R._load(folder / 'index.json')
                path = folder / index['chunks'][-1]['path']
                chunk = R._load(path)
                if case == 'unreviewed_tail':
                    chunk['items'][-1]['decision']['reviewed'] = False
                elif case == 'empty_reason':
                    chunk['items'][-1]['decision'] = {'reviewed': True, 'action': 'preserve', 'reason': ''}
                elif case == 'unknown_role':
                    chunk['items'][-1]['decision'] = {'reviewed': True, 'action': 'format', 'role': 'invented-role'}
                else:
                    protected = next(i for i in chunk['items'] if i['object_hints']['inside_protected_object'])
                    protected['decision'] = {'reviewed': True, 'action': 'format', 'role': 'body'}
                path.write_text(json.dumps(chunk), encoding='utf-8')
                with self.assertRaises(ValueError):
                    self.assemble(folder, case)
        only_first = self.create('only-first')
        first_path = only_first / R._load(only_first / 'index.json')['chunks'][0]['path']
        first = R._load(first_path)
        first['items'][0]['decision'] = {'reviewed': True, 'action': 'format', 'role': 'body'}
        first_path.write_text(json.dumps(first), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'Unreviewed item'):
            self.assemble(only_first, 'only-first')

    def test_source_change_and_tampered_evidence_fail_verification(self):
        folder = self.create()
        self.finish(folder)
        _, mapping_path, coverage_path = self.assemble(folder)
        changed = self.work / 'changed-source.docx'
        changed.write_bytes(self.source.read_bytes())
        with zipfile.ZipFile(changed, 'a') as archive:
            archive.writestr('extra.txt', '源文件已变化')
        with self.assertRaisesRegex(ValueError, 'Source SHA changed'):
            R.verify_coverage(changed, coverage_path)
        coverage_original = coverage_path.read_bytes()
        coverage = R._load(coverage_path)
        coverage['decisions'].pop()
        coverage_path.write_text(json.dumps(coverage), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'decisions'):
            R.verify_coverage(self.source, coverage_path)
        coverage_path.write_bytes(coverage_original)
        mapping_original = mapping_path.read_bytes()
        mapping = R._load(mapping_path)
        mapping['paragraphs']['p000001']['role'] = 'chapter'
        mapping_path.write_text(json.dumps(mapping), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'map hash changed'):
            R.verify_coverage(self.source, coverage_path)
        mapping_path.write_bytes(mapping_original)
        path = folder / 'chunks/chunk-0001.json'
        path.write_bytes(path.read_bytes() + b'\n')
        with self.assertRaisesRegex(ValueError, 'chunk hashes'):
            R.verify_coverage(self.source, coverage_path)


if __name__ == '__main__':
    unittest.main()

"""Content acceptance must never trust text equality alone."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document

SCRIPTS = Path(__file__).resolve().parents[1] / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F
import delivery_check as D
import review_plan as R


class DeliveryContentChainTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='fit-chain-')
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name).resolve()
        self.source = self.work / 'source.docx'
        doc = Document()
        doc.add_paragraph('合成原稿中的链接、编号与最后一个字符都必须保留。')
        doc.add_paragraph('第二段是末尾正文，不能因为只处理第一段就被遗漏。')
        doc.save(self.source)
        self.source_bytes = self.source.read_bytes()
        self.addCleanup(lambda: self.assertEqual(self.source.read_bytes(), self.source_bytes))

    def test_changed_relationship_or_numbering_needs_replay_even_if_text_matches(self):
        for member, selector, attribute in [
            ('word/_rels/document.xml.rels', None, 'Target'),
            ('word/numbering.xml', './/w:lvlText', F.q('val')),
        ]:
            with self.subTest(member=member):
                parts = F.package(self.source)
                root = F.parse(parts[member])
                element = list(root)[0] if selector is None else root.find(selector, F.NS)
                element.set(attribute, element.get(attribute, '') + '_changed')
                parts[member] = F.xml(root)
                final = self.work / ('candidate-' + str(len(member)) + '.docx')
                with zipfile.ZipFile(final, 'w') as output:
                    for name, data in parts.items(): output.writestr(name, data)
                self.assertEqual(F.compare_content(F.inventory(F.package(self.source)), F.inventory(parts)), [])
                result = D.check(self.source, final)
                content = [item for item in result['blocking_items'] if item['check'] == 'content_preservation']
                self.assertEqual(len(content), 1)
                self.assertIn('replayable stage chain', content[0]['reason'])
                self.assertFalse(result['delivery_ready'])

    def test_unchanged_source_still_requires_whole_document_review(self):
        result = D.check(self.source, self.source)
        checks = {item['check'] for item in result['blocking_items']}
        self.assertNotIn('content_preservation', checks)
        self.assertIn('whole_document_coverage', checks)
        self.assertFalse(result['delivery_ready'])

    def test_populated_unreviewed_draft_is_not_applied(self):
        final = self.work / 'unreviewed.docx'
        mapping = {'source_sha256': hashlib.sha256(self.source_bytes).hexdigest(),
                   'reviewed': False, 'paragraphs': {'p000001': 'body'}}
        with self.assertRaisesRegex(ValueError, 'unreviewed draft'):
            F.apply_doc(self.source, final, mapping, self.work / 'report.json')
        self.assertFalse(final.exists())

    def test_full_review_map_must_be_the_map_actually_applied(self):
        plan = self.work / 'review'
        R.create_review_plan(self.source, plan, chunk_size=1)
        for path in sorted((plan / 'chunks').glob('*.json')):
            chunk = json.loads(path.read_text(encoding='utf-8'))
            for item in chunk['items']:
                item['decision'] = {'reviewed': True, 'action': 'format', 'role': 'body'}
            path.write_text(json.dumps(chunk, ensure_ascii=False), encoding='utf-8')
        mapping_path, coverage = self.work / 'map.json', self.work / 'coverage.json'
        R.assemble_review_plan(self.source, plan, mapping_path, coverage)
        mapping = json.loads(mapping_path.read_text(encoding='utf-8'))
        for complete in (True, False):
            selected = json.loads(json.dumps(mapping))
            used_map = mapping_path
            if not complete:
                selected['paragraphs'].pop('p000002')
                used_map = self.work / 'partial-map.json'
                used_map.write_text(json.dumps(selected), encoding='utf-8')
            final = self.work / ('full.docx' if complete else 'partial.docx')
            F.apply_doc(self.source, final, selected, self.work / ('apply-' + str(complete) + '.json'))
            chain = self.work / ('chain-' + str(complete) + '.json')
            chain.write_text(json.dumps({'original_sha256': hashlib.sha256(self.source_bytes).hexdigest(),
                'stages': [{'scope': 'styles', 'source': str(self.source), 'output': str(final), 'mapping': str(used_map)}]}), encoding='utf-8')
            result = D.check(self.source, final, chain=chain, coverage=coverage)
            checks = {item['check'] for item in result['blocking_items']}
            self.assertNotIn('content_preservation', checks)
            if complete:
                self.assertNotIn('whole_document_coverage', checks)
                self.assertEqual(result['details']['whole_document_coverage']['counts']['reviewed'], 2)
            else:
                self.assertIn('whole_document_coverage', checks)


if __name__ == '__main__':
    unittest.main()

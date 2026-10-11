"""Full-length, entirely synthetic DOCX regression tests; no Word required.

Run: python -m unittest discover -s tests -p test_long_document.py -v
Keep a fixture for independent review:
    python tests/test_long_document.py --write-fixture NEW_DIRECTORY
The exported paper contains no user manuscript or school-template content.
"""
import copy
import hashlib
import io
import json
from pathlib import Path
import re
import struct
import subprocess
import sys
import tempfile
import unittest
import zipfile
import zlib

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
NS = {'w': W, 'm': M}
TEXT_TAGS = {qn('w:t'), qn('w:instrText'), qn('w:delText'), qn('m:t')}
BODY_PARAGRAPHS = 1602
END_MARKER = '全文末尾唯一哨兵：长文核验必须到达此处，朱雀玄武青龙白虎。'
BODY_SENTENCE = ('这是完全合成的长论文回归材料，用于检查篇幅增加之后文字顺序与研究内容是否保持完整。'
                 '每一段都必须被清单读取并应用已确认的语义样式，末尾内容与开头内容同样重要。')


def digest(data):
    return hashlib.sha256(data).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def load_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def read_parts(path):
    with zipfile.ZipFile(path) as bundle:
        return {name: bundle.read(name) for name in bundle.namelist()}


def write_parts(path, parts):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as bundle:
        for name, data in parts.items():
            bundle.writestr(name, data)


def paragraph_records(document_xml):
    """Read XML directly so the expected coverage does not use the inspector."""
    root = etree.fromstring(document_xml)
    return [{'id': 'p%06d' % index,
             'text': ''.join(node.text or '' for node in paragraph.iter() if node.tag in TEXT_TAGS),
             'top_level': paragraph.getparent().tag == qn('w:body')}
            for index, paragraph in enumerate(root.iter(qn('w:p')), 1)]


def text_nodes(parts):
    return {name: [(node.tag, node.text) for node in etree.fromstring(data).iter()
                   if node.tag in TEXT_TAGS]
            for name, data in parts.items() if name.endswith('.xml') and name != 'word/styles.xml'}


def xml_objects(document_xml, xpath):
    """Canonical object bytes include equations, drawing geometry and field data."""
    root = etree.fromstring(document_xml)
    return [etree.tostring(node, method='c14n') for node in root.xpath(xpath, namespaces=NS)]


def synthetic_png():
    """Tiny valid RGB PNG made from fixed bytes, without imaging dependencies."""
    def chunk(kind, payload):
        return struct.pack('>I', len(payload)) + kind + payload + struct.pack('>I', zlib.crc32(kind + payload))
    rows = b''.join(b'\x00' + bytes((40 + row * 10, 100, 170)) * 8 for row in range(8))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 8, 8, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b''))


def add_math(paragraph, number):
    math = OxmlElement('m:oMath')
    fraction = OxmlElement('m:f')
    for name, value in [('num', '合成分子' + str(number)), ('den', '合成分母' + str(number))]:
        container = OxmlElement('m:' + name)
        run = OxmlElement('m:r')
        token = OxmlElement('m:t'); token.text = value
        run.append(token); container.append(run); fraction.append(container)
    math.append(fraction)
    subscript = OxmlElement('m:sSub')
    for name, value in [('e', 'x'), ('sub', str(number))]:
        container = OxmlElement('m:' + name)
        run = OxmlElement('m:r'); token = OxmlElement('m:t'); token.text = value
        run.append(token); container.append(run); subscript.append(container)
    math.append(subscript); paragraph._p.append(math)


def add_bookmark_and_field(paragraph, number):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(number)); start.set(qn('w:name'), 'SyntheticAnchor%d' % number)
    paragraph._p.append(start)
    paragraph.add_run('书签锚点%d' % number)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(number)); paragraph._p.append(end)
    for kind in ('begin', 'instruction', 'separate', 'cached', 'end'):
        run = OxmlElement('w:r')
        if kind == 'instruction':
            node = OxmlElement('w:instrText')
            node.set(qn('xml:space'), 'preserve')
            node.text = ' REF SyntheticAnchor%d \\h ' % number
        elif kind == 'cached':
            node = OxmlElement('w:t'); node.text = '保留域缓存%d' % number
        else:
            node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), kind)
        run.append(node); paragraph._p.append(run)
    note = OxmlElement('w:footnoteReference'); note.set(qn('w:id'), str(number))
    paragraph.add_run()._r.append(note)


def add_footnotes(parts):
    root = etree.Element(qn('w:footnotes'), nsmap={'w': W})
    for number, kind in [(-1, 'separator'), (0, 'continuationSeparator')]:
        note = etree.SubElement(root, qn('w:footnote'), {qn('w:id'): str(number), qn('w:type'): kind})
        run = etree.SubElement(etree.SubElement(note, qn('w:p')), qn('w:r'))
        etree.SubElement(run, qn('w:' + kind))
    for number in range(1, 4):
        note = etree.SubElement(root, qn('w:footnote'), {qn('w:id'): str(number)})
        paragraph = etree.SubElement(note, qn('w:p'))
        run = etree.SubElement(paragraph, qn('w:r')); etree.SubElement(run, qn('w:footnoteRef'))
        run = etree.SubElement(paragraph, qn('w:r'))
        etree.SubElement(run, qn('w:t')).text = '合成脚注%d：保留脚注正文、引用位置和编号。' % number
    parts['word/footnotes.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    relationships = etree.fromstring(parts['word/_rels/document.xml.rels'])
    etree.SubElement(relationships, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                     Id='rIdSyntheticFootnotes',
                     Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',
                     Target='footnotes.xml')
    parts['word/_rels/document.xml.rels'] = etree.tostring(relationships, xml_declaration=True, encoding='UTF-8')
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                     PartName='/word/footnotes.xml',
                     ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8')


def create_fixture(directory):
    """Export source, complete reviewed synthetic map and independent summary."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=False)
    source = directory / 'synthetic-long-paper.docx'
    document = Document()
    document.core_properties.title = '完全合成长论文压力测试原稿'
    document.core_properties.author = 'Synthetic regression fixture'
    roles = {}
    anchors = {}

    def add(text, role='body'):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        run = paragraph.add_run(text); run.font.size = Pt(11)
        roles[paragraph._p] = role
        return paragraph

    add('完全合成长论文：仅用于全文覆盖和内容保留测试', 'cn_title')
    for number in range(1, BODY_PARAGRAPHS + 1):
        paragraph = add('合成正文第%04d段：' % number + BODY_SENTENCE + '  原始空格  [1]。')
        if number in (1, BODY_PARAGRAPHS // 2, BODY_PARAGRAPHS):
            phase = len(anchors) + 1
            anchors[('first', 'middle', 'last')[phase - 1]] = paragraph._p
            paragraph.runs[0].bold = True
            paragraph.add_run('斜体强调保留').italic = True
            paragraph.add_run('2').font.superscript = True
            # Every supported role occurs at the beginning, middle and end.
            # These explicitly labelled samples test formatting, not thesis prose.
            for role in F.ROLES:
                add('合成角色样例%d：%s。' % (phase, role), role)
            add('表%d-1 合成数据表' % phase, 'table_caption')
            table = document.add_table(rows=3, cols=3)
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    cell.text = '合成表%d行%d列%d：数值%d' % (phase, row_index, column_index, phase * 100 + row_index * 10 + column_index)
                    roles[cell.paragraphs[0]._p] = 'table_text'
            equation = add('公式（%d-1）' % phase, 'equation'); add_math(equation, phase)
            figure = add('合成图片锚点%d' % phase, 'figure')
            figure.add_run().add_picture(io.BytesIO(synthetic_png()), width=Inches(0.3))
            add('图%d-1 合成像素图' % phase, 'figure_caption')
            field = add('合成域和脚注%d：' % phase); add_bookmark_and_field(field, phase)
    # This paragraph is longer than typical previews: suffix truncation must fail.
    add('长段落覆盖：' + BODY_SENTENCE * 24 + '长段落尾部唯一标记。')
    add('致谢', 'acknowledgement_heading')
    add(END_MARKER)
    header = document.sections[0].header.paragraphs[0]
    header.text = '合成页眉：原文不得改变'
    footer = document.sections[0].footer.paragraphs[0]
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), ' PAGE ')
    run = OxmlElement('w:r'); token = OxmlElement('w:t'); token.text = '1'; run.append(token)
    field.append(run); footer._p.append(field)
    ordered = list(document.element.iter(qn('w:p')))
    position_roles = [roles[paragraph] for paragraph in ordered]
    anchor_ids = {label: 'p%06d' % (ordered.index(node) + 1) for label, node in anchors.items()}
    document.save(source)
    parts = read_parts(source); add_footnotes(parts); write_parts(source, parts)
    records = paragraph_records(parts['word/document.xml'])
    mapping = {'source_sha256': digest(source.read_bytes()), 'reviewed': True,
               'paragraphs': {record['id']: {'role': role, 'expected_text': record['text']}
                              for record, role in zip(records, position_roles)},
               'tables': {'t%04d' % number: {'kind': 'three-line', 'header_rows': 1} for number in range(1, 4)},
               'story_paragraphs': {'word/header1.xml': {'p000001': 'header'},
                                    'word/footer1.xml': {'p000001': 'page_number'}}}
    all_text = ''.join(record['text'] for record in records)
    summary = {'schema_version': 1, 'synthetic_only': True, 'source': source.name,
               'source_sha256': mapping['source_sha256'],
               'paragraph_count': len(records), 'body_paragraph_count': BODY_PARAGRAPHS,
               'top_level_paragraph_count': sum(record['top_level'] for record in records),
               'cjk_character_count': len(re.findall(r'[\u3400-\u9fff]', all_text)),
               'whole_text_sha256': digest(all_text.encode('utf-8')),
               'end_marker': END_MARKER, 'anchor_ids': anchor_ids,
               'role_count': len(F.ROLES), 'table_count': 3, 'math_count': 3, 'drawing_count': 3,
               'footnote_count': 3, 'bookmark_count': 3, 'complex_field_count': 3,
               'paragraphs': [{'id': record['id'], 'role': role, 'length': len(record['text']),
                               'text_sha256': digest(record['text'].encode('utf-8'))}
                              for record, role in zip(records, position_roles)],
               'media_sha256': {name: digest(data) for name, data in parts.items() if name.startswith('word/media/')},
               'note': 'Synthetic fixture only. No rendering or school-format acceptance is claimed.'}
    save_json(directory / 'expected-summary.json', summary)
    save_json(directory / 'reviewed-map.json', mapping)
    return source, summary, mapping


def cli(request, work):
    completed = subprocess.run([sys.executable, str(SCRIPTS / 'agent_cli.py'), '--request', '-'],
                               input=json.dumps(request, ensure_ascii=False), cwd=work,
                               capture_output=True, text=True, encoding='utf-8', timeout=120)
    if completed.returncode != 0:
        raise AssertionError(completed.stdout + completed.stderr)
    return json.loads(completed.stdout)


class LongDocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix='fit long 中文 ')
        cls.addClassCleanup(cls.temp.cleanup)
        cls.work = Path(cls.temp.name).resolve()
        cls.source, cls.expected, cls.mapping = create_fixture(cls.work / 'fixture')
        cls.before = read_parts(cls.source)
        cls.records = paragraph_records(cls.before['word/document.xml'])
        cls.output = cls.work / 'formatted.docx'
        cls.applied = cli({'action': 'apply_styles', 'source': str(cls.source), 'output': str(cls.output),
                           'mapping': str(cls.source.parent / 'reviewed-map.json'),
                           'report': str(cls.work / 'applied.json')}, cls.work)

    def tearDown(self):
        self.assertEqual(digest(self.source.read_bytes()), self.expected['source_sha256'])

    def test_prepare_and_inspect_include_every_paragraph_and_final_citation(self):
        self.assertGreaterEqual(self.expected['cjk_character_count'], 50000)
        self.assertGreaterEqual(self.expected['paragraph_count'], 1500)
        prepared = self.work / 'prepared'
        result = cli({'action': 'prepare', 'source': str(self.source), 'work_dir': str(prepared)}, self.work)
        self.assertTrue(result['operation_succeeded'], result)
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(digest((prepared / 'original.docx').read_bytes()), self.expected['source_sha256'])
        baseline = load_json(prepared / 'baseline.json')
        self.assertEqual(baseline['source_sha256'], self.expected['source_sha256'])
        styles = load_json(prepared / 'styles-inventory.json')
        expected_pairs = [(item['id'], item['text']) for item in self.records]
        self.assertEqual([(item['id'], item['text']) for item in styles['paragraphs']], expected_pairs)
        self.assertEqual(styles['paragraphs'][-1]['text'], END_MARKER)
        self.assertEqual(styles['counts']['math'], 3)
        self.assertEqual(styles['counts']['drawing'], 3)
        self.assertEqual(len(styles['tables']), 3)
        self.assertIn('word/footnotes.xml', styles['stories'])
        layout = load_json(prepared / 'layout-inventory.json')
        self.assertEqual([(item['id'], item['text']) for item in layout['top_level_paragraphs']],
                         [(item['id'], item['text']) for item in self.records if item['top_level']])
        citations = load_json(prepared / 'citation-inventory.json')['candidates']
        self.assertEqual(len(citations), BODY_PARAGRAPHS)
        self.assertEqual({item['pid'] for item in citations},
                         {item['id'] for item in self.records if '[1]' in item['text']})
        for item in citations:
            text = self.mapping['paragraphs'][item['pid']]['expected_text']
            self.assertEqual(text[item['start']:item['end']], '[1]')
        self.assertIn(self.expected['anchor_ids']['last'], {item['pid'] for item in citations})
        draft = load_json(prepared / 'style-map.draft.json')
        self.assertFalse(draft['reviewed'])
        self.assertEqual(draft['paragraphs'], {})

    def test_all_roles_apply_through_the_end_without_content_or_object_loss(self):
        result = self.applied
        self.assertTrue(result['ok'], result)
        self.assertFalse(result['delivery_ready'])
        self.assertTrue(result['result']['content_preserved'])
        self.assertEqual(result['result']['selected_paragraphs'], self.expected['paragraph_count'])
        # The full report must expose coverage even when the CLI returns a compact summary.
        report = load_json(self.work / 'applied.json')
        self.assertEqual(report['unmapped_paragraphs'], [])
        after = read_parts(self.output)
        self.assertEqual(set(after), set(self.before))
        self.assertEqual(text_nodes(after), text_nodes(self.before))
        self.assertEqual(paragraph_records(after['word/document.xml']), self.records)
        for name, data in self.before.items():
            if name not in ('word/document.xml', 'word/styles.xml', 'word/header1.xml', 'word/footer1.xml'):
                self.assertEqual(after[name], data, 'Unrelated package part changed: ' + name)
        for xpath in ('.//m:oMath', './/w:drawing', './/w:bookmarkStart', './/w:bookmarkEnd',
                      './/w:fldChar', './/w:instrText', './/w:footnoteReference'):
            self.assertEqual(xml_objects(after['word/document.xml'], xpath),
                             xml_objects(self.before['word/document.xml'], xpath), xpath)
        root = etree.fromstring(after['word/document.xml'])
        used_roles = set()
        for index, paragraph in enumerate(root.iter(qn('w:p')), 1):
            pid = 'p%06d' % index
            role = self.mapping['paragraphs'][pid]['role']; used_roles.add(role)
            self.assertEqual(paragraph.find('w:pPr/w:pStyle', NS).get(qn('w:val')), 'FIT_' + role, pid)
        self.assertEqual(used_roles, set(F.ROLES))
        for label, pid in self.expected['anchor_ids'].items():
            paragraph = list(root.iter(qn('w:p')))[int(pid[1:]) - 1]
            self.assertEqual(paragraph.find('w:pPr/w:pStyle', NS).get(qn('w:val')), 'FIT_body', label)
            self.assertIsNotNone(paragraph.find('w:r/w:rPr/w:b', NS), label)
            self.assertIsNotNone(paragraph.find('w:r/w:rPr/w:i', NS), label)
            self.assertIsNotNone(paragraph.find('w:r/w:rPr/w:vertAlign', NS), label)
        self.assertTrue(F.verify_doc(self.output, self.source, self.mapping)['ok'])

    def test_verification_rejects_tail_text_deletion_and_late_equation_damage(self):
        clean = read_parts(self.output)
        for corruption in ('last_character', 'last_equation'):
            with self.subTest(corruption=corruption):
                parts = copy.copy(clean)
                root = etree.fromstring(parts['word/document.xml'])
                if corruption == 'last_character':
                    node = list(root.iter(qn('w:t')))[-1]
                    self.assertEqual(node.text, END_MARKER)
                    node.text = node.text[:-1]
                else:
                    node = root.xpath('.//m:oMath', namespaces=NS)[-1]
                    node.getparent().remove(node)
                parts['word/document.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8')
                damaged = self.work / (corruption + '.docx'); write_parts(damaged, parts)
                result = F.verify_doc(damaged, self.source, self.mapping)
                self.assertFalse(result['ok'], result)
                self.assertFalse(result['content_preserved'], result)
                self.assertTrue(any('Original content changed:' in error for error in result['errors']))


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) == 3 and sys.argv[1] == '--write-fixture':
        source, expected, _ = create_fixture(Path(sys.argv[2]))
        print(json.dumps({'source': str(source.resolve()), 'paragraphs': expected['paragraph_count'],
                          'cjk_characters': expected['cjk_character_count'],
                          'summary': str((source.parent / 'expected-summary.json').resolve())}, ensure_ascii=False))
    else:
        unittest.main()

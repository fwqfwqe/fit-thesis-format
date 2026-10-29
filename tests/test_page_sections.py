"""School running furniture across synthetic cover, front and body sections.

These checks exercise actual DOCX transforms and source-bound verification.
They do not claim to validate Word's rendered page count or visual pagination.
"""
from pathlib import Path
import posixpath
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Cm

SCRIPTS = Path(__file__).resolve().parents[1] / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F
import thesis_layout as L


class PageSectionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='fit-page-sections-')
        self.addCleanup(temporary.cleanup)
        self.work = Path(temporary.name)
        self.source = self.work / 'synthetic.docx'
        self.title = '合成论文：页眉页脚回归测试'
        document = Document()
        blocks = (
            (self.title, '合成封面信息'),
            ('诚信承诺书', '合成承诺书正文'),
            ('中文摘要', '合成中文摘要正文'),
            ('Abstract', 'Synthetic abstract and contents.'),
            ('一、绪论', '正文第一节，含有公式对象。'),
            ('二、后续章节', '参考文献、附录及致谢沿用正文页码。'),
        )
        for number, paragraphs in enumerate(blocks):
            if number:
                document.add_section(WD_SECTION_START.NEW_PAGE)
            for text in paragraphs:
                document.add_paragraph(text)
            if number == 4:
                equation = document.add_paragraph()
                math = F.E.SubElement(equation._p, '{' + F.M + '}oMath')
                run = F.E.SubElement(math, '{' + F.M + '}r')
                F.E.SubElement(run, '{' + F.M + '}t').text = 'x = 1'

        # Deliberately wrong inherited furniture and numbering must be replaced.
        for number, section in enumerate(document.sections, 1):
            section.different_first_page_header_footer = True
            section.header_distance = Cm(0.7)
            section.footer_distance = Cm(0.9)
            for story in (section.header, section.even_page_header,
                          section.first_page_header, section.footer,
                          section.even_page_footer, section.first_page_footer):
                story.is_linked_to_previous = False
                story.paragraphs[0].text = '旧版布局文字 %d' % number
            F.child(section._sectPr, 'pgNumType', fmt='lowerRoman', start=8)
        document.save(self.source)
        self.source_bytes = self.source.read_bytes()

        paragraphs, _ = F.indexed(F.parse(F.package(self.source)[L.DOC]))
        self.pids = {F.ptext(node): pid for pid, node in paragraphs.items()}
        style_map = {
            'source_sha256': F.sha(self.source_bytes),
            'reviewed': True,
            'paragraphs': {self.pids[self.title]: {
                'role': 'cn_title', 'expected_text': self.title}},
            'page_setup': True,
        }
        self.styled = self.work / 'styled.docx'
        styled = F.apply_doc(self.source, self.styled, style_map, self.work / 'styles.json')
        self.assertTrue(styled['ok'], styled)
        self.assertTrue(F.verify_doc(self.styled, self.source, style_map)['ok'])
        self.mapping = {
            'source_sha256': F.sha(self.styled.read_bytes()),
            'title': {'paragraph_id': self.pids[self.title], 'expected_text': self.title},
            'regions': [
                {'start': 'start', 'kind': 'unpaged'},
                {'start': self.pids['中文摘要'], 'kind': 'front', 'expected_text': '中文摘要'},
                {'start': self.pids['一、绪论'], 'kind': 'body', 'expected_text': '一、绪论'},
            ],
        }
        self.output = self.work / 'layout.docx'
        self.result = L.apply(self.styled, self.output, self.mapping)
        self.assertTrue(self.result['ok'], self.result)
        self.addCleanup(lambda: self.assertEqual(self.source.read_bytes(), self.source_bytes))

    def story(self, parts, section, tag, kind):
        reference = section.find('w:%s[@w:type="%s"]' % (tag, kind), F.NS)
        self.assertIsNotNone(reference, (tag, kind))
        relationships = F.parse(parts[L.RELS])
        rid = reference.get('{' + L.R + '}id')
        relation = next(node for node in relationships if node.get('Id') == rid)
        path = posixpath.normpath(posixpath.join('word', relation.get('Target')))
        return F.parse(parts[path])

    def test_cover_front_and_multiple_body_sections(self):
        verified = L.verify(self.output, self.styled, self.mapping)
        self.assertTrue(verified['ok'], verified)
        parts = F.package(self.output)
        root = F.parse(parts[L.DOC])
        sections = L.sections(root)[1]
        self.assertEqual(len(sections), 6)
        self.assertEqual([row['kind'] for row in self.result['sections']],
                         ['unpaged', 'unpaged', 'front', 'front', 'body', 'body'])
        self.assertEqual([row['restart'] for row in self.result['sections']],
                         [False, False, True, False, True, False])
        self.assertIsNotNone(F.parse(parts[L.SETTINGS]).find(F.q('evenAndOddHeaders')))

        for index, (_, section) in enumerate(sections):
            with self.subTest(section=index + 1):
                self.assertIsNone(section.find(F.q('titlePg')))
                margins = section.find(F.q('pgMar'))
                self.assertEqual(margins.get(F.q('header')), '850')
                self.assertEqual(margins.get(F.q('footer')), '850')
                numbering = section.find(F.q('pgNumType'))
                if index < 2:
                    self.assertIsNone(numbering)
                else:
                    self.assertEqual(numbering.get(F.q('fmt')), 'upperRoman' if index < 4 else 'decimal')
                    self.assertEqual(numbering.get(F.q('start')), '1' if index in (2, 4) else None)
                for kind in ('default', 'even', 'first'):
                    header = self.story(parts, section, 'headerReference', kind)
                    footer = self.story(parts, section, 'footerReference', kind)
                    header_text = ''.join(node.text or '' for node in header.iter(F.q('t')))
                    self.assertEqual(header_text,
                                     (L.SCHOOL if kind == 'even' else self.title) if index >= 4 else '')
                    page_instructions = [node.text.strip() for node in footer.iter(F.q('instrText'))]
                    self.assertEqual(page_instructions, [] if index < 2 else ['PAGE'])
                    if index < 2:
                        self.assertEqual(''.join(node.text or '' for node in footer.iter(F.q('t'))), '')
                    else:
                        self.assertEqual([node.get(F.q('fldCharType')) for node in footer.iter(F.q('fldChar'))],
                                         ['begin', 'separate', 'end'])
                        self.assertEqual(footer.find('w:p/w:pPr/w:pStyle', F.NS).get(F.q('val')),
                                         F.style_id('page_number'))
                    if index >= 4:
                        self.assertEqual(header.find('w:p/w:pPr/w:pStyle', F.NS).get(F.q('val')),
                                         F.style_id('header'))

        styles = F.parse(parts['word/styles.xml'])
        for role, east_asia, size in (('header', '宋体', '21'), ('page_number', 'Times New Roman', '18')):
            style = styles.find('w:style[@w:styleId="%s"]' % F.style_id(role), F.NS)
            self.assertEqual(style.find('w:pPr/w:jc', F.NS).get(F.q('val')), 'center')
            fonts = style.find('w:rPr/w:rFonts', F.NS)
            self.assertEqual(fonts.get(F.q('eastAsia')), east_asia)
            self.assertEqual(fonts.get(F.q('ascii')), 'Times New Roman')
            self.assertEqual(style.find('w:rPr/w:sz', F.NS).get(F.q('val')), size)
        border = styles.find('w:style[@w:styleId="FIT_header"]/w:pPr/w:pBdr/w:bottom', F.NS)
        self.assertEqual((border.get(F.q('val')), border.get(F.q('sz')), border.get(F.q('space'))),
                         ('single', '4', '1'))
        original = F.package(self.source)
        self.assertEqual(F.content_skeleton(F.parse(original[L.DOC])), F.content_skeleton(root))
        for name, data in original.items():
            if name.startswith(('word/header', 'word/footer')) and name.endswith('.xml'):
                self.assertEqual(parts[name], data, 'Existing story bytes must remain preserved: ' + name)

    def test_verify_rejects_missing_even_footer_and_restarted_later_body(self):
        for mutation in ('missing_even_footer', 'body_restart'):
            with self.subTest(mutation=mutation):
                parts = F.package(self.output)
                root = F.parse(parts[L.DOC])
                section = L.sections(root)[1][-1][1]
                if mutation == 'missing_even_footer':
                    section.remove(section.find('w:footerReference[@w:type="even"]', F.NS))
                else:
                    section.find(F.q('pgNumType')).set(F.q('start'), '1')
                parts[L.DOC] = F.xml(root)
                damaged = self.work / (mutation + '.docx')
                with zipfile.ZipFile(damaged, 'w', zipfile.ZIP_DEFLATED) as archive:
                    for name, data in parts.items():
                        archive.writestr(name, data)
                result = L.verify(damaged, self.styled, self.mapping)
                self.assertFalse(result['ok'])
                self.assertIn('Unexpected change: word/document.xml', result['errors'])


if __name__ == '__main__':
    unittest.main()

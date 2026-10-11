"""Full-document audit coverage must not silently omit an unclassified tail."""
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "fit-thesis-format" / "scripts"))
from format_audit import Audit


class LongAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="fit-long-audit-")
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "anonymous.docx"

    def audit(self, document, roles=None):
        document.save(self.source)
        before = self.source.read_bytes()
        mapping = None if roles is None else {
            "source_sha256": hashlib.sha256(before).hexdigest(), "paragraphs": roles}
        result = Audit(self.source, mapping=mapping).run()
        self.assertEqual(self.source.read_bytes(), before)
        self.assertTrue(result["source_unchanged"])
        return result

    def coverage_finding(self, result):
        findings = [item for item in result["issues"]
                    if item["category"] == "coverage.unclassified_content"]
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["severity"], "warning")
        self.assertEqual(findings[0]["status"], "requires_role_review")
        return findings[0]

    def test_six_hundred_appendix_paragraphs_require_explicit_review(self):
        doc = Document()
        doc.add_paragraph("一、研究背景")
        doc.add_paragraph("正文内容。")
        doc.add_paragraph("附录")
        for index in range(600):
            paragraph = doc.add_paragraph(f"附录第{index + 1}段：访谈记录原文。")
            paragraph.runs[0].font.size = Pt(36)
        result = self.audit(doc)
        finding = self.coverage_finding(result)
        self.assertEqual(finding["actual"]["unclassified_count"], 600)
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"],
                         [f"p{index:06d}" for index in range(4, 604)])
        self.assertEqual(result["coverage"]["unclassified_ranges"],
                         [{"first": "p000004", "last": "p000603", "count": 600}])
        self.assertNotIn("p000603", result["role_identification"])
        self.assertTrue(result["coverage"]["role_review_required"])

    def test_unknown_decimal_heading_does_not_hide_long_body(self):
        doc = Document()
        doc.add_paragraph("1 引言")
        for index in range(600):
            doc.add_paragraph(f"第{index + 1}段研究正文，保留原文。")
        result = self.audit(doc)
        finding = self.coverage_finding(result)
        self.assertEqual(finding["actual"]["unclassified_count"], 601)
        self.assertEqual(result["summary"]["unclassified_substantive_paragraph_count"], 601)
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"][-1], "p000601")
        self.assertEqual(result["role_identification"], {})

    def test_uncertain_toc_boundary_cannot_hide_later_prose(self):
        doc = Document()
        doc.add_paragraph("目录")
        doc.add_paragraph("1 绪论")
        for index in range(200):
            doc.add_paragraph(f"目录之后第{index + 1}段正文，需要确认正文边界。")
        result = self.audit(doc)
        self.coverage_finding(result)
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"][-1], "p000202")
        self.assertNotIn("p000202", result["role_identification"])

    def test_blank_spacing_is_excluded_but_unclassified_field_is_not(self):
        doc = Document()
        doc.add_paragraph()
        doc.add_paragraph(" \t ")
        paragraph = doc.add_paragraph()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " PAGE ")
        paragraph._p.append(field)
        result = self.audit(doc)
        self.coverage_finding(result)
        self.assertEqual(result["coverage"]["substantive_paragraph_count"], 1)
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"], ["p000003"])

    def test_heading_style_on_long_prose_is_only_a_review_candidate(self):
        doc = Document()
        doc.add_paragraph("这是误套标题样式的研究正文，包含分析过程和访谈结果，并不是已确认的章节名称。" * 4,
                          style="Heading 1")
        result = self.audit(doc)
        findings = [item for item in result["issues"]
                    if item["category"] == "heading.long_text_requires_review"]
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["status"], "requires_role_review")
        self.assertEqual(findings[0]["location"]["paragraph_id"], "p000001")
        self.assertEqual(result["role_identification"]["p000001"]["role"], "chapter")
        self.assertEqual(result["role_identification"]["p000001"]["confidence"], "medium")

    def test_explicit_roles_resolve_coverage_without_guessing_or_editing(self):
        doc = Document()
        doc.add_paragraph("附录")
        doc.add_paragraph("附录访谈记录。")
        doc.add_paragraph("已经阅读并确认的长篇正文。" * 15, style="Heading 1")
        result = self.audit(doc, {"p000001": "appendix_heading", "p000002": "body", "p000003": "body"})
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"], [])
        self.assertFalse(result["coverage"]["role_review_required"])
        self.assertFalse(any(item["category"] in {"coverage.unclassified_content", "heading.long_text_requires_review"}
                             for item in result["issues"]))
        self.assertEqual(result["role_identification"]["p000003"]["role"], "body")

    def test_toc_closing_marker_alone_has_no_displayed_role(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        run = OxmlElement('w:r')
        marker = OxmlElement('w:fldChar')
        marker.set(qn('w:fldCharType'), 'end')
        run.append(marker)
        paragraph._p.append(run)
        result = self.audit(doc)
        self.assertEqual(result['coverage']['unclassified_paragraph_ids'], [])

    def test_known_long_body_checks_last_paragraph_format(self):
        doc = Document()
        doc.add_paragraph("一、研究背景")
        for index in range(600):
            paragraph = doc.add_paragraph(f"正文第{index + 1}段，包含研究信息。")
        paragraph.runs[0].font.size = Pt(36)
        result = self.audit(doc)
        self.assertEqual(result["coverage"]["unclassified_paragraph_ids"], [])
        self.assertEqual(result["role_identification"]["p000601"]["role"], "body")
        self.assertTrue(any(item["category"] == "paragraph.format" and
                            item["location"].get("paragraph_id") == "p000601"
                            for item in result["issues"]))


if __name__ == "__main__":
    unittest.main()

"""Portable integration tests: stdlib unittest plus the package's lxml dependency.

Run at repository root: python -m unittest discover -s tests -p test_agent_cli.py -v
The included school template is read only; all outputs use temporary folders.
"""
import hashlib
from contextlib import contextmanager
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "fit-thesis-format" / "scripts" / "agent_cli.py"
TEMPLATE = ROOT / "fit-thesis-format" / "assets" / "fit-economics-management-template.docx"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class AgentCliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="fit agent 中文 ")
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name).resolve()
        self.original_hash = digest(TEMPLATE)
        self.addCleanup(lambda: self.assertEqual(digest(TEMPLATE), self.original_hash))

    def call(self, request, file_mode=False):
        args = [sys.executable, str(CLI), "--request", "-"]
        data = json.dumps(request, ensure_ascii=False)
        if file_mode:
            request_file = self.work / "请求子目录" / "请求 中文.json"
            request_file.parent.mkdir(parents=True, exist_ok=True)
            request_file.write_text(data, encoding="utf-8-sig")
            args[-1] = str(request_file)
        completed = subprocess.run(args, input=None if file_mode else data, cwd=self.work,
                                   capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(completed.stderr, "")
        result = json.loads(completed.stdout)
        self.assertEqual(completed.returncode, 0 if result["ok"] else 2)
        self.assertFalse(result["render_review_performed"])
        self.assertFalse(result["delivery_ready"])
        self.assertEqual(result["operation_succeeded"], result["ok"])
        self.assertIn("remaining_steps", result)
        return result

    def test_doctor_no_document_needed(self):
        result = self.call({"action": "doctor", "legacy": True}, file_mode=True)
        self.assertTrue(result["ok"])
        self.assertIn("lxml", result["result"]["dependencies"])
        self.assertFalse(result["result"]["render_candidates"]["operationally_tested"])

    def test_inspect_each_scope_template_read_only(self):
        for scope in ("styles", "layout", "citations"):
            with self.subTest(scope=scope):
                result = self.call({"action": "inspect", "scope": scope,
                                    "source": str(TEMPLATE), "report": scope + " 检查.json"})
                self.assertTrue(result["ok"], result)
                report = Path(result["outputs"]["report"])
                self.assertEqual(report.parent, self.work)
                data = json.loads(report.read_text(encoding="utf-8"))
                self.assertEqual(data["source_sha256"], self.original_hash)
                self.assertEqual(result["checks"]["rendering"], "not_run")

    def test_real_audit_reports_not_visual_acceptance(self):
        result = self.call({"action": "audit", "source": str(TEMPLATE),
                            "report": "审核.json", "markdown": "审核.md"}, file_mode=True)
        self.assertTrue(result["ok"], result)
        self.assertTrue(result["needs_render_review"])
        self.assertTrue(result["checks"]["input_files_unchanged"]["source"])
        data = json.loads(Path(result["outputs"]["report"]).read_text(encoding="utf-8"))
        self.assertEqual(Path(result["outputs"]["report"]).parent, self.work)
        self.assertIn("summary", data)
        self.assertGreater(Path(result["outputs"]["markdown"]).stat().st_size, 0)

    def test_reject_unknown_action_fields_and_existing_report(self):
        self.assertEqual(self.call({"action": "arbitrary_shell"})["error"]["code"], "UNKNOWN_ACTION")
        self.assertEqual(self.call({"action": "doctor", "legacy": True, "shell": "anything"})["error"]["code"], "INVALID_REQUEST")
        existing = self.work / "keep.json"
        existing.write_text("keep exactly", encoding="utf-8")
        result = self.call({"action": "inspect", "source": str(TEMPLATE), "report": str(existing)})
        self.assertEqual(result["error"]["code"], "OUTPUT_EXISTS")
        self.assertEqual(existing.read_text(encoding="utf-8"), "keep exactly")
        result = self.call({"action": "inspect", "source": str(TEMPLATE), "report": str(TEMPLATE)})
        self.assertEqual(result["error"]["code"], "OUTPUT_CONFLICT")

    def test_mapping_hash_is_mandatory_and_verify_needs_baseline(self):
        mapping = self.work / "stale.json"
        mapping.write_text(json.dumps({"source_sha256": "0" * 64, "paragraphs": {}}), encoding="utf-8")
        result = self.call({"action": "apply_styles", "source": str(TEMPLATE), "output": "new.docx",
                            "mapping": str(mapping), "report": "new.json"})
        self.assertEqual(result["error"]["code"], "STALE_MAPPING")
        self.assertFalse((self.work / "new.docx").exists())
        result = self.call({"action": "verify", "output": str(TEMPLATE), "report": "verify.json"})
        self.assertEqual(result["error"]["code"], "INVALID_REQUEST")

    def test_empty_map_is_not_reported_as_a_formatted_thesis(self):
        mapping = self.work / "empty-map.json"
        mapping.write_text(json.dumps({"source_sha256": self.original_hash, "paragraphs": {}}), encoding="utf-8")
        result = self.call({"action": "apply_styles", "source": str(TEMPLATE), "output": "empty.docx",
                            "mapping": str(mapping), "report": "empty-report.json"})
        self.assertFalse(result["ok"])
        self.assertIn("Empty formatting map", result["result"]["error"])
        self.assertEqual(result["status"], "failed")
        self.assertFalse((self.work / "empty.docx").exists())

    def test_zero_exit_missing_or_nonboolean_success_is_not_success(self):
        sys.path.insert(0, str(CLI.parent))
        self.addCleanup(lambda: sys.path.remove(str(CLI.parent)))
        import agent_cli
        for value in ({"message": "done"}, {"ok": "true"}, {"passed": True}, [True]):
            with self.subTest(value=value):
                completed = subprocess.CompletedProcess([], 0, json.dumps(value), "")
                with patch.object(agent_cli.subprocess, "run", return_value=completed):
                    result = agent_cli.execute({"action": "inspect", "source": str(TEMPLATE),
                                                "report": str(self.work / "protocol.json")})
                self.assertFalse(result["ok"])
                self.assertFalse(result["operation_succeeded"])
                self.assertFalse(result["delivery_ready"])
                self.assertEqual(result["error"]["code"], "TOOL_PROTOCOL_ERROR")
                self.assertFalse((self.work / "protocol.json").exists())

    def test_apply_and_verify_keep_content_and_new_paths(self):
        inspected = self.call({"action": "inspect", "source": str(TEMPLATE), "report": "roles.json"})
        data = json.loads(Path(inspected["outputs"]["report"]).read_text(encoding="utf-8"))
        paragraph = next(p for p in data["paragraphs"] if p["text"].strip() and
                         not any(p[key] for key in ("has_math", "has_drawing", "has_field", "in_textbox")))
        mapping = self.work / "reviewed-map.json"
        mapping.write_text(json.dumps({"source_sha256": self.original_hash,
                                      "paragraphs": {paragraph["id"]: {"role": "body", "expected_text": paragraph["text"]}}}), encoding="utf-8")
        request = {"action": "apply_styles", "source": str(TEMPLATE), "output": "formatted.docx",
                   "mapping": str(mapping), "report": "applied.json"}
        result = self.call(request)
        self.assertTrue(result["ok"], result)
        self.assertTrue(result["result"]["content_preserved"])
        self.assertEqual(result["result"]["output"], str(self.work / "formatted.docx"))
        verified = self.call(dict(request, action="verify", report="verified.json"))
        self.assertTrue(verified["ok"], verified)
        self.assertTrue(verified["result"]["content_preserved"])
        output_hash = digest(self.work / "formatted.docx")
        conflict = self.call(dict(request, report="another.json"))
        self.assertEqual(conflict["error"]["code"], "OUTPUT_EXISTS")
        self.assertEqual(digest(self.work / "formatted.docx"), output_hash)

    def test_staging_path_alias_is_replaced_in_result_and_published_report(self):
        from docx import Document
        sys.path.insert(0,str(CLI.parent))
        self.addCleanup(lambda:sys.path.remove(str(CLI.parent)))
        import agent_cli
        source=self.work/'alias-source.docx'
        document=Document();document.add_paragraph('Temporary path alias regression.');document.save(source)
        mapping=self.work/'alias-map.json'
        mapping.write_text(json.dumps({'source_sha256':digest(source),'paragraphs':{'p000001':'body'}}),encoding='utf-8')
        destination=self.work/'alias-formatted.docx';report=self.work/'alias-report.json'
        original_temporary_directory=tempfile.TemporaryDirectory
        roots=[]

        @contextmanager
        def aliased_directory(*args,**kwargs):
            with original_temporary_directory(*args,**kwargs) as directory:
                canonical=Path(directory).resolve()
                (canonical/'alias-anchor').mkdir()
                actual=canonical/'actual-stage';actual.mkdir();roots.append(actual)
                # An actual filesystem alias usable on Linux and Windows,
                # without symlink privileges. The child formatter resolves it,
                # reproducing the 8.3-short-name -> long-name normalization.
                yield str(canonical/'alias-anchor'/'..'/'actual-stage')

        with patch.object(agent_cli.tempfile,'TemporaryDirectory',aliased_directory):
            result=agent_cli.execute({'action':'apply_styles','source':str(source),'output':str(destination),
                                      'mapping':str(mapping),'report':str(report)})
        self.assertTrue(result['ok'],result)
        self.assertTrue(all(not path.exists() for path in roots))
        self.assertEqual(Path(result['result']['output']),destination)
        self.assertTrue(Path(result['result']['output']).is_file(), 'Result must never point at deleted staging output')
        self.assertEqual(Path(result['result']['full_report']),report)
        saved=json.loads(report.read_text(encoding='utf-8'))
        self.assertEqual(Path(saved['output']),destination)
        self.assertTrue(Path(saved['output']).is_file(), 'Published report must never point at deleted staging output')

    def test_citation_layout_stages_and_failed_control_report(self):
        from docx import Document
        source = self.work / "small.docx"
        document = Document()
        document.add_paragraph("人工审核测试标题")
        document.add_paragraph("张三（2020）提出这一观点[1]。")
        document.save(source)
        base_map = self.work / "base-map.json"
        base_map.write_text(json.dumps({"source_sha256": digest(source), "paragraphs": {
            "p000001": {"role": "cn_title", "expected_text": "人工审核测试标题"},
            "p000002": {"role": "body", "expected_text": "张三（2020）提出这一观点[1]。"}}}), encoding="utf-8")
        applied = self.call({"action": "apply_styles", "source": str(source), "output": "base.docx",
                             "mapping": str(base_map), "report": "base-report.json"})
        self.assertTrue(applied["ok"], applied)
        base = self.work / "base.docx"
        citation_map = self.work / "citation-map.json"
        # The fixture's span is explicitly reviewed here, not auto-approved by
        # the adapter from arbitrary candidate matches.
        text = "张三（2020）提出这一观点[1]。"
        mapping = {"schema_version": 1, "reviewed": True, "source_sha256": digest(base),
            "spans": [{"pid": "p000002", "start": text.index("[1]"), "end": text.index("[1]") + 3,
                       "expected_text": "[1]", "role": "body_citation"}]}
        citation_map.write_text(json.dumps(mapping), encoding="utf-8")
        request = {"action": "apply_citations", "source": str(base), "output": "cited.docx",
                   "mapping": str(citation_map), "report": "citation-report.json"}
        self.assertTrue(self.call(request)["ok"])
        self.assertTrue(self.call(dict(request, action="verify", scope="citations", report="citation-verify.json"))["ok"])
        cited = self.work / "cited.docx"
        layout_map = self.work / "layout-map.json"
        layout_map.write_text(json.dumps({"source_sha256": digest(cited),
            "title": {"paragraph_id": "p000001", "expected_text": "人工审核测试标题"},
            "regions": [{"start": "start", "kind": "body"}]}), encoding="utf-8")
        request = {"action": "apply_layout", "source": str(cited), "output": "layout.docx",
                   "mapping": str(layout_map), "report": "layout-report.json"}
        result = self.call(request)
        self.assertTrue(result["ok"], result)
        self.assertTrue(result["result"]["article_content_preserved"])
        self.assertTrue(self.call(dict(request, action="verify", scope="layout", report="layout-verify.json"))["ok"])
        result = self.call({"action": "validate_controls", "source": str(self.work / "layout.docx"), "report": "controls.json"})
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"]["code"], "CHECK_FAILED")
        self.assertTrue(Path(result["outputs"]["report"]).is_file())
        self.assertFalse(result["result"]["passed"])


if __name__ == "__main__":
    unittest.main()

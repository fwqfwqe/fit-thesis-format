"""Regression tests for formatting acceptance, separate from text preservation."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SCRIPTS = Path(__file__).resolve().parents[1] / "fit-thesis-format" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F


class FormatVerificationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="fit-verification-")
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.source = self.work / "source.docx"
        document = Document()
        body = document.add_paragraph()
        body.add_run("正文保留的加粗").bold = True
        body.add_run("和斜体").italic = True
        body.add_run("1").font.superscript = True
        cover = document.add_paragraph("原封面信息")
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover.paragraph_format.line_spacing = Pt(30)
        document.add_paragraph("本轮不处理的另一个段落")
        math = document.add_paragraph()
        omath = F.E.SubElement(math._p, '{'+F.M+'}oMath')
        mr = F.E.SubElement(omath, '{'+F.M+'}r')
        rp = F.E.SubElement(mr, F.q('rPr'))
        F.child(rp, 'rFonts', ascii='Cambria Math', hAnsi='Cambria Math')
        F.child(rp, 'sz', val='28')
        F.E.SubElement(mr, '{'+F.M+'}t').text = 'x'
        document.save(self.source)
        self.original_hash = F.sha(self.source.read_bytes())
        self.addCleanup(lambda: self.assertEqual(F.sha(self.source.read_bytes()), self.original_hash))
        self.mapping = {'source_sha256': self.original_hash,
                        'paragraphs': {'p000001': 'body', 'p000002': 'cover_info', 'p000004': 'equation'}}

    def apply(self, mapping=None, name='applied.docx'):
        path=self.work/name
        result=F.apply_doc(self.source,path,self.mapping if mapping is None else mapping,self.work/(name+'.json'))
        self.assertTrue(result['ok'],result)
        self.assertTrue(result['content_preserved'])
        return path

    def mutate_style(self, source, change):
        parts=F.package(source)
        styles=F.parse(parts['word/styles.xml'])
        change(styles)
        parts['word/styles.xml']=F.xml(styles)
        path=self.work/'mutated.docx'
        with zipfile.ZipFile(source) as zin, zipfile.ZipFile(path,'w') as zout:
            for info in zin.infolist():zout.writestr(info,parts[info.filename])
        return path

    def test_empty_map_rejected_before_output(self):
        destination=self.work/'empty.docx'
        with self.assertRaisesRegex(ValueError,'Empty formatting map'):
            F.apply_doc(self.source,destination,{'source_sha256':self.original_hash,'paragraphs':{}},self.work/'empty.json')
        self.assertFalse(destination.exists())
        self.assertFalse((self.work/'empty.json').exists())

    def test_partial_mapping_and_legal_local_formats_are_preserved(self):
        path=self.apply()
        result=F.verify_doc(path,self.source,self.mapping)
        self.assertTrue(result['ok'],result)
        doc=Document(path)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertTrue(doc.paragraphs[0].runs[1].italic)
        self.assertTrue(doc.paragraphs[0].runs[2].font.superscript)
        self.assertEqual(doc.paragraphs[1].alignment,WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(doc.paragraphs[1].paragraph_format.line_spacing,Pt(30))
        self.assertEqual(F.inventory(F.package(path))['protected'],F.inventory(F.package(self.source))['protected'])
        self.assertEqual(doc.paragraphs[2].text,'本轮不处理的另一个段落')

    def test_50pt_font_and_5pt_fixed_line_do_not_pass_verify(self):
        base=self.apply()
        def change(styles):
            style=styles.find('w:style[@w:styleId="FIT_body"]',F.NS)
            style.find('w:rPr/w:sz',F.NS).set(F.q('val'),'100')
            style.find('w:pPr/w:spacing',F.NS).set(F.q('line'),'100')
        bad=self.mutate_style(base,change)
        result=F.verify_doc(bad,self.source,self.mapping)
        self.assertFalse(result['ok'])
        self.assertTrue(result['content_preserved'])
        self.assertTrue(any('FIT_body/pPr' in error for error in result['errors']))
        self.assertTrue(any('FIT_body/rPr' in error for error in result['errors']))

    def test_unrecorded_font_alignment_indent_and_break_changes_fail(self):
        base=self.apply()
        def change(styles):
            style=styles.find('w:style[@w:styleId="FIT_body"]',F.NS)
            style.find('w:rPr/w:rFonts',F.NS).set(F.q('eastAsia'),'黑体')
            style.find('w:pPr/w:jc',F.NS).set(F.q('val'),'right')
            style.find('w:pPr/w:ind',F.NS).set(F.q('firstLineChars'),'900')
            style.find('w:pPr/w:pageBreakBefore',F.NS).set(F.q('val'),'1')
        result=F.verify_doc(self.mutate_style(base,change),self.source,self.mapping)
        self.assertFalse(result['ok'])
        self.assertTrue(result['content_preserved'])

    def test_reviewed_explicit_exception_must_match_in_both_apply_and_verify(self):
        mapping=copy.deepcopy(self.mapping)
        mapping['style_overrides']={'body':{'reviewed':True,'reason':'该段落需要最小行高，保留可读显示',
                                         'properties':{'line_rule':'atLeast','line_pt':24,'after_pt':6}}}
        path=self.apply(mapping)
        self.assertTrue(F.verify_doc(path,self.source,mapping)['ok'])
        self.assertFalse(F.verify_doc(path,self.source,self.mapping)['ok'])
        styles=F.parse(F.package(path)['word/styles.xml'])
        line=styles.find('w:style[@w:styleId="FIT_body"]/w:pPr/w:spacing',F.NS)
        self.assertEqual(line.get(F.q('line')),'480')
        self.assertEqual(line.get(F.q('lineRule')),'atLeast')

    def test_unreviewed_or_nonlayout_exceptions_rejected(self):
        for exception in (
            {'reviewed':False,'reason':'未审核','properties':{'size_pt':14}},
            {'reviewed':True,'reason':'','properties':{'size_pt':14}},
            {'reviewed':True,'reason':'改大纲','properties':{'outline':0}},
            {'reviewed':True,'reason':'严重裁切','properties':{'size_pt':50,'line_pt':5,'line_rule':'exact'}},
        ):
            with self.subTest(exception=exception):
                mapping=copy.deepcopy(self.mapping)
                mapping['style_overrides']={'body':exception}
                with self.assertRaises(ValueError):F.resolved_style_specs(mapping)

    def test_reviewed_tab_stops_are_applied_and_verified(self):
        mapping=copy.deepcopy(self.mapping)
        mapping['style_overrides']={'body':{'reviewed':True,'reason':'本测试检查明确指定的段落制表位',
            'properties':{'tabs':[{'align':'center','position_twips':2300},
                                  {'align':'right','position_twips':4600,'leader':'dot'}]}}}
        path=self.apply(mapping)
        self.assertTrue(F.verify_doc(path,self.source,mapping)['ok'])
        styles=F.parse(F.package(path)['word/styles.xml'])
        properties=styles.find('w:style[@w:styleId="FIT_body"]/w:pPr',F.NS)
        self.assertLess(list(properties).index(properties.find(F.q('tabs'))),
                        list(properties).index(properties.find(F.q('spacing'))))
        self.assertEqual([(t.get(F.q('val')),t.get(F.q('pos'))) for t in properties.find(F.q('tabs'))],
                         [('center','2300'),('right','4600')])
        self.assertFalse(F.verify_doc(path,self.source,self.mapping)['ok'])


if __name__=='__main__':
    unittest.main()

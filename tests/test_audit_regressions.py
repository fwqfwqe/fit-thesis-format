"""Anonymous structural regressions; no submitted thesis is a test fixture."""
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "fit-thesis-format" / "scripts"))
from format_audit import Audit


def element(tag, **attributes):
    result = OxmlElement(tag)
    for name, value in attributes.items():
        result.set(qn("w:" + name), str(value))
    return result


def styled_document():
    doc = Document()
    for sid, name, level in [("FIT_body", "论文-正文", 9), ("FIT_chapter", "论文-章标题", 0),
                             ("FIT_section", "论文-一级节标题", 1), ("FIT_toc1", "论文-一级目录", 9),
                             ("FIT_toc_heading", "论文-目录标题", 9)]:
        style = doc.styles.add_style(sid, WD_STYLE_TYPE.PARAGRAPH)
        style.name = name
        style.element.get_or_add_pPr().append(element("w:outlineLvl", val=level))
    return doc


def numbered_heading(doc, title="测试章", bound=False):
    numbers = doc.part.numbering_part.element
    abstract = element("w:abstractNum", abstractNumId=42)
    abstract.append(element("w:multiLevelType", val="multilevel"))
    level = element("w:lvl", ilvl=0)
    level.append(element("w:start", val=1))
    level.append(element("w:numFmt", val="chineseCounting"))
    if bound:
        level.append(element("w:pStyle", val="FIT_chapter"))
    level.append(element("w:lvlText", val="%1、"))
    abstract.append(level)
    first_num = next((x for x in numbers if x.tag == qn("w:num")), None)
    numbers.insert(numbers.index(first_num) if first_num is not None else len(numbers), abstract)
    number = element("w:num", numId=42)
    number.append(element("w:abstractNumId", val=42))
    numbers.append(number)
    heading = doc.add_paragraph(title, style="论文-章标题")
    props = heading._p.get_or_add_pPr()
    props.append(element("w:numPr"))
    props[-1].append(element("w:ilvl", val=0))
    props[-1].append(element("w:numId", val=42))
    return heading


class AuditRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="fit-audit-regression-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def audit(self, doc):
        source = self.folder / "anonymous.docx"
        doc.save(source)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result = Audit(source).run()
        self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertTrue(result["source_unchanged"])
        self.assertTrue(result["summary"]["not_a_final_compliance_certificate"])
        return result

    def test_school_name_alone_does_not_certify_original_cover(self):
        doc = styled_document()
        for value in ["福州理工学院", "本科生毕业论文（设计）", "（正本）", "研究题目", "作者：____", "专业：____"]:
            doc.add_paragraph(value)
        doc.add_paragraph("摘要")
        report = self.audit(doc)
        finding = next(x for x in report["issues"] if x["category"] == "front.school_cover_missing_or_unconfirmed")
        self.assertTrue(finding["actual"]["school_name_found"])
        self.assertFalse(finding["actual"]["native_template_cover_media_match"])
        self.assertIn("student_id", finding["actual"]["missing_field_labels"])
        self.assertEqual(finding["actual"]["cover_table_ids"], [])

    def test_original_school_template_cover_remains_recognized(self):
        source = ROOT / "fit-thesis-format" / "assets" / "fit-economics-management-template.docx"
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        report = Audit(source).run()
        self.assertNotIn("front.school_cover_missing_or_unconfirmed", {x["category"] for x in report["issues"]})
        self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())

    def test_fullwidth_inline_equation_is_a_review_candidate(self):
        doc = styled_document()
        doc.add_paragraph("模型关系为y_it＝α＋βx_it，变量含义另述。", style="论文-正文")
        report = self.audit(doc)
        finding = next(x for x in report["issues"] if x["category"] == "math.plain_text_candidate")
        self.assertEqual(finding["status"], "requires_role_review")
        self.assertEqual(finding["location"]["paragraph_id"], "p000001")
        self.assertIn("y_it＝α", finding["actual"]["ordinary_text_math_candidates"])

    def test_underscore_variables_need_mathematical_context(self):
        doc = styled_document()
        doc.add_paragraph("变量x_i表示观测值。", style="论文-正文")
        report = self.audit(doc)
        self.assertIn("math.plain_text_candidate", {x["category"] for x in report["issues"]})

    def test_paths_files_numbers_and_code_are_not_math_candidates(self):
        doc = styled_document()
        for value in ["保存为data_file.csv，地址https://example.org/x_i=3。", "文件在C:\\work\\x_i=3。",
                      "1.23是普通测量数据。", "def calculate(x_i=3):"]:
            doc.add_paragraph(value, style="论文-正文")
        report = self.audit(doc)
        self.assertNotIn("math.plain_text_candidate", {x["category"] for x in report["issues"]})
        self.assertNotIn("heading.not_true_automatic_numbering", {x["category"] for x in report["issues"]})

    def test_operation_and_case_names_are_not_underscore_variables(self):
        doc = styled_document()
        doc.add_paragraph("模型采用COPY_IN与Cache_READ操作，并分析case_001的运行时间。", style="论文-正文")
        report = self.audit(doc)
        self.assertNotIn("math.plain_text_candidate", {x["category"] for x in report["issues"]})

    def test_real_omml_is_not_reported_as_plain_math(self):
        doc = styled_document()
        paragraph = doc.add_paragraph("模型如下。", style="论文-正文")
        math = OxmlElement("m:oMath"); run = OxmlElement("m:r"); token = OxmlElement("m:t")
        token.text = "x_i=y_i"; run.append(token); math.append(run); paragraph._p.append(math)
        report = self.audit(doc)
        self.assertNotIn("math.plain_text_candidate", {x["category"] for x in report["issues"]})

    def test_nonempty_preceding_section_boundary_is_a_new_page(self):
        doc = styled_document()
        prior = doc.add_paragraph("目录最后一个条目")
        section = element("w:sectPr"); section.append(element("w:type", val="nextPage"))
        prior._p.get_or_add_pPr().append(section)
        heading = numbered_heading(doc, bound=True)
        heading._p.get_or_add_pPr().append(element("w:pageBreakBefore", val=0))
        report = self.audit(doc)
        wrong = [x for x in report["issues"] if x["category"] == "paragraph.format" and x["location"].get("paragraph_id") == "p000002" and "new_page" in x["actual"]]
        self.assertEqual(wrong, [])

    def test_appendix_container_label_is_not_a_main_section(self):
        doc = styled_document()
        doc.add_paragraph("附录")
        table = doc.add_table(rows=1, cols=1)
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.style = doc.styles["论文-一级节标题"]
        paragraph.add_run("附录1")
        report = self.audit(doc)
        self.assertEqual(report["role_identification"]["p000002"]["role"], "appendix_subheading")
        self.assertNotIn("heading.not_true_automatic_numbering", {x["category"] for x in report["issues"]})

    def test_direct_numbering_does_not_imply_style_binding(self):
        doc = styled_document(); numbered_heading(doc, bound=False)
        report = self.audit(doc)
        self.assertIn("heading.style_numbering_binding_unconfirmed", {x["category"] for x in report["issues"]})
        self.assertNotIn("heading.not_true_automatic_numbering", {x["category"] for x in report["issues"]})

    def test_bound_numbering_has_no_style_binding_warning(self):
        doc = styled_document(); numbered_heading(doc, bound=True)
        report = self.audit(doc)
        self.assertNotIn("heading.style_numbering_binding_unconfirmed", {x["category"] for x in report["issues"]})

    def test_toc_decimal_cache_disagrees_with_chinese_target(self):
        doc = styled_document()
        doc.add_paragraph("目录", style="论文-目录标题")
        entry = doc.add_paragraph("1. 测试章", style="论文-一级目录")
        reference = element("w:fldSimple", instr=" PAGEREF _TocAnonymous \\h ")
        entry._p.append(reference)
        heading = numbered_heading(doc, bound=True)
        heading._p.insert(1, element("w:bookmarkStart", id=1, name="_TocAnonymous"))
        heading._p.append(element("w:bookmarkEnd", id=1))
        report = self.audit(doc)
        finding = next(x for x in report["issues"] if x["category"] == "toc.cached_number_prefix_mismatch")
        self.assertEqual(finding["location"]["target_paragraph_id"], "p000003")
        self.assertEqual(finding["status"], "requires_content_authorization")


if __name__ == "__main__":
    unittest.main()

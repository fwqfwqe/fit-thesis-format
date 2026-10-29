"""Offline regression checks for preparation and delivery evidence.

All manuscripts, packages, PDFs and evidence are synthetic temporary fixtures.
The repository's retained school template is only read/copied into a temp package.
"""
from contextlib import redirect_stdout
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from PIL import Image, UnidentifiedImageError
import pypdfium2 as pdfium

REPO = Path(__file__).resolve().parents[1]
PACKAGE = REPO / 'fit-thesis-format'
SCRIPTS = PACKAGE / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_format as F
import delivery_check as DELIVERY
import prepare as PREPARE


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')


def load_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def blank_pdf(page_count):
    """Small, valid uncompressed PDF; no non-core PDF writer dependency."""
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               ('<< /Type /Pages /Count %d /Kids [%s] >>' %
                (page_count, ' '.join('%d 0 R' % (index + 3) for index in range(page_count)))).encode()]
    objects += [b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] /Resources << >> >>'] * page_count
    data = bytearray(b'%PDF-1.4\n'); offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(data)); data += ('%d 0 obj\n' % index).encode() + value + b'\nendobj\n'
    start = len(data)
    data += ('xref\n0 %d\n0000000000 65535 f \n' % len(offsets)).encode()
    for offset in offsets[1:]: data += ('%010d 00000 n \n' % offset).encode()
    data += ('trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(offsets), start)).encode()
    return bytes(data)


class PortableEntryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='fit portable 中文 ')
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name).resolve()
        self.source = self.work / 'paper.docx'
        document = Document()
        document.add_paragraph('合成测试论文')
        document.add_paragraph('这是仅用于代码回归的合成正文，不包含真实论文资料。')
        document.save(self.source)
        self.source_hash = sha(self.source)
        self.addCleanup(lambda: self.assertEqual(sha(self.source), self.source_hash))

    def local_package(self):
        target = self.work / 'local-skill'
        shutil.copytree(PACKAGE, target, ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
        # Documentation may be concurrently edited; preparation tests concern
        # the required-entry existence contract, not its prose.
        if not (target / 'START_HERE.md').exists():
            (target / 'START_HERE.md').write_text('Synthetic offline entry fixture.\n', encoding='utf-8')
        return target

    def evidence(self, actual_pages=2, declared_pages=None):
        count = actual_pages if declared_pages is None else declared_pages
        pdf = self.work / 'rendered.pdf'
        pdf.write_bytes(blank_pdf(actual_pages))
        final_hash, pdf_hash = sha(self.source), sha(pdf)
        render = self.work / 'render.json'; pages = self.work / 'pages.json'; review = self.work / 'review.json'
        image_items, review_items = [], []
        for number in range(1, count + 1):
            image = self.work / ('page-%d.png' % number)
            # Use actual PDF rasterization so matching fixture hashes cannot
            # substitute an unrelated image for a page in the evidence.
            with pdfium.PdfDocument(str(pdf)) as document:
                page = document[number - 1]
                bitmap = page.render(scale=1)
                picture = bitmap.to_pil()
                width, height = picture.size
                picture.save(image, format='PNG')
                picture.close(); bitmap.close(); page.close()
            image_items.append({'page': number, 'file': image.name, 'sha256': sha(image), 'width': width, 'height': height})
            review_items.append({'page': number, 'image_sha256': sha(image), 'status': 'pass',
                                 'checks': {check: 'pass' for check in DELIVERY.PAGE_CHECKS},
                                 'notes': 'Synthetic evidence fixture; no real thesis acceptance is claimed.'})
        save_json(render, {'engine': 'synthetic-test-engine', 'engine_version': 'test',
                          'source_sha256_before': final_hash, 'source_unchanged': True,
                          'pdf': pdf.name, 'pdf_sha256': pdf_hash, 'pages': count})
        save_json(pages, {'engine': 'pypdfium2', 'dpi': 72,
                          'input_sha256': pdf_hash, 'page_count': count, 'pages': image_items})
        save_json(review, {'final_docx_sha256': final_hash, 'pdf_sha256': pdf_hash,
                          'reviewer': 'synthetic-test', 'pages': review_items})
        return render, pages, review

    def test_bootstrap_offline_hash_validation_and_no_overwrite(self):
        package = self.local_package()
        payload = io.BytesIO(); entries = []
        with zipfile.ZipFile(payload, 'w', zipfile.ZIP_DEFLATED) as bundle:
            for path in sorted(package.rglob('*')):
                if path.is_file():
                    name = 'fit-thesis-format/' + path.relative_to(package).as_posix()
                    data = path.read_bytes(); bundle.writestr(name, data)
                    entries.append({'path': name, 'size': len(data), 'sha256': hashlib.sha256(data).hexdigest()})
        archive = self.work / 'offline.zip'; archive.write_bytes(payload.getvalue())
        manifest = self.work / 'manifest.json'
        record = {'skill': 'fit-thesis-format', 'version': 'test', 'files': entries,
                  'archive': {'path': 'dist/test.zip', 'size': archive.stat().st_size, 'sha256': sha(archive)}}
        save_json(manifest, record)
        destination = self.work / 'installed'
        args = [sys.executable, str(REPO / 'bootstrap.py'), '--dest', str(destination),
                '--manifest', str(manifest), '--archive', str(archive)]
        result = subprocess.run(args, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertFalse(json.loads(result.stdout)['delivery_ready'])
        installed = destination / 'fit-thesis-format' / 'scripts' / 'agent_cli.py'
        installed_hash = sha(installed)
        repeated = subprocess.run(args, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(repeated.returncode, 2)
        self.assertEqual(sha(installed), installed_hash)
        record['archive']['sha256'] = '0' * 64; save_json(manifest, record)
        broken_destination = self.work / 'bad-install'
        args[args.index('--dest') + 1] = str(broken_destination)
        failed = subprocess.run(args, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(failed.returncode, 2)
        self.assertFalse(broken_destination.exists())

    def test_prepare_rejects_missing_clean_template_before_creating_work(self):
        package = self.local_package()
        (package / 'assets' / 'fit-economics-management-template-clean.docx').unlink()
        work = self.work / 'incomplete-package-job'
        with patch.object(PREPARE, 'ROOT', package):
            with self.assertRaisesRegex(ValueError, 'Incomplete skill package.*template-clean.docx'):
                PREPARE.prepare(self.source, work)
        self.assertFalse(work.exists())

    def test_prepare_is_read_only_and_draft_cannot_be_applied(self):
        package = self.local_package(); target = self.work / 'prepared'
        with patch.object(PREPARE, 'ROOT', package): result = PREPARE.prepare(self.source, target)
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(sha(target / 'original.docx'), self.source_hash)
        mapping = load_json(target / 'style-map.draft.json')
        self.assertEqual(mapping['paragraphs'], {})
        self.assertFalse(mapping['reviewed'])
        request = {'action': 'apply_styles', 'source': str(target / 'original.docx'),
                   'output': str(self.work / 'bad-completion.docx'),
                   'mapping': str(target / 'style-map.draft.json'), 'report': str(self.work / 'bad-report.json')}
        run = subprocess.run([sys.executable, str(SCRIPTS / 'agent_cli.py'), '--request', '-'],
                             input=json.dumps(request), capture_output=True, text=True, encoding='utf-8')
        self.assertNotEqual(run.returncode, 0)
        self.assertFalse(json.loads(run.stdout)['delivery_ready'])
        self.assertFalse((self.work / 'bad-completion.docx').exists())

    def test_prepare_missing_dependency_stops_before_outputs(self):
        package = self.local_package(); target = self.work / 'missing-dependency'
        def invoke(request):
            if request['action'] == 'doctor':
                return {'ok': True, 'operation_succeeded': True, 'result': {'static_cli_ready': False,
                        'dependencies': {'lxml': {'available': False}}}}
            raise ValueError('DEPENDENCY_MISSING: lxml')
        with patch.object(PREPARE, 'ROOT', package), patch.object(PREPARE, 'invoke', side_effect=invoke), \
                patch.object(sys, 'argv', ['prepare.py', '--source', str(self.source), '--work-dir', str(target)]), \
                redirect_stdout(io.StringIO()) as captured:
            self.assertEqual(PREPARE.main(), 2)
        self.assertFalse(json.loads(captured.getvalue())['operation_succeeded'])
        self.assertFalse(target.exists(), 'Missing dependencies must be caught before creating a half-prepared job')

    def test_prepare_step_failure_leaves_explicit_failed_state(self):
        package = self.local_package(); target = self.work / 'failed-job'
        def invoke(request):
            if request['action'] == 'doctor':
                return {'ok': True, 'result': {'static_cli_ready': True}}
            raise ValueError('Synthetic inspect failure')
        with patch.object(PREPARE, 'ROOT', package), patch.object(PREPARE, 'invoke', side_effect=invoke), \
                patch.object(sys, 'argv', ['prepare.py', '--source', str(self.source), '--work-dir', str(target)]), \
                redirect_stdout(io.StringIO()) as captured:
            self.assertEqual(PREPARE.main(), 2)
        self.assertFalse(json.loads(captured.getvalue())['operation_succeeded'])
        if target.exists():
            result_file = target / 'prepare-result.json'
            self.assertTrue(result_file.exists(), 'Preserved partial jobs require an explicit failed result')
            result = load_json(result_file)
            self.assertFalse(result['operation_succeeded'])
            self.assertFalse(result['delivery_ready'])

    def test_delivery_fresh_chain_rejects_text_and_style_corruption(self):
        mapping = self.work / 'map.json'
        data = {'source_sha256': self.source_hash, 'paragraphs': {'p000002': 'body'}}
        save_json(mapping, data)
        base = self.work / 'base.docx'
        F.apply_doc(self.source, base, data, self.work / 'apply.json')
        chain = self.work / 'chain.json'
        save_json(chain, {'original_sha256': self.source_hash, 'stages': [
            {'scope': 'styles', 'source': str(self.source), 'output': str(base), 'mapping': str(mapping)}]})
        self.assertEqual(len(DELIVERY.verify_chain(self.source, base, chain)), 1)
        clean = base.read_bytes()
        for kind in ('font', 'text'):
            with self.subTest(kind=kind):
                with zipfile.ZipFile(io.BytesIO(clean)) as zipped:
                    parts = {name: zipped.read(name) for name in zipped.namelist()}
                if kind == 'font':
                    styles = F.parse(parts['word/styles.xml'])
                    styles.find('w:style[@w:styleId="FIT_body"]/w:rPr/w:sz', F.NS).set(F.q('val'), '100')
                    parts['word/styles.xml'] = F.xml(styles)
                else:
                    root = F.parse(parts['word/document.xml'])
                    next(root.iter(F.q('t'))).text += '篡改'
                    parts['word/document.xml'] = F.xml(root)
                with zipfile.ZipFile(base, 'w') as output:
                    for name, value in parts.items(): output.writestr(name, value)
                with self.assertRaises(ValueError): DELIVERY.verify_chain(self.source, base, chain)

    def test_delivery_accepts_complete_two_page_synthetic_evidence(self):
        render, pages, review = self.evidence(actual_pages=2)
        result = DELIVERY.check_render(self.source, render, pages, review)
        self.assertTrue(result['evidence_consistent'])
        self.assertEqual(result['page_count'], 2)
        self.assertEqual(result['pdf_sha256'], sha(self.work / 'rendered.pdf'))
        self.assertFalse(result['visual_inspection_performed_by_this_script'])

    def test_delivery_rejects_invalid_image_even_with_matching_hashes(self):
        render, pages, review = self.evidence(actual_pages=2)
        image = self.work / 'page-2.png'
        image.write_bytes(b'This file is not a PNG despite its extension.')
        image_manifest = load_json(pages)
        image_manifest['pages'][1]['sha256'] = sha(image)
        save_json(pages, image_manifest)
        page_review = load_json(review)
        page_review['pages'][1]['image_sha256'] = sha(image)
        save_json(review, page_review)
        with self.assertRaisesRegex(UnidentifiedImageError, 'cannot identify image file'):
            DELIVERY.check_render(self.source, render, pages, review)

    def test_delivery_rejects_old_document_and_missing_review_page(self):
        render, pages, review = self.evidence()
        value = load_json(render); value['source_sha256_before'] = '0' * 64; save_json(render, value)
        with self.assertRaisesRegex(ValueError, 'final DOCX'):
            DELIVERY.check_render(self.source, render, pages, review)
        render, pages, review = self.evidence()
        value = load_json(review); value['pages'].pop(); save_json(review, value)
        with self.assertRaisesRegex(ValueError, 'Every page'):
            DELIVERY.check_render(self.source, render, pages, review)

    def test_delivery_fresh_audit_rejects_wrong_toc_depth(self):
        paper=self.work/'wrong-toc.docx'
        document=Document(self.source)
        paragraph=document.add_paragraph()
        field=F.E.SubElement(paragraph._p,F.q('fldSimple'))
        field.set(F.q('instr'),' TOC \\o "1-3" \\h ')
        run=F.E.SubElement(field,F.q('r')); F.E.SubElement(run,F.q('t')).text='错误的三级目录缓存'
        document.save(paper)
        parts=F.package(paper); paragraphs,_=F.indexed(F.parse(parts['word/document.xml']))
        mapping=self.work/'final-map.json'
        save_json(mapping, {'source_sha256': sha(paper), 'paragraphs': {
            pid: {'role': 'toc1' if p.find('.//' + F.q('fldSimple')) is not None else 'body',
                  'expected_text': F.ptext(p)} for pid,p in paragraphs.items()}, 'tables': {}})
        result=DELIVERY.check(paper,paper,final_mapping=mapping)
        self.assertFalse(result['delivery_ready'])
        audit=result['details']['fresh_audit']
        self.assertIn('toc.depth',{item['category'] for item in audit['blocking_findings']})

    def test_delivery_counts_pages_from_actual_pdf_not_only_claims(self):
        render, pages, review = self.evidence(actual_pages=2, declared_pages=1)
        with self.assertRaisesRegex(ValueError, 'actual PDF'):
            DELIVERY.check_render(self.source, render, pages, review)

    def test_json_prepare_entry_and_delivery_needs_review_are_distinct_from_execution_failure(self):
        package=self.local_package(); target=self.work/'json-prepared'
        request={'action':'prepare','source':str(self.source),'work_dir':str(target)}
        arguments=[sys.executable,str(package/'scripts'/'agent_cli.py'),'--request','-']
        run=subprocess.run(arguments,input=json.dumps(request),capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(run.returncode,0,run.stdout+run.stderr)
        result=json.loads(run.stdout)
        self.assertTrue(result['operation_succeeded'])
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(result['status'],'awaiting_mapping_review')
        self.assertEqual(Path(result['outputs']['work_dir']),target)
        request={'action':'delivery_check','source':str(self.source),'output':str(self.source),
                 'report':str(self.work/'delivery-result.json')}
        run=subprocess.run(arguments,input=json.dumps(request),capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(run.returncode,0,run.stdout+run.stderr)
        result=json.loads(run.stdout)
        self.assertTrue(result['ok'])
        self.assertTrue(result['operation_succeeded'])
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(result['tool_exit_code'],3)
        self.assertEqual(result['status'],'needs_review')
        self.assertTrue(result['blocking_items'])
        report=load_json(result['outputs']['report'])
        self.assertEqual(report['final_docx_sha256'],self.source_hash)
        self.assertEqual(report['blocking_items'],result['blocking_items'])
        # JSON envelope adds fields; the actual checker report remains its own
        # full evidence format and hashes, not an edited success summary.
        self.assertIn('details',report)
        self.assertNotIn('protocol_version',report)

    def test_json_prepare_failure_exposes_partial_diagnostic_workspace(self):
        package=self.local_package(); target=self.work/'json-failed'
        invalid=self.work/'invalid.docx'; invalid.write_bytes(b'This is not a DOCX ZIP package.')
        request={'action':'prepare','source':str(invalid),'work_dir':str(target)}
        run=subprocess.run([sys.executable,str(package/'scripts'/'agent_cli.py'),'--request','-'],
                           input=json.dumps(request),capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(run.returncode,2,run.stdout+run.stderr)
        result=json.loads(run.stdout)
        self.assertFalse(result['operation_succeeded'])
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(result['status'],'preparation_failed')
        self.assertEqual(Path(result['outputs']['work_dir']),target)
        stored=load_json(result['outputs']['report'])
        self.assertFalse(stored['operation_succeeded'])
        self.assertIn('original.docx',stored['partial_outputs'])


if __name__ == '__main__':
    unittest.main()

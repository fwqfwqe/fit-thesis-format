"""Audit integration for template review metadata; only anonymous fixtures."""
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'fit-thesis-format' / 'scripts'))
from format_audit import Audit
import template_hygiene as H


class ReviewHygieneAuditTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='fit-review-audit-')
        self.addCleanup(temporary.cleanup)
        self.work = Path(temporary.name)

    def audit(self, document):
        source = self.work / 'anonymous.docx'
        document.save(source)
        before = source.read_bytes()
        result = Audit(source).run()
        self.assertEqual(source.read_bytes(), before)
        return result

    def test_matching_template_comment_is_review_blocker(self):
        document = Document()
        run = document.add_paragraph('Anonymous manuscript.').runs[0]
        note = next(iter(H.template_comment_texts()))
        document.add_comment(run, text=note, author='Anonymous')
        report = self.audit(document)
        self.assertEqual(report['review_hygiene']['comment_count'], 1)
        findings = [i for i in report['issues'] if i['category'] == 'review.template_comment_match']
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]['severity'], 'warning')
        self.assertEqual(findings[0]['status'], 'requires_role_review')

    def test_genuine_manuscript_comment_is_preserved_and_informational(self):
        document = Document()
        run = document.add_paragraph('Anonymous manuscript.').runs[0]
        document.add_comment(run, text='Please clarify the sample definition.', author='Reviewer')
        report = self.audit(document)
        review_issues = [i for i in report['issues'] if i['category'].startswith('review.')]
        self.assertEqual(len(review_issues), 1)
        self.assertEqual(review_issues[0]['category'], 'review.manuscript_comments_preserved')
        self.assertEqual(review_issues[0]['severity'], 'info')
        self.assertEqual(report['review_hygiene']['template_comment_matches'], [])

    def test_four_dangling_comment_overrides_block_audit_without_comment_parts(self):
        document = Document()
        document.add_paragraph('Anonymous manuscript.')
        source = self.work / 'source.docx'
        document.save(source)
        parts = H.read_package(source)
        root = H.parse(parts['[Content_Types].xml'])
        for part, kind in [
            ('comments.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml'),
            ('commentsExtended.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml'),
            ('commentsIds.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml'),
            ('people.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml'),
        ]:
            node = H.E.SubElement(root, '{' + H.CT + '}Override')
            node.set('ContentType', kind)
            node.set('PartName', '/word/' + part)
        parts['[Content_Types].xml'] = H.serialize(root)
        broken = self.work / 'broken.docx'
        with zipfile.ZipFile(broken, 'x') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        before = broken.read_bytes()
        report = Audit(broken).run()
        self.assertEqual(broken.read_bytes(), before)
        inventory = report['review_hygiene']
        self.assertEqual(inventory['comment_count'], 0)
        self.assertEqual(inventory['comment_parts'], [])
        self.assertEqual(len(inventory['dangling_comment_content_types']), 4)
        self.assertTrue(any(i['category'] == 'review.dangling_comment_metadata' and
                            i['severity'] == 'error' for i in report['issues']))

    def test_field_instruction_does_not_create_a_revision_issue(self):
        document = Document()
        run = document.add_paragraph().add_run()
        instruction = OxmlElement('w:instrText')
        instruction.text = ' PAGE '
        run._r.append(instruction)
        report = self.audit(document)
        self.assertFalse(any(report['review_hygiene']['revisions'].values()))
        self.assertFalse(any(i['category'] == 'review.tracked_revisions_preserved'
                             for i in report['issues']))


if __name__ == '__main__':
    unittest.main()

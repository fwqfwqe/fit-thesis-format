"""Read-only citation review catches missing marks and split-run baseline errors."""
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
from format_audit import Audit


class NarrativeAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='fit-citation-audit-')
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'synthetic.docx'

    def document(self):
        document = Document()
        style = document.styles.add_style('FIT_body', WD_STYLE_TYPE.PARAGRAPH)
        style.name = '论文-正文'
        return document

    def audit(self, document, references=()):
        if references:
            document.add_paragraph('参考文献')
            for reference in references:
                document.add_paragraph(reference)
        document.save(self.source)
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        result = Audit(self.source).run()
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())
        return result

    def citations(self, result):
        return [i for i in result['issues'] if i['category'].startswith('citation.')]

    def test_missing_marks_and_one_two_three_author_review(self):
        doc = self.document()
        for body in ['张甲等（2020）指出合成论述甲。', '李乙等（2021）指出合成论述乙。',
                     '黄丙、赵丁、钱戊（2022）指出合成论述丙。']:
            doc.add_paragraph(body, '论文-正文')
        result = self.audit(doc, ['[1] 张甲.合成题名[J].合成期刊,2020,1:1-2.',
                                 '[2] 李乙,孙己.合成题名[J].合成期刊,2021,1:1-2.',
                                 '[3] 黄丙,赵丁,钱戊.合成题名[J].合成期刊,2022,1:1-2.'])
        issues = self.citations(result)
        for category in ['citation.missing_numeric_marker', 'citation.author_count_needs_review']:
            found = [i for i in issues if i['category'] == category]
            self.assertEqual(len(found), 3, found)
            self.assertTrue(all(i['status'] == 'requires_role_review' for i in found))
        self.assertFalse(result['citation_review']['automatic_reference_assignment'])

    def test_split_year_and_author_position_are_detected(self):
        doc = self.document()
        p = doc.add_paragraph(style='论文-正文')
        author = p.add_run('张甲')
        position = OxmlElement('w:position'); position.set(qn('w:val'), '4')
        author._r.get_or_add_rPr().append(position)
        for fragment in ['（', '20', '20', '）']:
            p.add_run(fragment).font.superscript = True
        p.add_run('指出合成论述')
        for fragment in ['[', '1', ']']:
            p.add_run(fragment).font.superscript = True
        p.add_run('。')
        result = self.audit(doc, ['[1] 张甲.合成题名[J].合成期刊,2020,1:1-2.'])
        categories = {i['category'] for i in self.citations(result)}
        self.assertIn('citation.year_not_baseline', categories)
        self.assertIn('citation.author_year_not_baseline', categories)
        self.assertNotIn('citation.number_position', categories)

    def test_correct_baseline_and_complete_superscript_have_no_citation_warning(self):
        doc = self.document()
        p = doc.add_paragraph('张甲（2020）指出合成论述', '论文-正文')
        for fragment in ['[', '1', ']']:
            p.add_run(fragment).font.superscript = True
        p.add_run('。')
        result = self.audit(doc, ['[1] 张甲.合成题名[J].合成期刊,2020,1:1-2.'])
        self.assertEqual(self.citations(result), [])

    def test_existing_wrong_number_requires_review(self):
        doc = self.document()
        p = doc.add_paragraph('张甲（2020）指出合成论述', '论文-正文')
        p.add_run('[9]').font.superscript = True; p.add_run('。')
        result = self.audit(doc, ['[1] 张甲.合成题名[J].合成期刊,2020,1:1-2.'])
        self.assertIn('citation.numeric_marker_needs_review', {i['category'] for i in self.citations(result)})

    def test_ambiguous_same_author_year_is_not_automatically_assigned(self):
        doc = self.document(); doc.add_paragraph('张甲（2020）指出合成论述。', '论文-正文')
        result = self.audit(doc, ['[1] 张甲.合成题名甲[J].合成期刊,2020,1:1-2.',
                                 '[2] 张甲.合成题名乙[J].合成期刊,2020,1:3-4.'])
        found = [i for i in self.citations(result) if i['category'] == 'citation.reference_ambiguous']
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]['actual']['reference_pids'], ['p000003', 'p000004'])

    def test_ordinary_data_year_is_not_assigned_a_citation(self):
        doc = self.document(); doc.add_paragraph('2020年样本数为100，2021年为120。', '论文-正文')
        result = self.audit(doc, ['[1] 张甲.合成题名[J].合成期刊,2020,1:1-2.'])
        self.assertEqual(self.citations(result), [])

    def test_bibliography_three_full_is_valid_four_full_requires_review(self):
        doc = self.document()
        result = self.audit(doc, ['[1] 张甲,李乙,黄丙.合成题名甲[J].合成期刊,2020,1:1-2.',
                                 '[2] 张甲,李乙,黄丙,孙己.合成题名乙[J].合成期刊,2021,1:1-2.',
                                 '[3] 张甲,李乙,黄丙,等.合成题名丙[J].合成期刊,2022,1:1-2.'])
        found = [i for i in result['issues'] if i['category'] == 'reference.author_listing_needs_review']
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]['location']['paragraph_id'], 'p000003')
        self.assertEqual(found[0]['actual']['listed_author_count'], 4)

    def test_bibliography_short_truncation_and_Chinese_separator_are_reviewed(self):
        doc = self.document()
        result = self.audit(doc, ['[1] 张甲,等.合成题名甲[J].合成期刊,2020,1:1-2.',
                                 '[2] 张甲，李乙.合成题名乙[J].合成期刊,2021,1:1-2.'])
        found = [i for i in result['issues'] if i['category'] == 'reference.author_listing_needs_review']
        self.assertEqual(len(found), 2)
        self.assertIn('truncation_does_not_follow_three_listed_authors', found[0]['actual']['problems'])
        self.assertIn('author_separator_is_not_an_English_comma', found[1]['actual']['problems'])

    def test_foreign_two_uses_Chinese_he_three_uses_et_al(self):
        doc = self.document()
        for value, number in [('Smith和Brown（2020）指出合成论述', 1),
                              ('Smith与Brown（2020）指出合成论述', 1),
                              ('Jones et al.（2021）指出合成论述', 2),
                              ('Jones等（2021）指出合成论述', 2)]:
            p = doc.add_paragraph(value, '论文-正文')
            p.add_run(f'[{number}]').font.superscript = True; p.add_run('。')
        result = self.audit(doc, ['[1] Smith A,Brown B. Synthetic title[J]. Journal,2020,1:1-2.',
                                 '[2] Jones A,Clark B,Davis C. Synthetic title[J]. Journal,2021,1:1-2.'])
        found = [i for i in self.citations(result) if i['category'] == 'citation.author_count_needs_review']
        self.assertEqual({i['location']['paragraph_id'] for i in found}, {'p000002', 'p000004'})


if __name__ == '__main__':
    unittest.main()

"""Reviewed narrative citations cannot become a general text-replacement channel."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from lxml import etree as E

SCRIPTS = Path(__file__).resolve().parents[1] / 'fit-thesis-format' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import citation_style as C
import narrative_citations as N


class NarrativeCitationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='fit-narrative-中文-')
        self.addCleanup(self.temporary.cleanup)
        self.work = Path(self.temporary.name)
        self.serial = 0

    def source(self, body, references, runs=None, extra=None):
        self.serial += 1
        path = self.work / ('source-' + str(self.serial) + '.docx')
        doc = Document()
        if runs:
            paragraph = doc.add_paragraph()
            for text, options in runs:
                run = paragraph.add_run(text)
                for name, value in options.items():
                    setattr(run.font, name, value)
        else:
            doc.add_paragraph(body)
        if extra:
            extra(doc)
        doc.add_paragraph('参考文献')
        for reference in references:
            doc.add_paragraph(reference)
        doc.save(path)
        return path

    def map(self, source, source_author, names, author_area=None, language='zh', year='2020', number=1,
            connector=None, abbreviation=None, reference_pid=None, claim_end=None):
        paragraphs = C.indexed(C.parse(C.package(source)['word/document.xml']))
        body = C.ptext(paragraphs['p000001'])
        if reference_pid is None:
            reference_pid = next(pid for pid, paragraph in paragraphs.items()
                                 if C.ptext(paragraph).startswith('[' + str(number) + ']'))
        reference = C.ptext(paragraphs[reference_pid])
        reference_start = N.NUMBER.match(reference).end()
        if author_area is None:
            author_area = ','.join(n[0] if isinstance(n, tuple) else n for n in names)
        author_end = reference_start + len(author_area)
        authors, cursor = [], reference_start
        for value in names:
            exact, display = value if isinstance(value, tuple) else (value, value)
            start = reference.index(exact, cursor)
            end = start + len(exact)
            authors.append({'start': start, 'end': end, 'exact_name': exact, 'display_name': display, 'language': language})
            cursor = end
        author_start = body.index(source_author)
        year_start = body.index(year, author_start + len(source_author))
        year_end = year_start + len(year)
        ay_end = year_end + (1 if body[year_end:year_end+1] in {'）', ')'} else 0)
        year_ref_start = reference.index(year, author_end)

        def span(text, start, end):
            return {'start': start, 'end': end, 'expected_text': text[start:end]}

        author_year = span(body, author_start, ay_end)
        author_year.update(author=span(body, author_start, author_start + len(source_author)),
                           year=span(body, year_start, year_end))
        record = {'pid': 'p000001', 'author_year': author_year,
                  'claim': span(body, author_start, len(body) - 1 if claim_end is None else claim_end),
                  'reference': {'pid': reference_pid, 'expected_text': reference, 'reference_number': number,
                                'author_list': span(reference, reference_start, author_end), 'authors': authors,
                                'year': span(reference, year_ref_start, year_ref_start + len(year))}}
        if connector is not None:
            record['connector'] = connector
        if abbreviation is not None:
            record['abbreviation'] = abbreviation
        return {'schema_version': 1, 'reviewed': True, 'source_sha256': C.sha(source.read_bytes()), 'records': [record]}

    def output(self, source, mapping):
        path = self.work / ('output-' + str(self.serial) + '.docx')
        before = source.read_bytes()
        result = N.apply(source, path, mapping)
        self.assertTrue(result['ok'], result)
        self.assertEqual(source.read_bytes(), before)
        self.assertTrue(N.verify(source, path, mapping)['ok'])
        return path

    def text(self, path, pid='p000001'):
        return C.ptext(C.indexed(C.parse(C.package(path)['word/document.xml']))[pid])

    def mutate(self, path, mutate):
        parts = C.package(path)
        mutate(parts)
        target = self.work / ('modified-' + str(self.serial) + '-' + str(len(list(self.work.iterdir()))) + '.docx')
        with zipfile.ZipFile(target, 'w') as zipped:
            for name, data in parts.items():
                zipped.writestr(name, data)
        return target

    def assert_rejected(self, source, mapping, pattern):
        path = self.work / ('rejected-' + str(self.serial) + '.docx')
        with self.assertRaisesRegex(ValueError, pattern):
            N.apply(source, path, mapping)
        self.assertFalse(path.exists())

    def test_chinese_one_two_three_and_four_listed_authors(self):
        cases = [(['张三'], '张三', None, None, '张三'),
                 (['张三', '李四'], '张三与李四', '和', None, '张三和李四'),
                 (['张三', '李四', '王五'], '张三、李四、王五', None, '等', '张三等'),
                 (['张三', '李四', '王五', '赵六'], '张三、李四、王五、赵六', None, '等', '张三等')]
        for names, source_author, connector, abbreviation, expected_author in cases:
            with self.subTest(count=len(names)):
                reference = '[1] ' + ','.join(names) + '. 研究成果[J]. 学报,2020:1-9.'
                source = self.source(source_author + '（2020）提出结论。', [reference])
                mapping = self.map(source, source_author, names, connector=connector, abbreviation=abbreviation)
                output = self.output(source, mapping)
                self.assertEqual(self.text(output), expected_author + '（2020）提出结论[1]。')
                self.assertEqual(self.text(output, 'p000003'), reference)

    def test_foreign_one_two_three_and_four_listed_authors(self):
        names = [('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown'), ('Green S', 'Green')]
        for count in range(1, 5):
            with self.subTest(count=count):
                selected = names[:count]
                source_author = ' and '.join(name[0] for name in selected)
                area = ','.join(name[0] for name in selected)
                reference = '[1] ' + area + '. Research[J]. Journal,2020:1-9.'
                source = self.source(source_author + ' (2020) reports results.', [reference])
                mapping = self.map(source, source_author, selected, language='en',
                                   connector='和' if count == 2 else None,
                                   abbreviation='et al.' if count >= 3 else None)
                output = self.output(source, mapping)
                author = 'Smith' if count == 1 else 'Smith和Jones' if count == 2 else 'Smith et al.'
                self.assertEqual(self.text(output), author + ' (2020) reports results[1].')
                self.assertEqual(self.text(output, 'p000003'), reference)

    def test_actual_three_author_bibliography_and_truncated_four_plus_unchanged(self):
        cases = [(['张三', '李四', '王五'], '张三,李四,王五', '张三、李四、王五', 'zh', '等'),
                 (['张三', '李四', '王五'], '张三,李四,王五,等', '张三等', 'zh', '等'),
                 ([('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')],
                  'Smith J,Jones K,Brown R', 'Smith and Jones and Brown', 'en', 'et al.'),
                 ([('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')],
                  'Smith J,Jones K,Brown R,et al.', 'Smith et al.', 'en', 'et al.')]
        for names, area, author, language, abbreviation in cases:
            with self.subTest(area=area):
                reference = '[1] ' + area + (' Research[J]. Journal,2020.' if area.endswith('al.') else '. Research[J]. Journal,2020.')
                body = author + '（2020）论述结论。'
                source = self.source(body, [reference])
                mapping = self.map(source, author, names, author_area=area, language=language, abbreviation=abbreviation)
                output = self.output(source, mapping)
                self.assertEqual(self.text(output, 'p000003'), reference)
                self.assertEqual(self.text(output), ('张三等' if language == 'zh' else 'Smith et al.') + '（2020）论述结论[1]。')

    def test_wrong_abbreviation_can_be_corrected_from_exact_reference(self):
        cases = [(['张三'], '张三等', {}, '张三'),
                 (['张三', '李四'], '张三等', {'connector': '和'}, '张三和李四'),
                 (['张三', '李四', '王五'], '张三', {'abbreviation': '等'}, '张三等'),
                 ([('Smith J', 'Smith')], 'Smith et al.', {'language': 'en'}, 'Smith'),
                 ([('Smith J', 'Smith'), ('Jones K', 'Jones')], 'Smith et al.', {'language': 'en', 'connector': '和'}, 'Smith和Jones'),
                 ([('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')], 'Smith', {'language': 'en', 'abbreviation': 'et al.'}, 'Smith et al.')]
        for names, author, options, desired in cases:
            with self.subTest(author=author, names=names):
                area = ','.join(value[0] if isinstance(value, tuple) else value for value in names)
                source = self.source(author + '（2020）提出结论。', ['[1] ' + area + '. Results[J]. Journal,2020.'])
                output = self.output(source, self.map(source, author, names, **options))
                self.assertEqual(self.text(output), desired + '（2020）提出结论[1]。')

    def test_exact_english_initial_punctuation_is_supported(self):
        source = self.source('Smith and Jones (2020) report results.',
                             ['[1] Smith J., Jones K. Research[J]. Journal,2020.'])
        mapping = self.map(source, 'Smith and Jones', [('Smith J.', 'Smith'), ('Jones K', 'Jones')],
                           author_area='Smith J., Jones K', language='en', connector='和')
        output = self.output(source, mapping)
        self.assertEqual(self.text(output), 'Smith和Jones (2020) report results[1].')

    def test_foreign_source_chinese_connector_and_abbreviation_are_normalized_to_current_targets(self):
        source = self.source('Acemoglu与Restrepo（2020）提出结论。',
                             ['[1] Acemoglu D,Restrepo P. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, 'Acemoglu与Restrepo', [('Acemoglu D', 'Acemoglu'), ('Restrepo P', 'Restrepo')],
                           language='en', connector='和')
        output = self.output(source, mapping)
        self.assertEqual(self.text(output), 'Acemoglu和Restrepo（2020）提出结论[1]。')
        source = self.source('Smith等（2020）提出结论[1]。', ['[1] Smith J,Jones K,Brown R. Research[J]. Journal,2020.'])
        mapping = self.map(source, 'Smith等', [('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')],
                           language='en', abbreviation='et al.')
        output = self.output(source, mapping)
        self.assertEqual(self.text(output), 'Smith et al.（2020）提出结论[1]。')
        self.assertEqual(N.inspect(source)['candidates'][0]['issues'], ['author_count_needs_review'])
        self.assertEqual(len(N.inspect(source)['matches']), 1)
        self.assertEqual(N.inspect(output)['candidates'], [])

    def test_targets_conflicting_with_current_user_standard_are_rejected(self):
        for language, names, author in [('zh', ['张三', '李四'], '张三与李四'),
                                        ('en', [('Smith J', 'Smith'), ('Jones K', 'Jones')], 'Smith and Jones')]:
            area = ','.join(name[0] if isinstance(name, tuple) else name for name in names)
            for old_connector in ('与', '、', 'and'):
                with self.subTest(language=language, old_connector=old_connector):
                    source = self.source(author + '（2020）提出结论。', ['[1] ' + area + '. Research[J]. Journal,2020.'])
                    mapping = self.map(source, author, names, language=language, connector=old_connector)
                    self.assert_rejected(source, mapping, 'target connector 和')
        for language, names, author, old_abbreviations in [
                ('zh', ['张三', '李四', '王五'], '张三等人', ('等人', 'et al.')),
                ('en', [('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')], 'Smith等', ('等', '等人'))]:
            area = ','.join(name[0] if isinstance(name, tuple) else name for name in names)
            for old_abbreviation in old_abbreviations:
                with self.subTest(language=language, old_abbreviation=old_abbreviation):
                    source = self.source(author + '（2020）提出结论。', ['[1] ' + area + '. Research[J]. Journal,2020.'])
                    mapping = self.map(source, author, names, language=language, abbreviation=old_abbreviation)
                    self.assert_rejected(source, mapping, 'target abbreviation')

    def test_inspection_reports_current_targets_and_warns_about_old_source_display(self):
        cases = [('zh', ['张三', '李四'], '张三与李四', '张三和李四', {'connector': '和'}),
                 ('en', [('Smith J', 'Smith'), ('Jones K', 'Jones')], 'Smith and Jones', 'Smith和Jones', {'connector': '和'}),
                 ('zh', ['张三', '李四', '王五'], '张三等人', '张三等', {'abbreviation': '等'}),
                 ('en', [('Smith J', 'Smith'), ('Jones K', 'Jones'), ('Brown R', 'Brown')], 'Smith等', 'Smith et al.', {'abbreviation': 'et al.'})]
        for language, names, old_author, expected_author, target in cases:
            with self.subTest(old_author=old_author):
                area = ','.join(name[0] if isinstance(name, tuple) else name for name in names)
                source = self.source(old_author + '（2020）提出结论[1]。', ['[1] ' + area + '. Research[J]. Journal,2020.'])
                inspected = N.inspect(source)
                self.assertEqual(inspected['expected_conventions'], {
                    'two_author_connector': '和', 'three_or_more_author_abbreviations': {'zh': '等', 'en': 'et al.'}})
                self.assertIn('author_count_needs_review', inspected['candidates'][0]['issues'])
                output = self.output(source, self.map(source, old_author, names, language=language, **target))
                self.assertEqual(self.text(output), expected_author + '（2020）提出结论[1]。')
                self.assertEqual(N.inspect(output)['candidates'], [])

    def test_existing_marker_across_runs_is_not_duplicated_and_year_is_baseline(self):
        runs = [('张三（', {}), ('2020', {'superscript': True}), ('）提出结论', {}),
                ('[', {}), ('1', {}), (']', {}), ('。末尾正文原样。', {})]
        body = ''.join(value for value, _ in runs)
        source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'], runs=runs)
        claim_end = body.index('。')
        mapping = self.map(source, '张三', ['张三'], claim_end=claim_end)
        output = self.output(source, mapping)
        self.assertEqual(self.text(output), body)
        paragraph = C.indexed(C.parse(C.package(output)['word/document.xml']))['p000001']
        for run in paragraph.findall(C.q('r')):
            text = C.run_text(run)
            if text in {'[', '1', ']'}:
                self.assertEqual(run.find('w:rPr/w:rStyle', C.NS).get(C.q('val')), C.STYLE_ID)
            if text == '2020':
                self.assertEqual(run.find('w:rPr/w:vertAlign', C.NS).get(C.q('val')), 'baseline')
        self.assertEqual(self.text(output).count('[1]'), 1)

    def test_marker_directly_after_year_or_before_sentence_punctuation(self):
        for body, claim_end, expected in [('张三（2020）。', 8, '张三（2020）[1]。'),
                                         ('张三（2020）提出结论。', 12, '张三（2020）提出结论[1]。')]:
            with self.subTest(body=body, claim_end=claim_end):
                source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'])
                mapping = self.map(source, '张三', ['张三'], claim_end=claim_end)
                output = self.output(source, mapping)
                self.assertEqual(self.text(output), expected)

    def test_missing_marker_requires_the_complete_claim_but_existing_year_marker_is_preserved(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        self.assert_rejected(source, self.map(source, '张三', ['张三'], claim_end=8), 'complete cited claim')
        body = '张三（2020）[1]提出结论。'
        source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'])
        output = self.output(source, self.map(source, '张三', ['张三'], claim_end=8))
        self.assertEqual(self.text(output), body)

    def test_missing_marker_cannot_be_inserted_after_sentence_punctuation(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        self.assert_rejected(source, self.map(source, '张三', ['张三'], claim_end=13), 'claim.end excludes punctuation')

    def test_year_suffix_and_explicit_same_year_reference_binding(self):
        source = self.source('张三（2020a）提出结论。', ['[1] 张三. 第一篇[J]. 学报,2020a.', '[2] 张三. 第二篇[J]. 学报,2020b.'])
        output = self.output(source, self.map(source, '张三', ['张三'], year='2020a'))
        self.assertEqual(self.text(output), '张三（2020a）提出结论[1]。')
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 第一篇[J]. 学报,2020.', '[2] 张三. 第二篇[J]. 学报,2020.'])
        result = N.inspect(source)
        self.assertEqual(len(result['candidates']), 1)
        candidate = result['candidates'][0]
        self.assertEqual(candidate['status'], 'reference_ambiguous')
        self.assertIn('missing_numeric_marker', candidate['issues'])
        self.assertEqual(candidate['reference_pids'], ['p000003', 'p000004'])
        self.assertFalse(candidate['reviewed'])
        # Inspection never accepts ambiguity; the reviewer may explicitly bind
        # the claim to the exact PID/full-text of the second actual reference.
        output = self.output(source, self.map(source, '张三', ['张三'], number=2))
        self.assertEqual(self.text(output), '张三（2020）提出结论[2]。')

    def test_inspection_missing_number_and_count_issues_are_advisory(self):
        source = self.source('张三、李四、王五（2020）提出结论。样本收集于2020年。',
                             ['[1] 张三,李四,王五. 研究成果[J]. 学报,2020.'])
        result = N.inspect(source)
        self.assertEqual(len(result['references']), 1)
        self.assertEqual(len(result['candidates']), 1)
        self.assertEqual(set(result['candidates'][0]['issues']), {'missing_numeric_marker', 'author_count_needs_review'})
        self.assertEqual(N.inspect_parts(C.package(source))['candidates'], result['candidates'])
        self.assertNotIn('records', result)

    def test_inspection_one_two_author_misdisplay_and_wrong_marker(self):
        for names, author in [(['张三'], '张三等'), (['张三', '李四'], '张三等'),
                              (['张三', '李四'], '张三')]:
            with self.subTest(names=names, author=author):
                source = self.source(author + '（2020）提出结论[1]。', ['[1] ' + ','.join(names) + '. 研究成果[J]. 学报,2020.'])
                result = N.inspect(source)
                self.assertEqual(len(result['candidates']), 1)
                self.assertEqual(result['candidates'][0]['issues'], ['author_count_needs_review'])
        for marker in ('[9]', '[1,9]', '[1][9]'):
            source = self.source('张三（2020）提出结论' + marker + '。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
            self.assertIn('numeric_marker_needs_review', N.inspect(source)['candidates'][0]['issues'])

    def test_inspection_matches_includes_correct_author_year_for_baseline_audit(self):
        source = self.source('张三（2020）提出结论[1]。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        result = N.inspect(source)
        self.assertEqual(result['candidates'], [])
        self.assertEqual(len(result['matches']), 1)
        match = result['matches'][0]
        self.assertEqual(match['expected_text'], '张三（2020）')
        self.assertEqual(match['issues'], [])
        self.assertEqual(match['author_start'], 0)
        self.assertEqual(match['author_end'], 2)
        self.assertEqual(match['year_start'], 3)
        self.assertEqual(match['year_end'], 7)

    def test_arbitrary_author_claim_year_and_reference_changes_are_rejected(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        original = self.map(source, '张三', ['张三'])
        mutations = [
            (lambda m: m.update(reviewed=False), 'reviewed=true'),
            (lambda m: m.update(source_sha256='0' * 64), 'stale'),
            (lambda m: m.update(allow_changes=True), 'arbitrary changes'),
            (lambda m: m['records'][0]['reference'].update(reference_number=2), 'Reference number'),
            (lambda m: m['records'][0]['reference'].update(expected_text='[1] 李四. 假文献,2020.'), 'full text mismatch'),
            (lambda m: m['records'][0]['reference']['authors'][0].update(display_name='李四'), 'Chinese author display'),
            (lambda m: m['records'][0]['reference']['authors'][0].update(exact_name='李四'), 'exact_name mismatch'),
            (lambda m: m['records'][0]['reference']['year'].update(expected_text='2021'), 'expected text mismatch'),
            (lambda m: m['records'][0]['claim'].update(expected_text='张三（2020）改变结论'), 'expected text mismatch'),
            (lambda m: m['records'][0]['author_year']['author'].update(expected_text='李四'), 'expected text mismatch'),
        ]
        for mutate, error in mutations:
            with self.subTest(error=error):
                mapping = copy.deepcopy(original)
                mutate(mapping)
                self.assert_rejected(source, mapping, error)

    def test_source_year_or_unknown_author_cannot_be_rebound(self):
        source = self.source('张三（2021）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020,2021.'])
        mapping = self.map(source, '张三', ['张三'], year='2021')
        reference = mapping['records'][0]['reference']
        reference['year']['start'] = reference['expected_text'].index('2020')
        reference['year']['end'] = reference['year']['start'] + 4
        reference['year']['expected_text'] = '2020'
        self.assert_rejected(source, mapping, 'Body year')
        source = self.source('李四（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        self.assert_rejected(source, self.map(source, '李四', ['张三']), 'Source authors')
        source = self.source('张三和陌生人（2020）提出结论。', ['[1] 张三,李四. 研究成果[J]. 学报,2020.'])
        self.assert_rejected(source, self.map(source, '张三和陌生人', ['张三', '李四'], connector='和'), 'Source authors')

    def test_author_list_cannot_be_shortened_or_reordered(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三,李四,王五. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, '张三', ['张三', '李四', '王五'], abbreviation='等')
        mapping['records'][0]['reference']['authors'].pop()
        self.assert_rejected(source, mapping, 'unproved name')
        mapping = self.map(source, '张三', ['张三'], author_area='张三')
        self.assert_rejected(source, mapping, 'bibliography sentence delimiter')
        mapping = self.map(source, '张三', ['张三', '李四', '王五'], abbreviation='等')
        mapping['records'][0]['reference']['authors'].reverse()
        self.assert_rejected(source, mapping, 'ordered exact spans')
        source = self.source('张三（2020）提出结论。', ['[1] 张三李四. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, '张三', ['张三', '李四'], author_area='张三李四', connector='和')
        self.assert_rejected(source, mapping, 'actual bibliography author delimiters')

    def test_truncated_bibliography_does_not_evidence_unknown_fourth_name(self):
        source = self.source('张三、李四、王五、赵六（2020）提出结论。', ['[1] 张三,李四,王五,等. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, '张三、李四、王五、赵六', ['张三', '李四', '王五'], author_area='张三,李四,王五,等', abbreviation='等')
        self.assert_rejected(source, mapping, 'Source authors')

    def test_foreign_surname_display_is_not_an_arbitrary_prefix(self):
        source = self.source('Smith (2020) reports results.', ['[1] Smith J. Research[J]. Journal,2020.'])
        mapping = self.map(source, 'Smith', [('Smith J', 'Smith')], language='en')
        mapping['records'][0]['reference']['authors'][0]['display_name'] = 'Smi'
        self.assert_rejected(source, mapping, 'surname prefix')
        mapping['records'][0]['reference']['authors'][0]['display_name'] = 'Smith J'
        self.assert_rejected(source, mapping, 'surname prefix')

    def test_year_only_is_not_a_channel_for_replacing_authors_or_years(self):
        source = self.source('样本于2020年收集。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        # Construct exact offsets deliberately without a real author-year span.
        paragraphs = C.indexed(C.parse(C.package(source)['word/document.xml']))
        body = C.ptext(paragraphs['p000001'])
        source2 = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source2, '张三', ['张三'])
        mapping['source_sha256'] = C.sha(source.read_bytes())
        record = mapping['records'][0]
        record['author_year'] = {'start': 3, 'end': 7, 'expected_text': '2020',
                                 'author': {'start': 3, 'end': 7, 'expected_text': '2020'},
                                 'year': {'start': 3, 'end': 7, 'expected_text': '2020'}}
        record['claim'] = {'start': 0, 'end': len(body)-1, 'expected_text': body[:-1]}
        self.assert_rejected(source, mapping, 'author_year must start')
        self.assertEqual(N.inspect(source)['candidates'], [])

    def test_different_grouped_and_adjacent_markers_are_rejected(self):
        for token in ('[2]', '[1,2]', '[1][2]', '[1] [1]'):
            with self.subTest(token=token):
                source = self.source('张三（2020）提出结论' + token + '。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
                self.assert_rejected(source, self.map(source, '张三', ['张三']), 'numeric marker')

    def test_noncanonical_numeric_markers_are_not_mistaken_for_missing_markers(self):
        for token in ('[ 1 ]', '[1 ]', '[ 1]', '［1］', '［1］[1]', '[1 2]'):
            with self.subTest(token=token):
                source = self.source('张三（2020）提出结论' + token + '。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
                self.assert_rejected(source, self.map(source, '张三', ['张三']), 'numeric marker')
                candidate = N.inspect(source)['candidates'][0]
                self.assertIn('numeric_marker_needs_review', candidate['issues'])
                self.assertNotIn('missing_numeric_marker', candidate['issues'])

    def test_existing_marker_after_sentence_punctuation_is_never_duplicated(self):
        for token in ('[1]', '[ 1 ]', '［1］'):
            with self.subTest(token=token):
                body = '张三（2020）提出结论。' + token
                source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'])
                mapping = self.map(source, '张三', ['张三'], claim_end=body.index('。'))
                self.assert_rejected(source, mapping, 'after sentence punctuation')
                candidate = N.inspect(source)['candidates'][0]
                self.assertIn('numeric_marker_needs_review', candidate['issues'])
                self.assertNotIn('missing_numeric_marker', candidate['issues'])

    def test_year_endpoint_cannot_duplicate_a_later_existing_marker(self):
        for body in ('张三（2020）提出结论[1]。', '张三（2020）提出结论。[1]', '张三（2020）提出结论。 [ 1 ]'):
            with self.subTest(body=body):
                source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'])
                self.assert_rejected(source, self.map(source, '张三', ['张三'], claim_end=8), 'outside the reviewed claim')

    def test_claim_cannot_end_inside_an_existing_marker(self):
        body = '张三（2020）提出结论[1]。'
        source = self.source(body, ['[1] 张三. 研究成果[J]. 学报,2020.'])
        self.assert_rejected(source, self.map(source, '张三', ['张三'], claim_end=body.index('[1]')+2), 'ends inside')

    def test_bookmark_field_revision_math_and_media_paragraphs_are_protected(self):
        for tag in ('bookmarkStart', 'fldSimple', 'ins', 'drawing', 'oMath'):
            with self.subTest(tag=tag):
                source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
                parts = C.package(source)
                root = C.parse(parts['word/document.xml'])
                paragraph = C.indexed(root)['p000001']
                if tag == 'drawing':
                    paragraph.find(C.q('r')).append(E.Element(C.q(tag)))
                elif tag == 'oMath':
                    paragraph.append(E.Element('{' + C.M + '}' + tag))
                else:
                    paragraph.append(E.Element(C.q(tag)))
                parts['word/document.xml'] = C.serialize(root)
                with zipfile.ZipFile(source, 'w') as zipped:
                    for name, data in parts.items():
                        zipped.writestr(name, data)
                self.assert_rejected(source, self.map(source, '张三', ['张三']), 'protected|complex structure|non-text content')

    def test_author_replacement_crossing_semantic_formatting_is_rejected(self):
        runs = [('张三', {'bold': True}), ('、李四、王五', {}), ('（2020）提出结论。', {})]
        body = ''.join(text for text, _ in runs)
        source = self.source(body, ['[1] 张三,李四,王五. 研究成果[J]. 学报,2020.'], runs=runs)
        self.assert_rejected(source, self.map(source, '张三、李四、王五', ['张三', '李四', '王五'], abbreviation='等'),
                             'crosses different source formatting')

    def test_unselected_text_runs_math_parts_and_reference_parts_are_preserved(self):
        def extra(document):
            paragraph = document.add_paragraph('未处理段落 100% 与末尾字符保留。')
            run = paragraph.add_run('语义强调')
            run.bold = True
            paragraph._p.append(E.Element('{' + C.M + '}oMath'))
        runs = [('前文数据99；', {'bold': True}), ('张三、李四、王五', {}), ('（2020）提出结论。', {}), ('后文数据100。', {'italic': True})]
        body = ''.join(text for text, _ in runs)
        source = self.source(body, ['[1] 张三,李四,王五. 研究成果[J]. 学报,2020.'], runs=runs, extra=extra)
        mapping = self.map(source, '张三、李四、王五', ['张三', '李四', '王五'], abbreviation='等', claim_end=body.index('。'))
        before = C.package(source)
        before_paragraph = C.indexed(C.parse(before['word/document.xml']))['p000001']
        output = self.output(source, mapping)
        after = C.package(output)
        for name in before:
            if name not in {'word/document.xml', 'word/styles.xml'}:
                self.assertEqual(after[name], before[name], name)
        before_ps, after_ps = C.indexed(C.parse(before['word/document.xml'])), C.indexed(C.parse(after['word/document.xml']))
        for pid in before_ps:
            if pid != 'p000001':
                self.assertEqual(C.style_signature(before_ps[pid]), C.style_signature(after_ps[pid]), pid)
        after_runs = after_ps['p000001'].findall(C.q('r'))
        self.assertEqual(C.style_signature(before_paragraph.findall(C.q('r'))[0]), C.style_signature(after_runs[0]))
        self.assertEqual(C.style_signature(before_paragraph.findall(C.q('r'))[-1]), C.style_signature(after_runs[-1]))
        self.assertEqual(self.text(output), '前文数据99；张三等（2020）提出结论[1]。后文数据100。')

    def test_verify_detects_modified_body_reference_and_unselected_package_parts(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, '张三', ['张三'])
        output = self.output(source, mapping)
        for pid in ('p000001', 'p000003'):
            def mutate(parts):
                root = C.parse(parts['word/document.xml'])
                C.indexed(root)[pid].find('.//w:t', C.NS).text += '修改'
                parts['word/document.xml'] = C.serialize(root)
            modified = self.mutate(output, mutate)
            self.assertFalse(N.verify(source, modified, mapping)['ok'])
        modified = self.mutate(output, lambda parts: parts.update({'word/_rels/document.xml.rels': b'changed'}))
        self.assertFalse(N.verify(source, modified, mapping)['ok'])

    def test_outputs_preserve_source_and_existing_paths(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        mapping = self.map(source, '张三', ['张三'])
        output = self.output(source, mapping)
        before = output.read_bytes()
        with self.assertRaisesRegex(ValueError, 'new path'):
            N.apply(source, output, mapping)
        with self.assertRaisesRegex(ValueError, 'new path'):
            N.apply(source, source, mapping)
        self.assertEqual(output.read_bytes(), before)

    def test_leading_number_without_reference_heading_or_style_is_not_bibliography_evidence(self):
        source = self.source('张三（2020）提出结论。', ['[1] 张三. 研究成果[J]. 学报,2020.'])
        parts = C.package(source)
        root = C.parse(parts['word/document.xml'])
        paragraph = C.indexed(root)['p000002']
        paragraph.getparent().remove(paragraph)
        parts['word/document.xml'] = C.serialize(root)
        with zipfile.ZipFile(source, 'w') as zipped:
            for name, data in parts.items():
                zipped.writestr(name, data)
        mapping = self.map(source, '张三', ['张三'])
        self.assert_rejected(source, mapping, 'existing bibliography entry')
        self.assertEqual(N.inspect(source)['references'], [])
        self.assertEqual(N.inspect(source)['matches'], [])

    def test_reference_entry_style_can_authorize_bibliography_without_heading(self):
        from docx.enum.style import WD_STYLE_TYPE
        source = self.work / 'styled-references.docx'
        document = Document()
        document.add_paragraph('张三（2020）提出结论。')
        style = document.styles.add_style('Reference entry', WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph('[1] 张三. 研究成果[J]. 学报,2020.', style=style)
        document.save(source)
        output = self.output(source, self.map(source, '张三', ['张三']))
        self.assertEqual(self.text(output), '张三（2020）提出结论[1]。')
        self.assertEqual(len(N.inspect(source)['references']), 1)

    def test_foreign_author_cannot_be_sliced_out_of_an_unknown_name_or_word(self):
        for author in ('BlackSmith', "O'Smith", '5Smith'):
            with self.subTest(author=author):
                source = self.source(author + ' (2020) reports results.', ['[1] Smith J. Research[J]. Journal,2020.'])
                mapping = self.map(source, 'Smith', [('Smith J', 'Smith')], language='en')
                self.assert_rejected(source, mapping, 'word boundary')
                self.assertEqual(N.inspect(source)['matches'], [])


if __name__ == '__main__':
    unittest.main()

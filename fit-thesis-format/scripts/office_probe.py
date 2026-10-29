#!/usr/bin/env python3
"""Portable capability discovery and native Word/WPS smoke tests.

No packages, fonts or applications are installed. Registration is reported
separately from a successful actual export. Rendering never overwrites input.
"""
from __future__ import annotations

import argparse
import copy
import ctypes
from collections import Counter,defaultdict
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile

SCRIPT = Path(__file__).resolve().with_name('office_render.ps1')
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}
LAYOUT_FIELDS = {'TOC','SEQ','REF','PAGEREF','STYLEREF','PAGE','NUMPAGES','SECTIONPAGES'}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _powershell():
    return shutil.which('pwsh') or shutil.which('powershell')


def _owned_receipt_candidates(state,expected_hash):
    if state.get('schema_version')!=10 or state.get('source_sha256')!=expected_hash or state.get('closed'):
        return []
    old=set(state.get('preexisting_pids',[]));exe=str(state.get('engine_executable','')).casefold()
    result=[]
    for item in state.get('owned_processes',[]):
        pid=item.get('id');start=item.get('start_filetime_utc');path=item.get('path','')
        if isinstance(pid,int) and pid>0 and pid not in old and isinstance(start,int) and start>0 and str(path).casefold()==exe and exe:
            result.append(item)
    return result


def cleanup_owned_timeout(state,expected_hash):
    """Terminate only receipt-bound processes whose live identity still matches.

    A missing/unproven process is left alone. No process-name lookup or broad
    taskkill is used; creation time is read from the opened process handle to
    prevent PID reuse between check and termination.
    """
    return _terminate_matching_handles(_owned_receipt_candidates(state,expected_hash))


def _terminate_matching_handles(candidates):
    """Internal primitive; callers must establish ownership before this step."""
    if os.name!='nt':return []
    from ctypes import wintypes
    kernel=ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.OpenProcess.argtypes=[wintypes.DWORD,wintypes.BOOL,wintypes.DWORD];kernel.OpenProcess.restype=wintypes.HANDLE
    kernel.GetProcessTimes.argtypes=[wintypes.HANDLE,*([ctypes.POINTER(wintypes.FILETIME)]*4)]
    kernel.QueryFullProcessImageNameW.argtypes=[wintypes.HANDLE,wintypes.DWORD,wintypes.LPWSTR,ctypes.POINTER(wintypes.DWORD)]
    kernel.TerminateProcess.argtypes=[wintypes.HANDLE,wintypes.UINT]
    kernel.WaitForSingleObject.argtypes=[wintypes.HANDLE,wintypes.DWORD]
    kernel.CloseHandle.argtypes=[wintypes.HANDLE]
    results=[]
    for item in candidates:
        handle=kernel.OpenProcess(0x1000|0x0001|0x100000,False,item['id'])
        if not handle:
            results.append({'pid':item['id'],'terminated':False,'reason':'process already gone or inaccessible'});continue
        try:
            created,ended,user,system=(wintypes.FILETIME() for _ in range(4))
            if not kernel.GetProcessTimes(handle,ctypes.byref(created),ctypes.byref(ended),ctypes.byref(user),ctypes.byref(system)):
                results.append({'pid':item['id'],'terminated':False,'reason':'creation time unavailable'});continue
            actual=(created.dwHighDateTime<<32)|created.dwLowDateTime
            buffer=ctypes.create_unicode_buffer(32768);size=wintypes.DWORD(len(buffer))
            ok=kernel.QueryFullProcessImageNameW(handle,0,buffer,ctypes.byref(size))
            if actual!=item['start_filetime_utc'] or not ok or buffer.value.casefold()!=item['path'].casefold():
                results.append({'pid':item['id'],'terminated':False,'reason':'live process identity differs from receipt'});continue
            success=bool(kernel.TerminateProcess(handle,1))
            if success:kernel.WaitForSingleObject(handle,5000)
            results.append({'pid':item['id'],'terminated':success,'reason':'verified owned timed-out instance'})
        finally:kernel.CloseHandle(handle)
    return results


def _parse_cached_pagerefs(source):
    """Read stored field RESULTS, including nested PAGEREFs inside a TOC."""
    from lxml import etree
    with zipfile.ZipFile(source) as z:
        root = etree.fromstring(z.read('word/document.xml'), etree.XMLParser(resolve_entities=False, no_network=True))
    frames, found = [], []
    for el in root.iter():
        if el.tag == f'{{{W}}}fldChar':
            kind = el.get(f'{{{W}}}fldCharType')
            if kind == 'begin':
                frames.append({'code': '', 'result': '', 'separate': False})
            elif kind == 'separate' and frames:
                frames[-1]['separate'] = True
            elif kind == 'end' and frames:
                frame = frames.pop()
                match = re.match(r'^\s*PAGEREF\s+(?:"([^"]+)"|(\S+))', frame['code'], re.I)
                if match:
                    found.append({'target_bookmark': match.group(1) or match.group(2), 'displayed_page': frame['result'].strip()})
        elif el.tag == f'{{{W}}}instrText' and frames and not frames[-1]['separate']:
            frames[-1]['code'] += el.text or ''
        elif el.tag == f'{{{W}}}t':
            for frame in frames:
                if frame['separate']:
                    frame['result'] += el.text or ''
        elif el.tag == f'{{{W}}}fldSimple':
            code = el.get(f'{{{W}}}instr', '')
            match = re.match(r'^\s*PAGEREF\s+(?:"([^"]+)"|(\S+))', code, re.I)
            if match:
                found.append({'target_bookmark': match.group(1) or match.group(2),
                              'displayed_page': ''.join(el.itertext()).strip()})
    return found


def _field_spans(root):
    """Complex fields with actual start/end nodes; nested fields are retained."""
    stack, found = [], []
    for el in root.iter():
        if el.tag == f'{{{W}}}fldSimple':
            code=el.get(f'{{{W}}}instr','');match=re.match(r'^\s*([A-Za-z]+)',code)
            found.append({'simple':el,'code':code,'kind':match.group(1).upper() if match else 'UNKNOWN',
                          'identity':' '.join(code.split()),'result':''.join(el.xpath('.//w:t/text()',namespaces=NS)),
                          'inside_layout':any(re.match(r'^\s*TOC\b',f['code'],re.I) for f in stack)})
        elif el.tag == f'{{{W}}}fldChar':
            typ = el.get(f'{{{W}}}fldCharType')
            if typ == 'begin':
                inside_layout=any(re.match(r'^\s*TOC\b',f['code'],re.I) for f in stack)
                stack.append({'begin':el,'code':'','result':'','separate':None,'inside_layout':inside_layout})
            elif typ == 'separate' and stack:
                stack[-1]['separate']=el
            elif typ == 'end' and stack:
                frame=stack.pop();frame['end']=el
                match=re.match(r'^\s*([A-Za-z]+)',frame['code'])
                frame['kind']=match.group(1).upper() if match else 'UNKNOWN'
                frame['identity']=' '.join(frame['code'].split())
                found.append(frame)
        elif el.tag==f'{{{W}}}instrText' and stack and stack[-1]['separate'] is None:
            stack[-1]['code']+=el.text or ''
        elif el.tag==f'{{{W}}}t':
            for frame in stack:
                if frame['separate'] is not None:
                    frame['result']+=el.text or ''
    return found


def restore_protected_field_caches(source,saved):
    """Restore non-layout field caches after native auto-update-on-open.

    The native save is a newly created output, never the source. Only simple
    text-result fields with unambiguous matching codes can be restored. Complex
    changed fields fail closed instead of losing embedded objects or links.
    """
    from lxml import etree
    parser=etree.XMLParser(resolve_entities=False,no_network=True)
    with zipfile.ZipFile(source) as z:before={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(saved) as z:after={n:z.read(n) for n in z.namelist()}
    changes=[]
    for name,data in before.items():
        if not re.fullmatch(r'word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml',name):continue
        oldroot=etree.fromstring(data,parser)
        old=[f for f in _field_spans(oldroot) if f['kind'] not in LAYOUT_FIELDS and not f['inside_layout']]
        if not old:continue
        if name not in after:raise RuntimeError('Native save lost a protected field story: '+name)
        newroot=etree.fromstring(after[name],parser)
        new=[f for f in _field_spans(newroot) if f['kind'] not in LAYOUT_FIELDS and not f['inside_layout']]
        # New TOC hyperlinks are layout output. Match original field code and
        # occurrence; every original non-layout field must still be present.
        groups={}
        for f in new:groups.setdefault(f['identity'],[]).append(f)
        changed=False
        for field in old:
            matches=groups.get(field['identity'],[])
            if not matches:raise RuntimeError('Native save lost a protected field: '+field['identity'])
            dest=matches.pop(0)
            if field['result']==dest['result']:continue
            def span(f):
                if 'simple' in f:
                    parent=f['simple'].getparent()
                    if parent.tag!=f'{{{W}}}p':raise RuntimeError('Protected field has non-paragraph container')
                    index=parent.index(f['simple']);return parent,index,index
                parents=[f[x].getparent().getparent() for x in ('begin','separate','end') if f[x] is not None]
                if len(parents)!=3 or any(p.tag!=f'{{{W}}}p' for p in parents) or parents[0]!=parents[1] or parents[1]!=parents[2]:
                    raise RuntimeError('Complex protected field changed; cannot restore safely: '+f['identity'])
                for boundary in ('begin','end'):
                    if any(c.tag not in {f'{{{W}}}rPr',f'{{{W}}}fldChar'} for c in f[boundary].getparent()):
                        raise RuntimeError('Protected field boundary shares non-field text')
                parent=parents[0];return parent,parent.index(f['begin'].getparent()),parent.index(f['end'].getparent())
            srcparent,start,end=span(field);dstparent,a,b=span(dest)
            nodes=list(srcparent)[start:end+1]
            allowed={f'{{{W}}}'+n for n in ('r','rPr','rFonts','sz','szCs','b','bCs','i','iCs','color','lang','noProof','rtl','vertAlign','fldChar','fldSimple','instrText','t','spacing','kern','position','u','highlight','shd','webHidden','vanish')}
            if any(n.tag not in allowed for node in nodes for n in node.iter()):
                raise RuntimeError('Protected field contains complex content; cannot restore safely: '+field['identity'])
            for node in list(dstparent)[a:b+1]:dstparent.remove(node)
            for offset,node in enumerate(nodes):dstparent.insert(a+offset,copy.deepcopy(node))
            changes.append({'part':name,'code':field['identity'],'native_result':dest['result'],'restored_result':field['result']})
            changed=True
        if changed:after[name]=etree.tostring(newroot,encoding='UTF-8',xml_declaration=True,standalone=True)
    if changes:
        temp=Path(saved).with_name(Path(saved).name+'.protected-'+uuid.uuid4().hex+'.tmp')
        with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in after.items():z.writestr(name,data)
        os.replace(temp,saved)
    return changes


def dynamic_display_changes(source,native_results):
    from lxml import etree
    old=defaultdict(list);new=defaultdict(list)
    with zipfile.ZipFile(source) as z:
        for name in z.namelist():
            if not re.fullmatch(r'word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml',name):continue
            root=etree.fromstring(z.read(name),etree.XMLParser(resolve_entities=False,no_network=True))
            for f in _field_spans(root):
                if f['kind'] not in LAYOUT_FIELDS and not f['inside_layout']:
                    old[f['identity']].append(f['result'])
    for field in native_results:
        if field.get('kind') not in LAYOUT_FIELDS:
            new[' '.join(field.get('code','').split())].append(field.get('result','').strip('\r\x07'))
    return [{'code':code,'stored_results':values,'native_display_results':new.get(code,[]),
             'requires_visible_content_review':True}
            for code,values in old.items() if Counter(values)!=Counter(new.get(code,[]))]


def _page_label(number, style):
    if number is None or style is None:
        return None
    if style in (1, 2):
        result = ''
        for value, letters in ((1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')):
            while number >= value:
                result += letters
                number -= value
        return result.lower() if style == 2 else result
    if style == 0:
        return str(number)
    return None


def validate_toc(source, bookmark_pages):
    entries = _parse_cached_pagerefs(source)
    for entry in entries:
        page = bookmark_pages.get(entry['target_bookmark'], {})
        entry.update(page)
        entry['actual_page'] = page.get('actual_page')
        expected = _page_label(page.get('actual_page'), page.get('number_style'))
        entry['expected_page_label'] = expected
        # An unavailable native number style cannot prove the cache is correct.
        entry['passed'] = expected is not None and entry['displayed_page'] == expected
    return entries


def validate_table_spans(source,spans,page_count):
    from lxml import etree
    with zipfile.ZipFile(source) as z:
        root=etree.fromstring(z.read('word/document.xml'),etree.XMLParser(resolve_entities=False,no_network=True))
    tables=root.xpath('./w:body//w:tbl[not(ancestor::w:txbxContent)]',namespaces=NS)
    if not isinstance(spans,list) or len(spans)!=len(tables):
        raise RuntimeError(f'Native table collection count differs from DOCX: {len(spans) if isinstance(spans,list) else "missing"} versus {len(tables)}')
    for index,span in enumerate(spans,1):
        if span.get('table_index')!=index or not isinstance(span.get('start_page'),int) or not isinstance(span.get('end_page'),int) or not 1<=span['start_page']<=span['end_page']<=page_count:
            raise RuntimeError('Missing or invalid native table physical page span')
    return True


def render(engine, source, pdf, saved_docx=None, timeout=180):
    if engine not in ('word', 'wps'):
        raise ValueError('engine must be word or wps')
    source, pdf = Path(source).resolve(), Path(pdf).resolve()
    if source.suffix.lower() != '.docx' or not source.is_file():
        raise ValueError('source must be an existing DOCX')
    if pdf.suffix.lower() != '.pdf' or pdf.exists() or pdf == source:
        raise ValueError('PDF output must be a new .pdf file')
    if saved_docx is not None:
        saved_docx = Path(saved_docx).resolve()
        if saved_docx.suffix.lower() != '.docx' or saved_docx.exists() or saved_docx == source:
            raise ValueError('saved_docx must be a new .docx path')
    native_saved=(saved_docx.with_name(saved_docx.stem+'.native-'+engine+'-'+uuid.uuid4().hex[:8]+'.docx')
                  if saved_docx is not None else None)
    shell = _powershell()
    if os.name != 'nt' or not shell:
        raise RuntimeError('Native Word/WPS render requires Windows PowerShell')
    before = sha(source)
    state_path=pdf.with_name(pdf.stem+'.office-state-'+uuid.uuid4().hex[:8]+'.json')
    command = [shell, '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File', str(SCRIPT),
               '-Engine', engine, '-InputPath', str(source), '-OutputPdf', str(pdf),'-StatePath',str(state_path)]
    if saved_docx is not None:
        command += ['-UpdateFields', '-SavedPath', str(native_saved)]
    try:
        proc = subprocess.run(command, capture_output=True, text=True, encoding='utf-8-sig', timeout=timeout)
    except subprocess.TimeoutExpired as error:
        state={}
        try:state=json.loads(state_path.read_text(encoding='utf-8-sig'))
        except (OSError,ValueError):pass
        cleanup=cleanup_owned_timeout(state,before)
        state['timeout_cleanup']=cleanup
        if state_path.exists():state_path.write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
        raise RuntimeError(f'{engine} timed out after {timeout}s at phase {state.get("phase","unknown")}; '
                           f'owned-instance diagnostics: {state_path}; cleanup={cleanup}. No other office process was terminated.') from error
    if sha(source) != before:
        raise RuntimeError('Native renderer modified its input')
    if proc.returncode:
        raise RuntimeError(f'{engine} native export failed: {(proc.stderr or proc.stdout)[-6000:]}')
    try:
        report = json.loads(proc.stdout)
    except json.JSONDecodeError as error:
        raise RuntimeError(f'{engine} returned invalid JSON: {proc.stdout[:2000]}') from error
    if report.get('engine_key') != engine or report.get('source_sha256_before', '').lower() != before:
        raise RuntimeError('Native renderer provenance mismatch')
    if report.get('source_sha256_after', '').lower() != before or not report.get('source_unchanged'):
        raise RuntimeError('Native renderer lacks matching source preservation evidence')
    if not pdf.is_file() or sha(pdf) != report.get('pdf_sha256'):
        raise RuntimeError('Native PDF missing or hash mismatch')
    import pypdfium2
    document = pypdfium2.PdfDocument(str(pdf))
    try:
        count = len(document)
    finally:
        document.close()
    if count < 1 or count != report.get('pages'):
        raise RuntimeError('Native/PDF actual page counts disagree')
    report['pdf_actual_pages'] = count
    if engine=='wps' and saved_docx is None:
        with pypdfium2.PdfDocument(str(pdf)) as actual:
            pdf_errors=[]
            for index in range(len(actual)):
                page=actual[index];tp=page.get_textpage()
                try:
                    for match in re.finditer(r'Error!|错误[!！]|未找到引用源|未定义书签|Reference source not found|Bookmark not defined',tp.get_text_range(),re.I):
                        pdf_errors.append({'physical_page':index+1,'result':match.group()})
                finally:tp.close();page.close()
        report['field_scan_mode']='actual_pdf_text_and_native_toc_targets'
        report['field_errors']=pdf_errors
    validate_table_spans(source,report.get('table_spans'),count)
    if saved_docx:
        if not native_saved.is_file() or sha(native_saved)!=report.get('saved_docx_sha256'):
            raise RuntimeError('Native saved intermediate missing or hash mismatch')
        from field_cache_merge import merge_field_caches
        report['native_saved_docx']=str(native_saved)
        report['native_saved_docx_sha256']=report['saved_docx_sha256']
        report['cache_merge']=merge_field_caches(source,native_saved,saved_docx)
        report['saved_docx']=str(saved_docx)
        report['saved_docx_sha256']=sha(saved_docx)
        report['pdf_represents_native_intermediate']=True
        report['requires_same_file_dual_rerender']=True
    cached_source = saved_docx or source
    report['toc_entries'] = validate_toc(cached_source, report.get('bookmark_pages', {}))
    report['toc_cache_verified'] = all(entry['passed'] for entry in report['toc_entries'])
    report['toc_cache_entry_count'] = len(report['toc_entries'])
    report['nonlayout_dynamic_display_changes']=dynamic_display_changes(cached_source,report.get('field_results') or [])
    report['visual_review_completed'] = False
    return report


def _installed_fonts():
    if os.name != 'nt':
        return {'checked': False, 'required': ['SimSun', 'Times New Roman', 'SimHei'], 'missing': None}
    import winreg
    values = []
    for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            with winreg.OpenKey(hive, r'SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts') as key:
                for i in range(winreg.QueryInfoKey(key)[1]):
                    values.append(winreg.EnumValue(key, i)[0])
        except OSError:
            pass
    aliases = {'SimSun': ('simsun','宋体'), 'Times New Roman': ('times new roman',), 'SimHei': ('simhei','黑体')}
    found = {name: any(any(alias in value.lower() for alias in patterns) for value in values) for name,patterns in aliases.items()}
    return {'checked': True, 'required': list(aliases), 'available': found, 'missing': [name for name,value in found.items() if not value]}


def _registered(progid):
    if os.name != 'nt':
        return False
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid + r'\CLSID') as key:
            return bool(winreg.QueryValueEx(key, '')[0])
    except OSError:
        return False


def make_smoke_fixture(path):
    """A synthetic, non-user document for testing actual installed capability."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.enum.style import WD_STYLE_TYPE
    d = Document()
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    # A formatter prepares TOC styles before refresh; the capability fixture
    # mirrors that contract rather than importing arbitrary native styles.
    for style_name in ('TOC 1','TOC 2'):
        if style_name not in d.styles:
            style=d.styles.add_style(style_name,WD_STYLE_TYPE.PARAGRAPH)
            style.base_style=normal
    if 'Hyperlink' not in d.styles:
        d.styles.add_style('Hyperlink',WD_STYLE_TYPE.CHARACTER)
    d.settings.odd_and_even_pages_header_footer = True
    s=d.sections[0]
    s.header.paragraphs[0].text='格式能力测试 · 奇数页'
    s.even_page_header.paragraphs[0].text='福州理工学院本科生毕业论文（设计）'
    def field(p, code, cache='1'):
        for typ in ('begin','code','separate','text','end'):
            r=OxmlElement('w:r')
            if typ in ('begin','separate','end'):
                el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),typ)
            elif typ=='code':
                el=OxmlElement('w:instrText');el.set(qn('xml:space'),'preserve');el.text=' '+code+' '
            else:
                el=OxmlElement('w:t');el.text=cache
            r.append(el);p._p.append(r)
    for footer in (s.footer,s.even_page_footer):
        field(footer.paragraphs[0],'PAGE')
    d.add_heading('目录',0)
    field(d.add_paragraph(),'TOC \\o "1-2" \\h \\z', '目录待更新')
    d.add_page_break()
    p=d.add_heading('一、中文字体与数学对象',1)
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'10');start.set(qn('w:name'),'FITSmokeHeading');p._p.insert(1,start)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'10');p._p.append(end)
    d.add_paragraph('宋体中文正文与 Times New Roman Western text，DML-DID，12345。')
    black = d.add_paragraph().add_run('黑体标题检查')
    black.font.name = 'SimHei'
    black._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimHei')
    p=d.add_paragraph()
    math=OxmlElement('m:oMath');fraction=OxmlElement('m:f')
    for part,text in [('num','x'),('den','y')]:
        el=OxmlElement('m:'+part);r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);el.append(r);fraction.append(el)
    math.append(fraction);p._p.append(math)
    p.add_run('\t公式（1-1）')
    d.add_paragraph('表1-1 双引擎测试')
    table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='变量';table.cell(0,1).text='数值';table.cell(1,0).text='DID';table.cell(1,1).text='0.004046***'
    d.add_paragraph('受保护日期域不得被强制更新：')
    field(d.add_paragraph(),'DATE \\@ "yyyy-MM-dd"','2000-01-01')
    d.save(path)


def validate_smoke_fonts(glyphs):
    """Require the actual marker glyphs, rather than any registered font."""
    joined = ''.join(char for char, font in glyphs)
    requirements = {'宋体中文正文': ('simsun',), '黑体标题检查': ('simhei',), 'DML-DID': ('timesnewroman',)}
    checks = {}
    for marker, families in requirements.items():
        hits = []
        start = joined.find(marker)
        while start >= 0:
            used = [font for char, font in glyphs[start:start+len(marker)] if not char.isspace()]
            valid = bool(used) and all(any(family in re.sub(r'[^a-z]', '', font.lower()) for family in families) for font in used)
            hits.append({'actual_fonts': sorted(set(used)), 'passed': valid})
            start = joined.find(marker, start+1)
        checks[marker] = {'occurrences': hits, 'passed': bool(hits) and all(h['passed'] for h in hits)}
    return {'markers':checks, 'passed':all(c['passed'] for c in checks.values())}

def inspect_smoke_pdf(pdf):
    """Read actual exported glyph/text and font names, not OS registration."""
    import pypdfium2
    raw=pypdfium2.raw
    doc=pypdfium2.PdfDocument(str(pdf));texts=[];fonts=set();glyphs=[]
    try:
        for page in doc:
            text=page.get_textpage()
            try:
                texts.append(text.get_text_range())
                for index in range(min(text.count_chars(),1500)):
                    font_name=''
                    flags=ctypes.c_int()
                    size=raw.FPDFText_GetFontInfo(text.raw,index,None,0,ctypes.byref(flags))
                    if size:
                        buf=ctypes.create_string_buffer(size)
                        raw.FPDFText_GetFontInfo(text.raw,index,buf,size,ctypes.byref(flags))
                        font_name=buf.value.decode('utf-8',errors='replace')
                        fonts.add(font_name)
                    glyphs.append((text.get_text_range(index,1),font_name))
            finally:text.close();page.close()
    finally:doc.close()
    joined=''.join(texts)
    expected=['中文字体与数学对象','福州理工学院本科生毕业论文','DML-DID','0.004046']
    present={v:v in joined for v in expected}
    font_checks=validate_smoke_fonts(glyphs)
    return {'text_checks':present,'sampled_actual_font_names':sorted(fonts),'actual_marker_font_checks':font_checks,
            'passed':all(present.values()) and font_checks['passed']}


def doctor(work_dir=None, smoke=False):
    deps={name: importlib.util.find_spec(name) is not None for name in ('lxml','docx','pypdfium2','PIL')}
    result={'schema_version':10,'platform':sys.platform,'powershell':_powershell(), 'fonts':_installed_fonts(),
            'python_dependencies':deps,'engines':{},'smoke_executed':False,'dual_render_capable':False,
            'visual_capability':'requires_host_image_viewer_confirmation'}
    for engine,prog in [('word','Word.Application'),('wps','KWPS.Application')]:
        result['engines'][engine]={'registered':_registered(prog),'export_tested':False,'available':False}
    if not smoke:
        return result
    if os.name != 'nt' or not result['powershell'] or not all(deps.values()):
        result['reason']='Native smoke test needs Windows, PowerShell and listed Python dependencies.'
        return result
    base=Path(work_dir) if work_dir else Path(tempfile.gettempdir())
    qa=base/('office-smoke-'+uuid.uuid4().hex[:12]);qa.mkdir(parents=True,exist_ok=False)
    fixture=qa/'fixture.docx';make_smoke_fixture(fixture)
    result.update(smoke_executed=True,smoke_directory=str(qa),fixture_sha256=sha(fixture))
    for engine in ('word','wps'):
        data=result['engines'][engine]
        try:
            refresh=render(engine,fixture,qa/f'{engine}-updated.pdf',qa/f'{engine}-updated.docx')
            reread=render(engine,qa/f'{engine}-updated.docx',qa/f'{engine}-readonly.pdf')
            data.update(available=True,export_tested=True,report=reread,refresh_report=refresh)
            data['actual_pdf_content']=inspect_smoke_pdf(qa/f'{engine}-readonly.pdf')
            if reread.get('field_errors') or not reread.get('toc_cache_verified') or not reread.get('native_math_count') or not data['actual_pdf_content']['passed']:
                data.update(available=False,error='Smoke output has field/cache/equation capability failures.')
        except Exception as error:
            data.update(export_tested=True,available=False,error=str(error))
    result['dual_render_capable']=all(result['engines'][k]['available'] for k in ('word','wps')) and not result['fonts'].get('missing')
    (qa/'doctor.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--work-dir');p.add_argument('--smoke',action='store_true')
    args=p.parse_args();print(json.dumps(doctor(args.work_dir,args.smoke),ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()

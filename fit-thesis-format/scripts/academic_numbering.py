#!/usr/bin/env python3
"""Native Word controls for explicitly authorized *new* FIT thesis examples.

This module is intentionally separate from the content-preserving converters.
It creates live heading lists, fields, bookmarks and native OMML.  Never call it
on an existing user's manuscript under the ordinary "format only" contract.
"""
from __future__ import annotations

import copy
import json
import re
from pathlib import Path

from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

SCHOOL = "福州理工学院本科生毕业论文（设计）"
STYLE_IDS = ("FIT_chapter", "FIT_section", "FIT_subsection", "FIT_subsubsection")
SCHOOL_LEVELS = (
    ("chineseCounting", "%1、"),
    ("chineseCounting", "（%2）"),
    ("decimal", "%3."),
    ("decimal", "（%4）"),
)
DECIMAL_LEVELS = tuple(("decimal", ".".join(f"%{i}" for i in range(1, n + 2))) for n in range(4))


def element(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn(key if ":" in key else "w:" + key), str(value))
    return node


def _replace(parent, name, node):
    for old in list(parent):
        if old.tag == qn(name):
            parent.remove(old)
    parent.append(node)
    return node


def _style(document, sid):
    for style in document.styles:
        if style.style_id == sid:
            return style
    raise ValueError(f"Install the named FIT paragraph styles first: missing {sid}")


def install_styles(document):
    """Install the current style-spec into a new document, preserving style IDs.

    Existing FIT styles are left alone so callers can deliberately customize
    them. A name collision with another ID is an error rather than a overwrite.
    """
    spec = json.loads((Path(__file__).resolve().parent.parent / "references" / "style-spec.json").read_text(encoding="utf-8"))
    existing_ids = {s.style_id for s in document.styles}
    for role, over in spec["roles"].items():
        sid = "FIT_" + role
        if sid in existing_ids:
            continue
        name = over["name"]
        if name in document.styles:
            raise ValueError(f"Style name already exists with another ID: {name}")
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style._element.set(qn("w:styleId"), sid)
        style.base_style = document.styles["Normal"]
        style.quick_style = True
        style.hidden = False
        style.locked = False
        _replace(style._element, "w:autoRedefine", element("w:autoRedefine", val=0))
        values = spec["defaults"] | over
        pp = style._element.get_or_add_pPr()
        for tag, attrs in (
            ("w:keepNext", {"val": int(bool(values.get("keep_next")))}),
            ("w:keepLines", {"val": int(bool(values.get("keep_lines")))}),
            ("w:pageBreakBefore", {"val": int(bool(values.get("page_break")))}),
            ("w:widowControl", {"val": 1}),
            ("w:snapToGrid", {"val": 0}),
            ("w:spacing", {"before": round(values["before_pt"] * 20), "after": round(values["after_pt"] * 20), "line": round(values["line_pt"] * 20), "lineRule": values["line_rule"]}),
            ("w:jc", {"val": values["align"]}),
            ("w:outlineLvl", {"val": values["outline"]}),
        ):
            _replace(pp, tag, element(tag, **attrs))
        indent = {"left": 0, "right": 0, "firstLine": 0, "firstLineChars": round(values["indent_chars"] * 100)}
        if values.get("left_chars"):
            indent["leftChars"] = round(values["left_chars"] * 100)
        if values.get("hanging_chars"):
            indent.pop("firstLine", None)
            indent.pop("firstLineChars", None)
            indent["hangingChars"] = round(values["hanging_chars"] * 100)
        _replace(pp, "w:ind", element("w:ind", **indent))
        if values.get("bottom_border_pt"):
            border = element("w:pBdr")
            border.append(element("w:bottom", val="single", sz=round(values["bottom_border_pt"] * 8), space=1, color="000000"))
            _replace(pp, "w:pBdr", border)
        if role != "equation":
            style.font.name = values["western"]
            style.font.size = Pt(values["size_pt"])
            style.font.bold = bool(values.get("bold"))
            fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
            for channel in ("ascii", "hAnsi", "cs"):
                fonts.set(qn("w:" + channel), values["western"])
            fonts.set(qn("w:eastAsia"), values["east_asia"])
    # Mathematics in a newly authored example uses a visible character style.
    # Its 12 pt / Cambria Math are an authoring choice, not a school mandate.
    # No per-run default font/size is written: editing this style changes all
    # formulas and inline variables produced by the constructors below.
    if "FIT_math" not in {s.style_id for s in document.styles}:
        math_style = document.styles.add_style("论文-数学变量", WD_STYLE_TYPE.CHARACTER)
        math_style._element.set(qn("w:styleId"), "FIT_math")
        math_style.base_style = document.styles["Default Paragraph Font"]
        math_style.quick_style = True
        math_style.hidden = False
        math_style.locked = False
        math_style.font.name = "Cambria Math"
        math_style.font.size = Pt(12)
        fonts = math_style._element.get_or_add_rPr().get_or_add_rFonts()
        for channel in ("ascii", "hAnsi", "cs"):
            fonts.set(qn("w:" + channel), "Cambria Math")
    if "FIT_equation_number" not in {s.style_id for s in document.styles}:
        label_style = document.styles.add_style("论文-公式编号", WD_STYLE_TYPE.PARAGRAPH)
        label_style._element.set(qn("w:styleId"), "FIT_equation_number")
        label_style.base_style = _style(document, "FIT_body")
        label_style.quick_style = True
        label_style.hidden = False
        label_style.locked = False
        label_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        label_style.paragraph_format.line_spacing = Pt(20)
        _replace(label_style._element.get_or_add_pPr(), "w:ind", element("w:ind", left=0, right=0, firstLine=0, firstLineChars=0))
    if "FIT_equation_label" not in {s.style_id for s in document.styles}:
        # A formula and its label now share a single ordinary paragraph. The
        # label therefore needs a character style rather than a second
        # paragraph style. This produces no table grid, frame or text box.
        label_values = spec["defaults"] | spec["roles"].get("equation_number", {})
        label_chars = document.styles.add_style("论文-公式编号文字", WD_STYLE_TYPE.CHARACTER)
        label_chars._element.set(qn("w:styleId"), "FIT_equation_label")
        label_chars.base_style = document.styles["Default Paragraph Font"]
        label_chars.quick_style = True
        label_chars.hidden = False
        label_chars.locked = False
        label_chars.font.name = label_values["western"]
        label_chars.font.size = Pt(label_values["size_pt"])
        fonts = label_chars._element.get_or_add_rPr().get_or_add_rFonts()
        for channel in ("ascii", "hAnsi", "cs"):
            fonts.set(qn("w:" + channel), label_values["western"])
        fonts.set(qn("w:eastAsia"), label_values["east_asia"])
    return document


def add_field(paragraph, instruction, result="", *, hidden=False, dirty=True):
    """Append a real complex Word field; result is only the display cache."""
    nodes = []
    for tag, text, attrs in (
        ("w:fldChar", None, {"fldCharType": "begin", "dirty": "true" if dirty else "false"}),
        ("w:instrText", " " + instruction.strip() + " ", {}),
        ("w:fldChar", None, {"fldCharType": "separate"}),
        ("w:t", str(result), {}),
        ("w:fldChar", None, {"fldCharType": "end"}),
    ):
        run = element("w:r")
        if hidden:
            properties = element("w:rPr")
            properties.append(element("w:vanish"))
            run.append(properties)
        child = element(tag, **attrs)
        if text is not None:
            child.text = text
            child.set(qn("xml:space"), "preserve")
        run.append(child)
        paragraph._p.append(run)
        nodes.append(run)
    return nodes


def add_bookmark(paragraph, name, nodes):
    """Bookmark the complete selected field/text span for a full-label REF."""
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", name):
        raise ValueError("Word bookmark must start with an ASCII letter and be at most 40 characters")
    root = paragraph._p.getroottree().getroot()
    if root.xpath(f'.//w:bookmarkStart[@w:name="{name}"]'):
        raise ValueError("Duplicate bookmark: " + name)
    ids = [int(n.get(qn("w:id"))) for n in root.xpath(".//w:bookmarkStart")]
    identifier = max(ids, default=-1) + 1
    start = element("w:bookmarkStart", id=identifier, name=name)
    end = element("w:bookmarkEnd", id=identifier)
    if not nodes or any(n.getparent() is not paragraph._p for n in nodes):
        raise ValueError("Bookmark span must contain top-level paragraph children")
    paragraph._p.insert(paragraph._p.index(nodes[0]), start)
    paragraph._p.insert(paragraph._p.index(nodes[-1]) + 1, end)
    return name


def add_equation_reference(paragraph, bookmark, cached_label):
    """REF copies the whole '公式（x-x）' bookmark, including its word/punctuation."""
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", bookmark):
        raise ValueError("Invalid bookmark")
    return add_field(paragraph, f"REF {bookmark} \\h \\* CHARFORMAT", cached_label)


class ThesisNumbering:
    """One shared multilevel list and chapter-aware equation sequences.

    Caller must explicitly authorize creating a new document. Existing-source
    conversions require a separately reviewed source-hash-bound mapping and
    visible-number comparison, and are not performed by this constructor.
    """
    def __init__(self, document, *, authorized_new_document=False, schema="school"):
        if not authorized_new_document:
            raise ValueError("This helper only creates an explicitly authorized new thesis/example")
        if schema not in {"school", "decimal"}:
            raise ValueError("schema must be school or decimal")
        for sid in STYLE_IDS:
            _style(document, sid)
        self.document = document
        self.chapter = 0
        self.equation_in_chapter = 0
        numbering = document.part.numbering_part.element
        abstracts = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
        nums = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
        aid = max(abstracts, default=-1) + 1
        self.num_id = max(nums, default=0) + 1
        abstract = element("w:abstractNum", abstractNumId=aid)
        abstract.append(element("w:multiLevelType", val="multilevel"))
        levels = SCHOOL_LEVELS if schema == "school" else DECIMAL_LEVELS
        for index, ((fmt, pattern), sid) in enumerate(zip(levels, STYLE_IDS)):
            level = element("w:lvl", ilvl=index)
            level.append(element("w:start", val=1))
            level.append(element("w:numFmt", val=fmt))
            # Omit lvlRestart: the OOXML default restarts a lower level when
            # the immediately preceding higher level occurs. This matches
            # the school hierarchy. Native Word validation is essential;
            # explicit lvlRestart values here caused its deeper levels to be
            # read with empty formats on the tested Windows Word build.
            level.append(element("w:pStyle", val=sid))
            level.append(element("w:suff", val="nothing"))
            level.append(element("w:lvlText", val=pattern))
            level.append(element("w:lvlJc", val="left"))
            abstract.append(level)
        # OOXML requires abstractNum definitions before concrete num entries.
        # Appending an abstract after existing nums can make Word silently
        # repair the list and discard its deeper levels.
        first_num = next((i for i, node in enumerate(numbering) if node.tag == qn("w:num")), len(numbering))
        numbering.insert(first_num, abstract)
        num = element("w:num", numId=self.num_id)
        num.append(element("w:abstractNumId", val=aid))
        numbering.append(num)
        for index, sid in enumerate(STYLE_IDS):
            pp = _style(document, sid)._element.get_or_add_pPr()
            properties = element("w:numPr")
            properties.append(element("w:ilvl", val=index))
            properties.append(element("w:numId", val=self.num_id))
            _replace(pp, "w:numPr", properties)

    def heading(self, title, level=0):
        if not 0 <= level <= 3:
            raise ValueError("School heading levels are 0 through 3")
        if re.match(r"^(?:[一二三四五六七八九十百]+、|（[一二三四五六七八九十百0-9]+）|\d+[.、])", title):
            raise ValueError("For a new example supply title words without manually typed numbering")
        paragraph = self.document.add_paragraph(title, style=_style(self.document, STYLE_IDS[level]))
        # Word does not consistently resolve ilvl from custom paragraph styles
        # even when numbering levels carry pStyle. Store the shared list and
        # exact level on each authored heading as well as on its named style.
        num_properties = element("w:numPr")
        num_properties.append(element("w:ilvl", val=level))
        num_properties.append(element("w:numId", val=self.num_id))
        _replace(paragraph._p.get_or_add_pPr(), "w:numPr", num_properties)
        if level == 0:
            self.chapter += 1
            self.equation_in_chapter = 0
            # Native heading list remains Chinese. Hidden chapter SEQ is an
            # independent Arabic counter for equations/figures/tables.
            add_field(paragraph, "SEQ FITChapter \\h", "", hidden=True)
            add_field(paragraph, "SEQ FITEquation \\r 0 \\h", "", hidden=True)
        return paragraph

    def equation(self, expression, *, bookmark=None, width_twips=9355, label_width_twips=1800):
        """Return (paragraph, label) without tables, borders, frames or boxes.

        The first tab centers native inline OMML at the page text midpoint;
        the second tab right-aligns the complete dynamic label at its edge.
        label_width_twips remains accepted for caller compatibility only.
        Wide expressions require rendered-page QA rather than being shrunk.
        """
        if self.chapter < 1:
            raise ValueError("Create a chapter heading before a chapter-numbered equation")
        if width_twips <= 0:
            raise ValueError("Equation text width must be positive")
        self.equation_in_chapter += 1
        label = f"公式（{self.chapter}-{self.equation_in_chapter}）"
        paragraph = self.document.add_paragraph(style=_style(self.document, "FIT_equation"))
        properties = paragraph._p.get_or_add_pPr()
        _replace(properties, "w:jc", element("w:jc", val="left"))
        _replace(properties, "w:ind", element("w:ind", left=0, right=0, firstLine=0, firstLineChars=0, leftChars=0, rightChars=0))
        _replace(properties, "w:spacing", element("w:spacing", before=0, after=0, line=400, lineRule="atLeast"))
        _replace(properties, "w:keepLines", element("w:keepLines", val=1))
        tabs = element("w:tabs")
        tabs.append(element("w:tab", val="center", pos=round(width_twips / 2)))
        tabs.append(element("w:tab", val="right", pos=width_twips))
        _replace(properties, "w:tabs", tabs)
        # No page, paragraph or frame border is introduced here.
        paragraph.add_run().add_tab()
        append_math(paragraph, expression)
        paragraph.add_run().add_tab()
        nodes = [paragraph.add_run("公式（")._r]
        nodes += add_field(paragraph, "SEQ FITChapter \\c", str(self.chapter))
        nodes.append(paragraph.add_run("-")._r)
        nodes += add_field(paragraph, "SEQ FITEquation", str(self.equation_in_chapter))
        nodes.append(paragraph.add_run("）")._r)
        for run in nodes:
            properties = run.find(qn("w:rPr"))
            if properties is None:
                properties = element("w:rPr")
                run.insert(0, properties)
            _replace(properties, "w:rStyle", element("w:rStyle", val="FIT_equation_label"))
        if bookmark:
            add_bookmark(paragraph, bookmark, nodes)
        return paragraph, label


# OMML constructors clone input nodes so the same symbol/expression can be
# reused safely in several formulas without being detached from the old one.
def _append_expression(parent, expression):
    if isinstance(expression, str):
        parent.append(math_text(expression))
    elif isinstance(expression, (list, tuple)):
        for item in expression:
            _append_expression(parent, item)
    else:
        parent.append(copy.deepcopy(expression))


def math_text(text, *, italic=True, size_pt=None):
    run = element("m:r")
    math_properties = element("m:rPr")
    math_properties.append(element("m:sty", **{"m:val": "i" if italic else "p"}))
    run.append(math_properties)
    word_properties = element("w:rPr")
    word_properties.append(element("w:rStyle", val="FIT_math"))
    if size_pt is not None:
        word_properties.append(element("w:sz", val=round(size_pt * 2)))
    run.append(word_properties)
    content = element("m:t")
    content.text = str(text)
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    return run


def math_sequence(*expressions):
    """Return a reusable list; wrap only once with append_math/equation."""
    return list(expressions)


def fraction(numerator, denominator):
    node = element("m:f")
    properties = element("m:fPr")
    properties.append(element("m:type", **{"m:val": "bar"}))
    node.append(properties)
    for name, value in (("m:num", numerator), ("m:den", denominator)):
        child = element(name)
        _append_expression(child, value)
        node.append(child)
    return node


def subscript(base, index):
    node = element("m:sSub")
    for name, value in (("m:e", base), ("m:sub", index)):
        child = element(name)
        _append_expression(child, value)
        node.append(child)
    return node


def superscript(base, exponent):
    node = element("m:sSup")
    for name, value in (("m:e", base), ("m:sup", exponent)):
        child = element(name)
        _append_expression(child, value)
        node.append(child)
    return node


def sub_superscript(base, index, exponent):
    node = element("m:sSubSup")
    for name, value in (("m:e", base), ("m:sub", index), ("m:sup", exponent)):
        child = element(name)
        _append_expression(child, value)
        node.append(child)
    return node


def summation(expression, lower="i=1", upper="n"):
    node = element("m:nary")
    properties = element("m:naryPr")
    properties.append(element("m:chr", **{"m:val": "∑"}))
    properties.append(element("m:limLoc", **{"m:val": "undOvr"}))
    properties.append(element("m:subHide", **{"m:val": 0}))
    properties.append(element("m:supHide", **{"m:val": 0}))
    node.append(properties)
    for name, value in (("m:sub", lower), ("m:sup", upper), ("m:e", expression)):
        child = element(name)
        _append_expression(child, value)
        node.append(child)
    return node


def brackets(expression, left="(", right=")"):
    node = element("m:d")
    properties = element("m:dPr")
    properties.append(element("m:begChr", **{"m:val": left}))
    properties.append(element("m:endChr", **{"m:val": right}))
    node.append(properties)
    child = element("m:e")
    _append_expression(child, expression)
    node.append(child)
    return node


def square_root(expression):
    node = element("m:rad")
    properties = element("m:radPr")
    properties.append(element("m:degHide", **{"m:val": 1}))
    node.append(properties)
    node.append(element("m:deg"))
    child = element("m:e")
    _append_expression(child, expression)
    node.append(child)
    return node


def math_function(name, argument):
    node = element("m:func")
    fn = element("m:fName")
    fn.append(math_text(name, italic=False))
    node.append(fn)
    child = element("m:e")
    _append_expression(child, argument)
    node.append(child)
    return node


def append_math(paragraph, expression):
    """Append a native editable equation, also for inline prose variables."""
    node = element("m:oMath")
    _append_expression(node, expression)
    paragraph._p.append(node)
    return node


def configure_running_furniture(document, plans, title, *, continuous_with_page_break=False):
    """Configure all final sections using real Word odd/even story parts.

    plans = [{'kind':'unpaged'}, {'kind':'front','start':1},
             {'kind':'body','start':1}, {'kind':'body','start':None}]
    The list must cover every section and be called after building sections.
    No conditional IF substitutes are used. All sections have distinct default,
    even and first headers/footers; first-page-different is disabled.

    Optional continuous mode is for audited documents where Word inserts a
    parity blank at a number-format restart. Caller must put pageBreakBefore on
    the first content paragraph of each subsequent section; render with Word
    to confirm the actual transition and logical page parity before delivery.
    """
    if not title.strip():
        raise ValueError("Provide the exact new document title")
    if len(plans) != len(document.sections):
        raise ValueError("Running-furniture plans must cover every final section")
    document.settings.odd_and_even_pages_header_footer = True
    for index, (section, plan) in enumerate(zip(document.sections, plans)):
        kind = plan["kind"]
        if kind not in {"unpaged", "front", "body"}:
            raise ValueError("Unknown region: " + kind)
        section.different_first_page_header_footer = False
        if index:
            section.start_type = WD_SECTION_START.CONTINUOUS if continuous_with_page_break else WD_SECTION_START.NEW_PAGE
        number = element("w:pgNumType", fmt="upperRoman" if kind == "front" else "decimal")
        if plan.get("start") is not None:
            number.set(qn("w:start"), str(plan["start"]))
        _replace(section._sectPr, "w:pgNumType", number)
        for furniture, text in (
            (section.header, title if kind == "body" else ""),
            (section.even_page_header, SCHOOL if kind == "body" else ""),
            (section.first_page_header, ""),
        ):
            furniture.is_linked_to_previous = False
            for child in list(furniture._element):
                furniture._element.remove(child)
            paragraph = furniture.add_paragraph()
            if text:
                paragraph.style = _style(document, "FIT_header")
                paragraph.add_run(text)
            else:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(1)
        for furniture in (section.footer, section.even_page_footer, section.first_page_footer):
            furniture.is_linked_to_previous = False
            for child in list(furniture._element):
                furniture._element.remove(child)
            paragraph = furniture.add_paragraph()
            if kind != "unpaged":
                paragraph.style = _style(document, "FIT_page_number")
                add_field(paragraph, "PAGE", "I" if kind == "front" else "1")
    return document


def add_region(document, *, continuous_with_page_break=False):
    """Start a region and return a paragraph for its first content.

    Use the same mode in configure_running_furniture. Building all regions
    before configuring avoids python-docx's reused final sectPr references.
    """
    document.add_section(WD_SECTION_START.CONTINUOUS if continuous_with_page_break else WD_SECTION_START.NEW_PAGE)
    paragraph = document.add_paragraph()
    if continuous_with_page_break:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


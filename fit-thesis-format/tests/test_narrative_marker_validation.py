import sys,tempfile,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import narrative_citations as N

class NarrativeMarkerValidationTests(unittest.TestCase):
    def inspect(self, marker):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';d=Document()
            d.add_paragraph('张三（2020）和李四（2021）讨论了共同问题'+marker+'。')
            d.add_heading('参考文献',1)
            d.add_paragraph('[1] 张三. 研究问题[J]. 期刊,2020.')
            d.add_paragraph('[2] 李四. 研究问题[J]. 期刊,2021.')
            d.save(p);return N.inspect(p)
    def test_shared_valid_marker_keeps_two_author_year_matches(self):
        result=self.inspect('[1,2]')
        self.assertEqual(len(result['matches']),2)
        self.assertFalse(any('numeric_marker_needs_review' in x['issues'] for x in result['matches']))
        self.assertFalse(any('missing_numeric_marker' in x['issues'] for x in result['matches']))
    def test_missing_reference_and_malformed_marker_are_reviewed(self):
        for marker in ('[1,9]','[ 1 ]','[1 2]','［1］','[1][2]'):
            with self.subTest(marker=marker):
                self.assertTrue(all('numeric_marker_needs_review' in x['issues'] for x in self.inspect(marker)['matches']))

if __name__=='__main__':unittest.main()

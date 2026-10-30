import sys,tempfile,unittest,zipfile
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F
from format_audit import Audit


class FrontSectionContextTests(unittest.TestCase):
    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.base=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()

    def sample(self):
        d=Document();d.add_paragraph('封面信息');d.add_section(WD_SECTION_START.NEW_PAGE)
        d.add_paragraph('真实论文题目');d.add_paragraph('经管类研究论文　　含模拟计算')
        d.add_paragraph('摘要');d.add_paragraph('中文摘要正文');d.add_paragraph('关键词：研究，实证')
        d.add_paragraph('Abstract');d.add_paragraph('English abstract');d.add_paragraph('Keywords: research, evidence')
        d.add_paragraph('目录');d.add_section(WD_SECTION_START.NEW_PAGE)
        d.add_paragraph('一、绪论');d.add_paragraph('实际正文');d.add_section(WD_SECTION_START.NEW_PAGE)
        d.add_paragraph('表1-1 只有表题和数据的后续节');d.add_table(rows=1,cols=1).cell(0,0).text='数据'
        for i,section in enumerate(d.sections):
            if i: F.child(section._sectPr,'pgNumType',fmt='upperRoman' if i==1 else 'decimal',**({'start':'1'} if i in (1,2) else {}))
        path=self.base/'source.docx';d.save(path)
        role_by_text={'封面信息':'cover_info','真实论文题目':'cn_title','经管类研究论文　　含模拟计算':'body',
            '摘要':'cn_abstract_heading','中文摘要正文':'cn_abstract','关键词：研究，实证':'cn_keywords',
            'Abstract':'en_abstract_heading','English abstract':'en_abstract','Keywords: research, evidence':'en_keywords',
            '目录':'toc_heading','一、绪论':'chapter','实际正文':'body','表1-1 只有表题和数据的后续节':'table_caption','数据':'table_text'}
        ps=F.indexed(F.parse(F.package(path)['word/document.xml']))[0]
        mapping={'source_sha256':F.sha(path.read_bytes()),'paragraphs':{pid:role_by_text[F.ptext(p)] for pid,p in ps.items() if F.ptext(p) in role_by_text}}
        return path,mapping

    def test_body_disclosure_on_abstract_page_does_not_reclassify_roman_section(self):
        path,mapping=self.sample();a=Audit(path,mapping);a.infer_roles();a.check_sections()
        self.assertEqual([s['kind'] for s in a.sections],['unpaged','front','body','body'])
        wrong=[i for i in a.issues if i['category'] in {'pagination.number_format','pagination.unwanted_restart','pagination.first_page_start'}]
        self.assertEqual(wrong,[])
        self.assertFalse(any(i['category']=='header.wrong_effective_text' and i['location']['section_id']=='s002' for i in a.issues))

    def test_real_later_body_restart_is_still_reported(self):
        path,mapping=self.sample();parts=F.package(path);root=F.parse(parts['word/document.xml'])
        list(root.iter(F.q('sectPr')))[-1].find(F.q('pgNumType')).set(F.q('start'),'1');parts['word/document.xml']=F.xml(root)
        with zipfile.ZipFile(path,'w') as z:
            for n,v in parts.items():z.writestr(n,v)
        mapping['source_sha256']=F.sha(path.read_bytes());a=Audit(path,mapping);a.infer_roles();a.check_sections()
        bad=[i for i in a.issues if i['category']=='pagination.unwanted_restart']
        self.assertEqual(len(bad),1);self.assertEqual(bad[0]['location']['section_id'],'s004')

    def test_wrong_front_number_format_is_not_hidden_by_context_fix(self):
        path,mapping=self.sample();parts=F.package(path);root=F.parse(parts['word/document.xml'])
        list(root.iter(F.q('sectPr')))[1].find(F.q('pgNumType')).set(F.q('fmt'),'decimal');parts['word/document.xml']=F.xml(root)
        with zipfile.ZipFile(path,'w') as z:
            for n,v in parts.items():z.writestr(n,v)
        mapping['source_sha256']=F.sha(path.read_bytes());a=Audit(path,mapping);a.infer_roles();a.check_sections()
        bad=[i for i in a.issues if i['category']=='pagination.number_format']
        self.assertEqual(len(bad),1);self.assertEqual(bad[0]['location']['section_role'],'front')

    def test_true_title_style_is_resolved_by_name_and_disclosure_is_preserved(self):
        d=Document();d.add_paragraph('数字化转型对制造业企业投资效率的影响','Title')
        d.add_paragraph('经管类研究论文　　含模拟计算','Subtitle');d.add_paragraph('摘要');d.add_paragraph('摘要正文')
        path=self.base/'title.docx';d.save(path);parts=F.package(path);root=F.parse(parts['word/document.xml']);styles=F.parse(parts['word/styles.xml'])
        st=styles.find('w:style[@w:styleId="Title"]',F.NS);st.set(F.q('styleId'),'aa')
        root.find('.//w:pPr/w:pStyle',F.NS).set(F.q('val'),'aa');parts['word/document.xml']=F.xml(root);parts['word/styles.xml']=F.xml(styles)
        with zipfile.ZipFile(path,'w') as z:
            for n,v in parts.items():z.writestr(n,v)
        before=path.read_bytes();a=Audit(path);a.infer_roles()
        self.assertEqual(a.roles['p000001'][0],'cn_title')
        self.assertNotEqual(a.roles.get('p000002',(None,))[0],'cn_title')
        self.assertEqual(a.title,'数字化转型对制造业企业投资效率的影响')
        self.assertEqual(path.read_bytes(),before)


if __name__=='__main__':unittest.main()

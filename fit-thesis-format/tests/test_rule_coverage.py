"""Rule coverage must require positive structural/native evidence."""
from pathlib import Path
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import rule_coverage as R
import thesis_format as F


class CoverageTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.base=Path(self.temp.name)
        self.path=self.base/'source.docx'
        d=Document();d.add_paragraph('正文');d.save(self.path)
        self.mapping={'source_sha256':F.sha(self.path.read_bytes()),'paragraphs':{'p000001':{'role':'body'}},'tables':{}}

    def tearDown(self):self.temp.cleanup()

    def results(self,**kwargs):
        report=R.evaluate(self.path,self.mapping,{'issues':[]},True,**kwargs)
        return {r['rule_id']:r for r in report['rule_results']}

    def test_empty_audit_does_not_prove_front_or_toc(self):
        rows=self.results()
        self.assertEqual(rows['front.components']['status'],'fail')
        self.assertEqual(rows['toc.two_levels']['status'],'fail')
        self.assertEqual(rows['section.pagination']['status'],'fail')
        self.assertEqual(rows['headings.numbering']['status'],'fail')

    def test_absent_panel_and_math_are_not_applicable(self):
        rows=self.results()
        for rid in ('table.panel','math.layout','math.editability'):
            self.assertEqual(rows[rid]['status'],'not_applicable')

    def test_missing_dual_native_evidence_blocks(self):
        rows=self.results()
        self.assertEqual(rows['delivery.dual_engine']['status'],'fail')

    def test_blank_appendix_spacer_fails_but_confirmed_code_blank_is_preserved(self):
        d=Document();d.add_paragraph('绪论');d.add_paragraph('附录');d.add_paragraph('');d.save(self.path)
        self.mapping['source_sha256']=F.sha(self.path.read_bytes())
        self.mapping['paragraphs']={'p000001':'chapter','p000002':'appendix_heading','p000003':'body'}
        evidence={key:{'all_pages_passed':True} for key in ('word','wps')}
        evidence['same_final_docx_sha256']=F.sha(self.path.read_bytes())
        spacer=self.results(dual_evidence=evidence)['whitespace.layout']
        self.assertEqual(spacer['status'],'fail')
        self.assertEqual(spacer['observed']['removable_body_spacer_indices'],[2])
        self.mapping['paragraphs']['p000003']='appendix_code'
        self.assertEqual(self.results(dual_evidence=evidence)['whitespace.layout']['status'],'pass')

    def test_same_final_hash_is_required(self):
        evidence={key:{'all_pages_passed':True} for key in ('word','wps')}
        evidence['same_final_docx_sha256']='stale'
        self.assertEqual(self.results(dual_evidence=evidence)['delivery.dual_engine']['status'],'fail')
        evidence['same_final_docx_sha256']=F.sha(self.path.read_bytes())
        self.assertEqual(self.results(dual_evidence=evidence)['delivery.dual_engine']['status'],'pass')

    def test_native_math_count_required_even_after_page_pass(self):
        d=Document(self.path);m=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x';r.append(t);m.append(r);d.paragraphs[0]._p.append(m);d.save(self.path)
        self.mapping['source_sha256']=F.sha(self.path.read_bytes())
        evidence={key:{'all_pages_passed':True,'native_math_count':0} for key in ('word','wps')}
        evidence['same_final_docx_sha256']=F.sha(self.path.read_bytes())
        self.assertEqual(self.results(dual_evidence=evidence)['math.editability']['status'],'fail')
        for key in ('word','wps'):evidence[key]['native_math_count']=1
        self.assertEqual(self.results(dual_evidence=evidence)['math.editability']['status'],'pass')

    def test_unknown_enabled_rule_blocks(self):
        registry=json.loads((R.ROOT/'references/rules-v10.json').read_text(encoding='utf-8'))
        registry['rules'].append({'id':'new.unimplemented.rule'})
        (self.base/'references').mkdir();(self.base/'references/rules-v10.json').write_text(json.dumps(registry),encoding='utf-8')
        with patch.object(R,'ROOT',self.base):
            self.assertEqual(self.results()['new.unimplemented.rule']['status'],'fail')

    def test_xml_comment_is_not_an_equation_object(self):
        from lxml import etree
        d=Document(self.path);d.paragraphs[0]._p.append(etree.Comment('source annotation'));d.save(self.path)
        self.mapping['source_sha256']=F.sha(self.path.read_bytes())
        self.assertEqual(self.results()['math.editability']['status'],'not_applicable')


if __name__=='__main__':unittest.main()

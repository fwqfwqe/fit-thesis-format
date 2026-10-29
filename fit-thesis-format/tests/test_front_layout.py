import sys,tempfile,unittest
from pathlib import Path
from docx import Document
from lxml import etree as E
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F,format_engine as G

class FrontLayoutTests(unittest.TestCase):
    def test_english_title_is_not_chinese_abstract_body(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'a.docx';d=Document();d.add_paragraph('这是中文论文题目');d.add_paragraph('摘要');d.add_paragraph('摘要正文。');d.add_paragraph('关键词：研究，实证')
            d.add_paragraph('Effect of Policy on Corporate Investment Efficiency: Empirical Evidence.')
            d.add_paragraph('Abstract');d.add_paragraph('English abstract.');d.save(p)
            inspection=G.inspect_document(p)
            roles=[v['role'] for v in inspection['map']['paragraphs'].values() if 'Effect of Policy' in v.get('expected_text','')]
            self.assertEqual(roles,['en_title'])
    def test_toc_right_aligned_dot_stop_in_named_styles(self):
        root=E.Element(F.q('styles'));F.make_styles(root)
        width=F.SPEC['page']['width_twips']-F.SPEC['page']['left_twips']-F.SPEC['page']['right_twips']
        for role in ('toc1','toc2'):
            tab=root.find(f'w:style[@w:styleId="FIT_{role}"]/w:pPr/w:tabs/w:tab',F.NS)
            self.assertEqual(tab.get(F.q('val')),'right');self.assertEqual(tab.get(F.q('pos')),str(width));self.assertEqual(tab.get(F.q('leader')),'dot')
    def test_named_appendix_is_a_heading(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'a.docx';d=Document();d.add_paragraph('附录 长安汽车财务绩效评价计算过程');d.save(p)
            roles=G.inspect_document(p)['map']['paragraphs']
            self.assertEqual(roles['p000001']['role'],'appendix_heading')
    def test_mixed_commitment_uses_separate_semantic_paragraphs(self):
        root=E.Element(F.q('body'));p=E.SubElement(root,F.q('p'))
        for n,text in enumerate(['学生承诺','我承诺遵守学校规定。','       学生（签名）：','      2026年10月28日']):
            if n:E.SubElement(E.SubElement(p,F.q('r')),F.q('br'))
            E.SubElement(E.SubElement(p,F.q('r')),F.q('t')).text=text
        roles={p:'commitment_body'};ops=[];G._commitment_blocks(root,roles,ops)
        self.assertEqual([roles[x] for x in root],['commitment_label','commitment_body','commitment_signature','commitment_signature'])
        self.assertEqual(G.text(root[-1]),'2026年10月28日')
        self.assertEqual(ops[0]['kind'],'commitment_split')

    def test_commitment_split_keeps_source_bound_citation_targets(self):
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.docx';out=Path(td)/'out.docx';d=Document()
            d.add_paragraph('测试论文题目');d.add_paragraph('诚信承诺书')
            d.add_paragraph('学生承诺\n我承诺遵守学校规定。\n学生（签名）：\n2026年10月28日')
            d.add_paragraph('摘要');d.add_paragraph('摘要正文。');d.add_heading('1 绪论',1)
            claim='张三（2020）讨论了相关问题';d.add_paragraph(claim+'。')
            d.add_heading('参考文献',1);ref='[1] 张三. 相关问题[J]. 期刊, 2020.';d.add_paragraph(ref);d.save(source)
            inv=G.inspect_document(source);ps=inv['map']['paragraphs']
            pid=next(k for k,v in ps.items() if v['expected_text']==claim+'。')
            rid=next(k for k,v in ps.items() if v['expected_text']==ref)
            dec={'source_sha256':inv['source_sha256'],'citation_insertions':[{'pid':pid,'expected_text':claim+'。','position':len(claim),'text':'[1]','claim':{'start':0,'end':len(claim),'expected_text':claim},'references':[{'pid':rid,'reference_number':1,'expected_text':ref}]}]}
            source_ps=F.indexed(F.parse(F.package(source)['word/document.xml']))[0]
            pledge_pid=next(k for k,p in source_ps.items() if '学生承诺' in G.text(p))
            dec['paragraphs']={pledge_pid:{'role':'commitment_body','expected_text':G.text(source_ps[pledge_pid])}}
            G.format_document(source,out,dec)
            final=F.parse(F.package(out)['word/document.xml'])
            self.assertIn(claim+'[1]。',G.text(final))
            self.assertTrue(any(G.text(p)=='学生承诺' for p in final.iter(F.q('p'))))

if __name__=='__main__':unittest.main()

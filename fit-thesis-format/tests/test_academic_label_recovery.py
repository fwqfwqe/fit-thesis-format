import sys,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F,format_engine as G
from integrity import verify_transition

class LabelRecovery(unittest.TestCase):
    def test_raw_numbers_toc_and_render_cache_recover(self):
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'raw.docx';out=Path(td)/'final.docx'
            d=Document();d.add_paragraph('测试论文题目',style='Title');d.add_paragraph('经管类研究论文　　含模拟计算')
            d.add_paragraph('摘要');d.add_paragraph('摘要正文。');d.add_paragraph('目录')
            d.add_heading('1 绪论',1);d.add_paragraph('表1报告数据，公式（1）定义变量。')
            p=d.add_paragraph('表1 数据表');p.runs[0]._r.insert(0,OxmlElement('w:lastRenderedPageBreak'))
            t=d.add_table(rows=2,cols=2);t.cell(0,0).text='变量';t.cell(1,0).text='x'
            shd=OxmlElement('w:shd');shd.set(F.q('fill'),'CCCCCC');t.cell(0,0)._tc.get_or_add_tcPr().append(shd)
            p=d.add_paragraph();m=OxmlElement('m:oMath');r=OxmlElement('m:r');v=OxmlElement('m:t');v.text='x';r.append(v);m.append(r);p._p.append(m);p.add_run('　（1）')
            d.save(source);i=G.inspect_document(source)
            decisions={'source_sha256':i['source_sha256'],'paragraphs':{}}
            for c in i['candidates']:
                decisions['paragraphs'][c['id']]={'role':'cn_title' if c['text']=='测试论文题目' else 'body','expected_text':c['text']}
            report=G.format_document(source,out,decisions);root=F.parse(F.package(out)['word/document.xml'])
            self.assertIn('表1-1报告数据，公式（1-1）定义变量。',G.text(root))
            self.assertNotIn('1 绪论',G.text(root));self.assertIn('含模拟计算',G.text(root))
            self.assertFalse(root.findall('.//w:tcPr/w:shd',F.NS))
            check=verify_transition(source,out,'format',report)
            self.assertTrue(check['passed'],check)
            # A replay is actually checked; forged label evidence cannot pass.
            label=next(x for x in report['operations'] if x['kind']=='academic_label')
            from academic_labels import verified
            label['after']='捏造内容'
            self.assertFalse(verified(label,F.parse(F.package(source)['word/document.xml']),report['source_mapping']))

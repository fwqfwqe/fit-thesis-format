import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import thesis_format as F
import format_engine as G
from integrity import verify_transition


class TailAndReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def sample(self, appendix=False):
        d = Document()
        d.add_paragraph('测试论文题目')
        d.add_paragraph('摘要')
        d.add_paragraph('摘要正文。')
        d.add_paragraph('目录')
        d.add_heading('一、绪论', 1)
        p = d.add_paragraph('原引文')
        G.field(p._p, ' REF _MissingTarget \\r \\h \\* MERGEFORMAT ', '[1]')
        p.add_run('保留的结论。')
        d.add_heading('参考文献', 1)
        d.add_paragraph('[1] 张三. 文献[J].期刊,2020.')
        d.add_heading('致谢', 1)
        d.add_paragraph('感谢导师。')
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if appendix:
            d.add_heading('附录', 1)
            d.add_paragraph('附录数据 123。')
        source = self.path / 'source.docx'
        d.save(source)
        inspected = G.inspect_document(source)
        title_id = next(k for k, v in inspected['map']['paragraphs'].items()
                        if v['expected_text'] == '测试论文题目')
        decisions = {'source_sha256': F.sha(source.read_bytes()),
            'paragraphs': {title_id: {'role': 'cn_title', 'expected_text': '测试论文题目'}},
            'metadata': {'title': '测试论文题目'}}
        return source, decisions

    def test_appendix_reorder_removes_only_terminal_separator(self):
        source, decisions = self.sample(appendix=True)
        result = self.path / 'a.docx'
        again = self.path / 'b.docx'
        source_bytes = source.read_bytes()
        report = G.format_document(source, result, decisions)
        root = F.parse(F.package(result)['word/document.xml'])
        body = root.find(F.q('body'))
        self.assertLess(G.text(root).index('附录数据'), G.text(root).index('感谢导师'))
        self.assertEqual(G.text(body[-2]), '感谢导师。')
        self.assertEqual(body[-1].tag, F.q('sectPr'))
        ops = [o for o in report['operations'] if o['kind'] == 'remove_terminal_page_break']
        self.assertEqual(len(ops), 1)
        self.assertEqual(ops[0]['page_breaks'], 1)
        self.assertEqual(len(ops[0]['before_sha256']), 64)
        self.assertEqual(source_bytes, source.read_bytes())
        self.assertTrue(verify_transition(source, result, 'format', report)['passed'])
        G.format_document(result, again)
        self.assertEqual(F.package(result), F.package(again))

    def test_terminal_cleanup_respects_text_controls_and_anchors(self):
        for element in ('bookmarkStart', 'commentReference', 'footnoteReference',
                        'sectPr', 'ins', 'fldChar', 'drawing'):
            with self.subTest(element=element):
                d = Document()
                p = d.add_paragraph()
                p.add_run().add_break(WD_BREAK.PAGE)
                target = p._p
                if element == 'sectPr':
                    F.child(F.ensure(target, 'pPr', True), element)
                else:
                    F.child(target, element)
                root = F.parse(F.xml(d._element))
                before = F.xml(root)
                G._trim_terminal_page_breaks(root, {}, [])
                self.assertEqual(before, F.xml(root))
        for text, kind in [('正文', WD_BREAK.PAGE), ('', WD_BREAK.LINE)]:
            d = Document()
            d.add_paragraph(text).add_run().add_break(kind)
            root = F.parse(F.xml(d._element))
            before = F.xml(root)
            G._trim_terminal_page_breaks(root, {}, [])
            self.assertEqual(before, F.xml(root))
        d = Document()
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_paragraph('后续正文')
        root = F.parse(F.xml(d._element))
        before = F.xml(root)
        G._trim_terminal_page_breaks(root, {}, [])
        self.assertEqual(before, F.xml(root))

    def test_strict_mode_retains_terminal_manual_break(self):
        source, decisions = self.sample()
        result = self.path / 'strict.docx'
        report = G.format_document(source, result, decisions, {'strict_text': True})
        root = F.parse(F.package(result)['word/document.xml'])
        self.assertTrue(root.findall('.//w:br[@w:type="page"]', F.NS))
        self.assertFalse(any(o['kind'] == 'remove_terminal_page_break' for o in report['operations']))

    def test_repaired_ref_has_native_character_format_and_intact_payload(self):
        source, decisions = self.sample()
        result = self.path / 'a.docx'
        again = self.path / 'b.docx'
        report = G.format_document(source, result, decisions)
        root = F.parse(F.package(result)['word/document.xml'])
        instruction = next(x for x in root.iter(F.q('instrText'))
                           if 'REF FITv10Reference1Full' in (x.text or ''))
        self.assertIn('\\* CHARFORMAT', instruction.text)
        self.assertNotIn('\\r', instruction.text)
        self.assertEqual(instruction.getparent().find('w:rPr/w:vertAlign', F.NS).get(F.q('val')), 'superscript')
        body = next(p for p in root.iter(F.q('p')) if G.text(p).startswith('原引文'))
        self.assertEqual(G.text(body), '原引文[1]保留的结论。')
        cache = next(r for r in body.findall(F.q('r')) if G.text(r) == '[1]')
        self.assertEqual(cache.find('w:rPr/w:rStyle', F.NS).get(F.q('val')), 'FIT_citation_superscript')
        self.assertIn('[1] 张三. 文献[J].期刊,2020.', G.text(root))
        self.assertTrue(verify_transition(source, result, 'format', report)['passed'])
        G.format_document(result, again)
        self.assertEqual(F.package(result), F.package(again))


if __name__ == '__main__':
    unittest.main()

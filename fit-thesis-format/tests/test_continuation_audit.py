"""Continuation recognition must prove a captioned compatible table chain."""
import tempfile
from pathlib import Path
import sys
import unittest
import zipfile

from docx import Document
from lxml import etree as E

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import thesis_format as F
import format_engine as G
from format_audit import Audit


class ContinuationAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.docx'
        self.output = self.root / 'formatted.docx'
        doc = Document()
        doc.add_paragraph('测试论文题目')
        doc.add_paragraph('摘要')
        doc.add_paragraph('摘要正文。')
        doc.add_heading('1 绪论', 1)
        doc.add_paragraph('表1-1 模型估计结果')
        table = doc.add_table(rows=7, cols=2)
        for i, row in enumerate(table.rows):
            row.cells[0].text = '变量' if i == 0 else '变量' + str(i)
            row.cells[1].text = '估计值' if i == 0 else str(i / 10)
        doc.save(self.source)
        table_node = F.indexed(F.parse(F.package(self.source)['word/document.xml']))[1]['t0001']
        decisions = {'source_sha256': F.sha(self.source.read_bytes()), 'tables': {
            't0001': {'kind': 'three-line', 'header_rows': 1, 'continuation_breaks': [3, 5],
                      'expected_sha256': F.sha(F.xml(table_node))}}}
        report = G.format_document(self.source, self.output, decisions)
        self.assertTrue(any(op['kind'] == 'table_continuation' for op in report['operations']))
        self.parts = F.package(self.output)
        self.document = F.parse(self.parts['word/document.xml'])
        self.tables = [t for t in self.document.iter(F.q('tbl')) if '估计值' in G.text(t)]
        self.assertEqual(len(self.tables), 3)

    def audit(self):
        target = self.root / 'audit.docx'
        self.parts['word/document.xml'] = F.xml(self.document)
        with zipfile.ZipFile(target, 'x', zipfile.ZIP_DEFLATED) as archive:
            for name, data in self.parts.items():
                archive.writestr(name, data)
        mapping = G.inspect_document(target)['map']
        audit = Audit(target, mapping)
        audit.infer_roles()
        audit.check_tables()
        return audit

    def caption_failures(self, audit):
        return [x for x in audit.issues if x['category'] == 'table.caption_missing_or_position']

    def test_source_bound_multi_segment_split_is_accepted(self):
        audit = self.audit()
        self.assertFalse(self.caption_failures(audit))
        evidence = audit.evidence['table_continuations']
        self.assertEqual(len(evidence), 2)
        self.assertEqual([len(x['links']) for x in evidence], [1, 2])
        self.assertEqual(len({x['origin_table_id'] for x in evidence}), 1)
        self.assertTrue(all('模型估计结果' in x['origin_caption'] for x in evidence))

    def test_unmanaged_words_do_not_exempt_missing_caption(self):
        marker = self.tables[1].getprevious()
        marker.find('w:pPr/w:pStyle', F.NS).set(F.q('val'), 'Normal')
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)
        self.assertFalse(audit.evidence['table_continuations'])

    def test_orphan_managed_marker_is_rejected(self):
        self.tables[0].getparent().remove(self.tables[0])
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)

    def test_missing_original_caption_cannot_be_hidden_by_continuation(self):
        self.tables[0].getprevious().find('.//w:t', F.NS).text = '没有编号的普通说明'
        # The managed caption has a native field; remove that field too so this
        # is truly an unnumbered prose paragraph rather than a stale cache.
        caption = self.tables[0].getprevious()
        for child in list(caption):
            if child.tag != F.q('pPr'):
                caption.remove(child)
        G.plain_run(caption, '没有编号的普通说明')
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 3)

    def test_mismatched_repeated_header_is_rejected(self):
        self.tables[1].find('w:tr/w:tc/w:p/w:r/w:t', F.NS).text = '其他指标'
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)

    def test_changed_grid_is_rejected_even_with_same_header_text(self):
        grid = self.tables[1].find('w:tblGrid/w:gridCol', F.NS)
        grid.set(F.q('w'), str(int(grid.get(F.q('w'))) + 10))
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)

    def test_repeated_header_flag_is_required(self):
        header = self.tables[1].find('w:tr/w:trPr/w:tblHeader', F.NS)
        header.getparent().remove(header)
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)

    def test_intervening_content_breaks_table_chain(self):
        self.tables[1].getprevious().addprevious(E.Element(F.q('p')))
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 2)

    def test_ordinary_table_without_caption_still_reports(self):
        for table in self.tables[1:]:
            marker = table.getprevious()
            table.getparent().remove(marker)
            table.getparent().remove(table)
        caption = self.tables[0].getprevious()
        caption.getparent().remove(caption)
        audit = self.audit()
        self.assertEqual(len(self.caption_failures(audit)), 1)


if __name__ == '__main__':
    unittest.main()

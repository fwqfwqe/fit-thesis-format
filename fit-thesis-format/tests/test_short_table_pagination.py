"""Compact table pagination policy; native Office verification is separate."""
from pathlib import Path
import sys
import unittest

from docx import Document
from lxml import etree as E

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import format_engine as G
import thesis_format as F


class ShortTablePaginationTests(unittest.TestCase):
    def table(self,rows=10,columns=4,value='0.004046***\n(0.001536)'):
        doc=Document();doc.add_paragraph('表1-1 回归结果')
        table=doc.add_table(rows=rows,cols=columns)
        for row in table.rows:
            for cell in row.cells:cell.text=value
        root=F.parse(F.xml(doc._element))
        return root,root.find('.//w:tbl',F.NS)

    def states(self,table):
        return [[p.find('w:pPr/w:keepNext',F.NS).get(F.q('val')) if p.find('w:pPr/w:keepNext',F.NS) is not None else None
                 for p in row.findall('./w:tc/w:p',F.NS)] for row in table.findall(F.q('tr'))]

    def test_ten_row_coefficient_table_moves_as_one_and_stops_before_notes(self):
        doc,t=self.table()
        before=G.text(t)
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        states=self.states(t)
        self.assertTrue(all(v=='1' for row in states[:-1] for v in row))
        self.assertTrue(all(v=='0' for v in states[-1]))
        self.assertEqual(G.text(t),before)
        first=F.xml(t)
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        self.assertEqual(F.xml(t),first)

    def test_small_row_count_with_long_cells_is_allowed_to_paginate(self):
        doc,t=self.table(rows=5,value='这段文字相当长，需要自动换行。'*40)
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        states=self.states(t)
        self.assertTrue(all(v=='1' for v in states[0]))
        self.assertTrue(all(v=='0' for row in states[1:] for v in row))

    def test_large_table_does_not_form_an_unbounded_chain(self):
        doc,t=self.table(rows=40,value='12.345')
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        self.assertTrue(all(v=='0' for row in self.states(t)[1:] for v in row))
        self.assertIsNotNone(t.find('w:tr/w:trPr/w:tblHeader',F.NS))

    def test_reviewed_continuation_boundaries_take_precedence(self):
        for label in ('explicit','previous','following'):
            with self.subTest(label=label):
                doc,t=self.table(rows=4,value='12.345')
                setting={'kind':'three-line','header_rows':1}
                if label=='explicit':setting['continuation_breaks']=[2]
                else:
                    p=E.Element(F.q('p'));G.plain_run(p,'续前表')
                    if label=='previous':t.addprevious(p)
                    else:t.addnext(p)
                G._table(t,setting,9355)
                self.assertTrue(all(v is None for row in self.states(t)[1:] for v in row))

    def test_merged_cells_and_nested_tables_are_not_grouped(self):
        doc,t=self.table(rows=4,value='12.345')
        cell=t.find('w:tr/w:tc',F.NS)
        F.child(F.ensure(cell,'tcPr',True),'vMerge',val='restart')
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        self.assertTrue(all(v=='0' for row in self.states(t)[1:] for v in row))
        doc,t=self.table(rows=4,value='12.345')
        nested=E.SubElement(t.findall(F.q('tr'))[1].find(F.q('tc')),F.q('tbl'))
        E.SubElement(E.SubElement(E.SubElement(nested,F.q('tr')),F.q('tc')),F.q('p'))
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        self.assertTrue(all(v=='0' for row in self.states(t)[1:] for v in row))

    def test_explicit_breaks_and_large_font_change_height_estimate(self):
        doc,t=self.table(rows=8,value='甲\n乙\n丙\n丁')
        for p in t.findall('.//w:p',F.NS):
            for r in p.findall(F.q('r')):F.child(F.ensure(r,'rPr',True),'sz',val=48)
        G._table(t,{'kind':'three-line','header_rows':1},9355)
        self.assertTrue(all(v=='0' for row in self.states(t)[1:] for v in row))


if __name__=='__main__':unittest.main()

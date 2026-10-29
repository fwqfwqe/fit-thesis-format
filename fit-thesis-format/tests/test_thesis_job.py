"""Orchestration tests use explicit mock Office calls, never claim native QA."""
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import thesis_job as J
import delivery_check as D
from docx import Document


class JobTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / 'input.docx'
        doc = Document(); doc.add_paragraph('受保护的论文正文'); doc.save(self.source)
        self.work = self.root / 'job'

    def tearDown(self):
        self.temp.cleanup()

    def start_blocked(self):
        with patch.object(J, 'doctor', return_value={'dual_render_capable': False}):
            return J.new_job({'source': str(self.source), 'work_dir': str(self.work)})

    def test_missing_wps_stops_before_format(self):
        result = self.start_blocked()
        self.assertEqual(result['status'], 'capability_required')
        self.assertFalse(result['delivery_ready'])
        self.assertFalse(list(self.work.glob('attempt-*')))
        self.assertEqual(J.sha(self.source), J.sha(self.work / 'original.docx'))

    def test_source_mutation_invalidates_job(self):
        self.start_blocked()
        self.source.write_bytes(self.source.read_bytes() + b'drift')
        with self.assertRaisesRegex(ValueError, 'Original input changed'):
            J.load_state(self.work)

    def test_original_copy_mutation_invalidates_job(self):
        self.start_blocked()
        (self.work / 'original.docx').write_bytes(b'drift')
        with self.assertRaisesRegex(ValueError, 'Immutable original'):
            J.load_state(self.work)

    def test_rule_drift_invalidates_job(self):
        self.start_blocked()
        with patch.object(J, 'rules_fingerprint', return_value={'different': 'rules'}):
            with self.assertRaisesRegex(ValueError, 'Rules or code changed'):
                J.load_state(self.work)

    def test_reject_existing_or_input_output(self):
        for path in (self.source, self.root):
            with self.assertRaises(ValueError):
                J.new_job({'source': str(self.source), 'work_dir': str(path)})
        with self.assertRaises(ValueError):
            J.bound_path(self.root, '../outside.docx')

    def test_atomic_state_and_resume_reprobe(self):
        self.start_blocked()
        with patch.object(J, 'doctor', return_value={'dual_render_capable': False}) as probe:
            result = J.execute({'action': 'resume', 'work_dir': str(self.work)})
            probe.assert_called_once()
            self.assertEqual(result['status'], 'capability_required')
        self.assertFalse(list(self.work.glob('*.tmp')))

    def test_final_mapping_mandatory_and_complete(self):
        with self.assertRaisesRegex(ValueError, 'Final semantic map'):
            J.verify_mapping_complete(self.source, None)
        with self.assertRaisesRegex(ValueError, 'Unclassified'):
            J.verify_mapping_complete(self.source, {'source_sha256': J.sha(self.source)})
        mapping = {'source_sha256': J.sha(self.source), 'paragraphs': {'p000001': {'role': 'body'}}, 'tables': {}}
        self.assertTrue(J.verify_mapping_complete(self.source, mapping)['passed'])

    def test_field_refresh_retains_reviewed_table_choices(self):
        import format_engine as E
        doc = Document(); doc.add_paragraph('正文内容')
        table = doc.add_table(rows=3, cols=2)
        for i, row in enumerate(table.rows):
            row.cells[0].text = '项目' + str(i)
            row.cells[1].text = '这是一段用于证明重新推断会得到复杂表的长文本。' * 6
        doc.save(self.source)
        mapping = E.inspect_document(self.source)['map']
        mapping['paragraphs']['p000001'] = {'role': 'body', 'expected_text': '正文内容'}
        choice = {'kind': 'three-line', 'header_rows': 2, 'column_widths_twips': [2000, 6000],
                  'review_evidence': 'confirmed from source table'}
        mapping['tables']['t0001'] = choice
        final = self.root / 'cache-refreshed.docx'; shutil.copyfile(self.source, final)
        self.assertEqual(E.inspect_document(final)['map']['tables']['t0001']['kind'], 'complex')
        rebound = J.rebind_final_mapping(self.source, final, mapping)
        self.assertEqual(rebound['tables']['t0001'], choice)
        self.assertEqual(mapping['tables']['t0001'], choice)

    def test_mapping_rebind_rejects_changed_body(self):
        import format_engine as E
        mapping = E.inspect_document(self.source)['map']
        mapping['paragraphs']['p000001'] = {'role': 'body', 'expected_text': '受保护的论文正文'}
        final = self.root / 'changed.docx'
        doc = Document(self.source); doc.paragraphs[0].text = '不允许改写的正文'; doc.save(final)
        with self.assertRaisesRegex(ValueError, 'paragraph changed'):
            J.rebind_final_mapping(self.source, final, mapping)

    def test_mapping_rebind_handles_expanded_toc_without_losing_roles(self):
        import format_engine as E
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from lxml import etree
        doc = Document()
        p = doc.add_paragraph()
        for kind in ('begin', 'instruction', 'separate', 'text', 'end'):
            r = OxmlElement('w:r'); p._p.append(r)
            if kind == 'instruction':
                n = OxmlElement('w:instrText'); n.text = ' TOC \\o "1-2" '
            elif kind == 'text':
                n = OxmlElement('w:t'); n.text = '第一章 1'
            else:
                n = OxmlElement('w:fldChar'); n.set(qn('w:fldCharType'), kind)
            r.append(n)
        doc.add_paragraph('明确审定的正文'); doc.save(self.source)
        mapping = E.inspect_document(self.source)['map']
        mapping['paragraphs']['p000002'] = {'role': 'body', 'expected_text': '明确审定的正文', 'review_note': 'manual confirmation'}
        with zipfile.ZipFile(self.source) as z: parts = {n: z.read(n) for n in z.namelist()}
        root = etree.fromstring(parts['word/document.xml']); first = root.find('.//' + qn('w:p'))
        end_run = first[-1]; first.remove(end_run)
        added = OxmlElement('w:p'); r = OxmlElement('w:r'); n = OxmlElement('w:t'); n.text = '第二章 2'; r.append(n); added.append(r); added.append(end_run)
        first.addnext(added)
        parts['word/document.xml'] = etree.tostring(root)
        final = self.root / 'expanded.docx'
        with zipfile.ZipFile(final, 'w') as z:
            for name, data in parts.items(): z.writestr(name, data)
        rebound = J.rebind_final_mapping(self.source, final, mapping)
        self.assertEqual(rebound['paragraphs']['p000003']['role'], 'body')
        self.assertEqual(rebound['paragraphs']['p000003']['review_note'], 'manual confirmation')
        self.assertEqual(rebound['paragraphs']['p000001']['role'], 'toc1')
        self.assertEqual(rebound['paragraphs']['p000002']['role'], 'toc1')

    def test_unresolved_candidate_preserves_source_inspection(self):
        def fmt(source, output, **kwargs):
            shutil.copyfile(source, output)
            return {'final_mapping': {'source_sha256': J.sha(output), 'paragraphs': {}, 'tables': {}},
                    'unresolved_issues': [{'code': 'title_or_body_boundary_missing', 'reason': '需要确认正文起点'}]}
        with patch.object(J, 'doctor', return_value={'dual_render_capable': True}), patch('format_engine.format_document', side_effect=fmt), patch.object(J, 'verify_stage', return_value={'passed': True}):
            result = J.new_job({'source': str(self.source), 'work_dir': str(self.work)})
        self.assertEqual(result['status'], 'semantic_review_required')
        task = result['next'][0]
        self.assertEqual(task['inspection']['source_sha256'], J.sha(self.source))
        candidate = next(c for c in task['inspection']['candidates'] if c['id'] == 'p000001')
        self.assertEqual(candidate['text'], '受保护的论文正文')
        self.assertEqual(task['decision_template']['paragraphs']['p000001']['expected_text'], '受保护的论文正文')

    def test_duplicate_corrections_stop_without_changing_source(self):
        self.start_blocked()
        request = {'kind': 'decisions', 'decisions': {'source_sha256': J.sha(self.source), 'paragraphs': {}}}
        with patch.object(J, 'doctor', return_value={'dual_render_capable': False}):
            J.submit(self.work, request)
            J.submit(self.work, request)
            result = J.submit(self.work, request)
        self.assertEqual(result['status'], 'no_progress')
        self.assertFalse(result['delivery_ready'])

    def test_review_bound_to_current_file(self):
        self.start_blocked()
        state = J.load_state(self.work)
        state['renders'] = {'word': {}}
        state['final_sha256'] = J.sha(self.source)
        J.save_state(self.work, state)
        with self.assertRaisesRegex(ValueError, 'actual final DOCX'):
            J.submit(self.work, {'kind': 'review', 'engine': 'word', 'review': {'final_docx_sha256': 'stale'}})

    def test_missing_engine_is_not_delivery(self):
        with self.assertRaisesRegex(ValueError, 'Both Word and WPS'):
            D.check_dual_render(self.source, {'word': {}})

    def test_final_mutation_blocks_verify(self):
        self.start_blocked()
        state = J.load_state(self.work)
        state.update(final=str(self.source), final_sha256='stale')
        result = J.verify_job(self.work, state)
        self.assertFalse(result['delivery_ready'])
        self.assertEqual(result['blocking_items'][0]['check'], 'final')

    def test_fresh_audit_receives_final_map(self):
        self.start_blocked()
        final = self.work / 'final.docx'; shutil.copyfile(self.source, final)
        report = self.work / 'stage.json'; J.atomic_json(report, {})
        mapping = {'source_sha256': J.sha(final), 'paragraphs': {'p000001': {'role': 'body'}}, 'tables': {}}
        map_path = self.work / 'final-map.json'; J.atomic_json(map_path, mapping)
        state = J.load_state(self.work)
        state.update(final=str(final), final_sha256=J.sha(final), final_mapping=str(map_path),
                     stages=[{'kind':'format','source':str(self.source),'output':str(final),'output_sha256':J.sha(final),'report':str(report)}])
        J.save_state(self.work, state)
        with patch.object(J, 'verify_stage', return_value={'passed':True}), patch('format_audit.Audit') as audit, patch.object(D, 'check_dual_render', return_value={'passed':True}), patch('rule_coverage.evaluate', return_value={'passed':True,'rule_results':[]}):
            audit.return_value.run.return_value = {'issues':[]}
            result = J.verify_job(self.work)
            audit.assert_called_once_with(final, mapping=mapping)
            self.assertTrue(result['delivery_ready'])

    def test_full_job_pauses_for_real_review_then_verifies_both(self):
        """Exercise orchestration; native Office and format correctness are separate tests."""
        from reportlab.pdfgen.canvas import Canvas
        def mapping(path):
            return {'source_sha256':J.sha(path),'paragraphs':{'p000001':{'role':'body'}},'tables':{}}
        def fmt(source, output, decisions=None, profile=None):
            shutil.copyfile(source,output)
            return {'source_sha256':J.sha(source),'output_sha256':J.sha(output),'operations':[],
                    'final_mapping':mapping(output),'unresolved':[]}
        def render(engine, source, pdf, saved_docx=None, **kwargs):
            if saved_docx:shutil.copyfile(source,saved_docx)
            c=Canvas(str(pdf));c.drawString(40,800,'UNIT TEST MOCK ENGINE');c.save()
            return {'engine_key':engine,'engine':engine,'engine_version':'unit-test',
                    'engine_executable':'C:/Office/WINWORD.EXE' if engine=='word' else 'C:/Kingsoft/WPS/wps.exe',
                    'source_sha256_before':J.sha(source),'source_unchanged':True,'pdf':str(pdf),
                    'pdf_sha256':J.sha(pdf),'pages':1,'field_errors':[],'toc_entries':[]}
        with patch.object(J,'rules_fingerprint',return_value={'test':'fixed'}), patch.object(J,'doctor',return_value={'dual_render_capable':True}), patch('format_engine.format_document',side_effect=fmt), patch('format_engine.inspect_document',side_effect=lambda path:{'mapping':mapping(path)}), patch('office_probe.render',side_effect=render), patch.object(J,'verify_stage',return_value={'passed':True}), patch('format_audit.Audit') as audit, patch('rule_coverage.evaluate',return_value={'passed':True,'rule_results':[]}):
            audit.return_value.run.return_value={'issues':[]}
            result=J.new_job({'source':str(self.source),'work_dir':str(self.work)})
            self.assertEqual(result['status'],'visual_review_required')
            self.assertFalse(result['delivery_ready'])
            self.assertEqual({x['engine'] for x in result['next']},{'word','wps'})
            for engine in ('word','wps'):
                state=J.load_state(self.work)
                review=J.read(state['renders'][engine]['review_template'])
                review['reviewer']='UNIT TEST SYNTHETIC REVIEW'
                for page in review['pages']:
                    page.update(status='pass',checks={key:'pass' for key in J.PAGE_CHECKS},notes='Synthetic fixture observation; no real thesis acceptance claimed.')
                result=J.execute({'action':'submit','work_dir':str(self.work),'kind':'review','engine':engine,'review':review})
                if engine=='word':self.assertFalse(result['delivery_ready'])
            self.assertTrue(result['delivery_ready'])
            self.assertEqual(result['status'],'complete')


class EvidenceTests(unittest.TestCase):
    def setUp(self):
        from reportlab.pdfgen.canvas import Canvas
        import pypdfium2
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.docx = self.root / 'final.docx'
        d = Document(); d.add_paragraph('Native evidence fixture'); d.save(self.docx)
        self.pdf = self.root / 'page.pdf'
        c = Canvas(str(self.pdf)); c.drawString(40,800,'Evidence fixture'); c.save()
        pdf = pypdfium2.PdfDocument(str(self.pdf)); page = pdf[0]; image = page.render(scale=1).to_pil()
        self.png = self.root / 'page.png'; image.save(self.png); image.close(); page.close(); pdf.close()
        self.export = {'engine':'Microsoft Word','engine_key':'word','engine_version':'test', 'engine_executable':'C:/Office/WINWORD.EXE',
                       'source_sha256_before':J.sha(self.docx),'source_unchanged':True,'pdf':str(self.pdf),'pdf_sha256':J.sha(self.pdf),'pages':1,'toc_entries':[],'field_errors':[]}
        self.pages = {'engine':'pypdfium2','dpi':72,'input_sha256':J.sha(self.pdf),'page_count':1,'pages':[{'page':1,'file':str(self.png),'sha256':J.sha(self.png)}]}
        self.review = {'reviewer':'UNIT TEST SYNTHETIC FIXTURE','final_docx_sha256':J.sha(self.docx),'pdf_sha256':J.sha(self.pdf),
                       'pages':[{'page':1,'image_sha256':J.sha(self.png),'status':'pass','checks':{k:'pass' for k in D.PAGE_CHECKS},'notes':'Synthetic test fixture; not a claim about any real manuscript.'}]}
        self.rp,self.pp,self.vp=[self.root/x for x in ('render.json','pages.json','review.json')]
        self.write()

    def write(self):
        for path,value in ((self.rp,self.export),(self.pp,self.pages),(self.vp,self.review)):J.atomic_json(path,value)

    def tearDown(self):
        self.temp.cleanup()

    def test_real_pdf_and_image_hashes_checked(self):
        self.assertTrue(D.check_render(self.docx,self.rp,self.pp,self.vp)['evidence_consistent'])
        self.pages['pages'][0]['sha256']='stale'; self.write()
        with self.assertRaisesRegex(ValueError,'image differs'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)

    def test_missing_page_review_rejected(self):
        self.review['pages']=[];self.write()
        with self.assertRaisesRegex(ValueError,'Every page'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)

    def test_consistent_fake_png_hashes_do_not_prove_pdf_page(self):
        from PIL import Image
        with Image.open(self.png) as old: size=old.size
        Image.new('RGB',size,'white').save(self.png)
        self.pages['pages'][0]['sha256']=J.sha(self.png)
        self.review['pages'][0]['image_sha256']=J.sha(self.png)
        self.write()
        with self.assertRaisesRegex(ValueError,'pixels do not reproduce'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)

    def test_pending_page_rejected(self):
        self.review['pages'][0]['status']='pending';self.write()
        with self.assertRaisesRegex(ValueError,'not recorded'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)

    def test_pending_review_can_be_saved_without_becoming_pass(self):
        self.review['pages'][0].update(status='pending',checks={key:'pending' for key in D.PAGE_CHECKS},notes='')
        self.write()
        result=D.check_render(self.docx,self.rp,self.pp,self.vp,allow_incomplete=True)
        self.assertFalse(result['all_pages_passed'])
        self.assertEqual(result['pending_pages'],[1])

    def test_failed_review_keeps_concrete_defect(self):
        self.review['pages'][0]['status']='fail'
        self.review['pages'][0]['checks']['math_tables']='fail'
        self.review['pages'][0]['notes']='Synthetic test: formula numerator is cut off.'
        self.write()
        result=D.check_render(self.docx,self.rp,self.pp,self.vp,allow_incomplete=True)
        self.assertFalse(result['all_pages_passed'])
        self.assertEqual(result['failed_pages'],[1])
        with self.assertRaisesRegex(ValueError,'not recorded'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)

    def test_fake_wps_name_rejected(self):
        proof={'render_manifest':str(self.rp),'page_manifest':str(self.pp),'review':str(self.vp)}
        self.export['engine_key']='wps';self.write()
        with self.assertRaisesRegex(ValueError,'identity'):
            D.check_dual_render(self.docx,{'word':proof,'wps':proof})

    def test_native_table_inventory_cannot_be_omitted(self):
        doc=Document(self.docx);doc.add_table(rows=1,cols=1);doc.save(self.docx)
        with self.assertRaisesRegex(ValueError,'all document tables'):
            D.continuation_requirements({'pages':2,'table_spans':[]},self.docx)
        expected=D.continuation_requirements({'pages':2,'table_spans':[{'table_index':1,'start_page':1,'end_page':2}]},self.docx)
        self.assertEqual(expected,{2:[1]})

    def test_generic_page_pass_cannot_hide_missing_continuation_label(self):
        from pypdf import PdfReader,PdfWriter
        reader=PdfReader(str(self.pdf));writer=PdfWriter();writer.add_page(reader.pages[0]);writer.add_page(reader.pages[0])
        replacement=self.root/'two.pdf'
        with replacement.open('wb') as out:writer.write(out)
        self.pdf=replacement
        doc=Document(self.docx);doc.add_table(rows=1,cols=1);doc.save(self.docx)
        self.export.update(source_sha256_before=J.sha(self.docx),pdf=str(self.pdf),pdf_sha256=J.sha(self.pdf),pages=2,
                           table_spans=[{'table_index':1,'start_page':1,'end_page':2}])
        self.pages.update(input_sha256=J.sha(self.pdf),page_count=2)
        self.pages['pages'].append({**self.pages['pages'][0],'page':2})
        self.review.update(final_docx_sha256=J.sha(self.docx),pdf_sha256=J.sha(self.pdf))
        self.review['pages'].append({**self.review['pages'][0],'page':2})
        self.write()
        with self.assertRaisesRegex(ValueError,'exactly the native continued tables'):
            D.check_render(self.docx,self.rp,self.pp,self.vp)
        self.review['pages'][1]['continued_tables']=[{'table_index':1,'caption_label':'续前表',
                                                     'header_repeated':True,'no_missing_rows':True}]
        self.write()
        self.assertEqual(D.check_render(self.docx,self.rp,self.pp,self.vp)['continued_page_checks'],1)


if __name__ == '__main__':unittest.main()

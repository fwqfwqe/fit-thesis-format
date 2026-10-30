import sys,tempfile,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from format_audit import Audit

class MissingYearAuditGateTests(unittest.TestCase):
    def test_missing_year_reaches_final_audit_even_with_correct_numeric_marker(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'source.docx';d=Document()
            d.add_heading('一、绪论',1);d.add_paragraph('Vial的综述讨论了数字化转型[5]。')
            d.add_heading('参考文献',1);d.add_paragraph('[5] VIAL G. Digital transformation[J]. Journal, 2019.')
            d.save(path)
            issues=Audit(path).run()['issues']
            self.assertTrue(any(i['category']=='citation.missing_year' for i in issues))

if __name__=='__main__':unittest.main()

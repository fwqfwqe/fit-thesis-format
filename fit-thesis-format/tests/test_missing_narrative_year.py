"""A numbered author mention must not disappear merely because its year is absent."""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import narrative_citations as N


class MissingNarrativeYearTests(unittest.TestCase):
    def inspect(self, body, refs):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'source.docx'
            d = Document()
            d.add_paragraph(body)
            d.add_heading('参考文献', 1)
            for ref in refs:
                d.add_paragraph(ref)
            d.save(source)
            before = source.read_bytes()
            result = N.inspect(source)
            self.assertEqual(source.read_bytes(), before)
            return result

    def test_missing_year_detected_despite_uppercase_bibliography(self):
        r = self.inspect('Vial的综述讨论了数字化转型[5]。',
                         ['[5] VIAL G. Understanding digital transformation[J]. Journal, 2019.'])
        self.assertEqual(len(r['candidates']), 1)
        candidate = r['candidates'][0]
        self.assertEqual(candidate['expected_text'], 'Vial')
        self.assertIn('missing_year', candidate['issues'])
        self.assertIsNone(candidate['year'])
        self.assertFalse(candidate['reviewed'])
        self.assertEqual(candidate['possible_years'], ['2019'])

    def test_two_authors_are_one_complete_review_span(self):
        r = self.inspect('Brynjolfsson和Hitt讨论组织投入[7]。',
                         ['[7] BRYNJOLFSSON E, HITT L M. Beyond Computation[J]. Journal, 2000.'])
        self.assertEqual([c['expected_text'] for c in r['candidates']], ['Brynjolfsson和Hitt'])

    def test_foreign_et_al_period_is_not_a_sentence_end(self):
        refs = ['[6] BHARADWAJ A, EL SAWY O A, PAVLOU P A, et al. Digital Business Strategy[J]. Journal, 2013.']
        for author in ('Bharadwaj等', 'Bharadwaj et al.'):
            with self.subTest(author=author):
                r = self.inspect(author + '讨论数字战略[6]。', refs)
                self.assertEqual(len(r['candidates']), 1)
                self.assertEqual(r['candidates'][0]['expected_text'], author)

    def test_real_existing_year_is_not_missing_or_wrong_author_case(self):
        r = self.inspect('Vial（2019）讨论了转型[5]。',
                         ['[5] VIAL G. Digital transformation[J]. Journal, 2019.'])
        self.assertEqual(len(r['matches']), 1)
        self.assertEqual(r['candidates'], [])

    def test_organization_with_citation_is_flagged(self):
        r = self.inspect('OECD报告讨论能力约束[2]。',
                         ['[2] OECD. Digital Transformation of SMEs[M]. Paris, 2021.'])
        self.assertEqual(r['candidates'][0]['expected_text'], 'OECD')

    def test_no_automatic_assignment_from_uncited_name_or_word_substring(self):
        refs = ['[5] VIAL G. Digital transformation[J]. Journal, 2019.']
        for body in ('这是Vial文件夹。', 'Trivial结果需要解释[5]。', 'Vial模型已经讨论。其他研究有结论[5]。'):
            with self.subTest(body=body):
                self.assertEqual(self.inspect(body, refs)['candidates'], [])

    def test_publication_year_remains_a_choice_when_same_claim_cites_multiple_works(self):
        r = self.inspect('Vial的综述分别讨论转型[5,6]。',
                         ['[5] VIAL G. First work[J]. Journal, 2019.',
                          '[6] VIAL G. Second work[J]. Journal, 2021.'])
        self.assertEqual(len(r['candidates']), 1)
        self.assertIn('reference_ambiguous', r['candidates'][0]['issues'])
        self.assertEqual(r['candidates'][0]['possible_years'], ['2019', '2021'])


if __name__ == '__main__':
    unittest.main()

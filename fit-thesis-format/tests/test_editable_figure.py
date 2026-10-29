import sys,tempfile,unittest
from pathlib import Path
from docx import Document
from lxml import etree as E
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F,format_engine as G

class EditableFigureTests(unittest.TestCase):
    def test_inline_vml_text_does_not_make_figure_a_body_paragraph(self):
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.docx';out=Path(td)/'formatted.docx'
            d=Document();d.add_paragraph('摘要');d.add_paragraph('摘要正文。');d.add_heading('1 绪论',1)
            p=d.add_paragraph();r=E.SubElement(p._p,F.q('r'));pict=E.SubElement(r,F.q('pict'))
            group=E.SubElement(pict,'{urn:schemas-microsoft-com:vml}group',style='width:386pt;height:520pt')
            box=E.SubElement(group,'{urn:schemas-microsoft-com:vml}textbox');content=E.SubElement(box,F.q('txbxContent'))
            inner=E.SubElement(content,F.q('p'));E.SubElement(E.SubElement(inner,F.q('r')),F.q('t')).text='研究设计与实证分析'
            before=F.xml(pict);d.add_paragraph('图1-1 研究技术路线');d.save(source)
            inv=G.inspect_document(source)
            outer=next(v for v in inv['map']['paragraphs'].values() if v['expected_text']=='研究设计与实证分析')
            self.assertEqual(outer['role'],'figure')
            self.assertEqual(len(inv['map']['preserve']),1)
            G.format_document(source,out)
            parts=F.package(out);root=F.parse(parts['word/document.xml']);newpict=root.find('.//w:pict',F.NS)
            self.assertEqual(F.xml(newpict),before)
            outerp=newpict.getparent().getparent()
            self.assertEqual(outerp.find('w:pPr/w:pStyle',F.NS).get(F.q('val')),'FIT_figure')
            styles=F.parse(parts['word/styles.xml']);sp=styles.find('w:style[@w:styleId="FIT_figure"]/w:pPr/w:spacing',F.NS)
            self.assertEqual(sp.get(F.q('lineRule')),'atLeast')

if __name__=='__main__':unittest.main()

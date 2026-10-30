import sys,tempfile,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F,format_engine as G

class HeadingDecoration(unittest.TestCase):
    def test_heading_decoration_removed_body_emphasis_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.docx';final=Path(td)/'final.docx'
            d=Document();d.add_paragraph('测试论文题目',style='Title');d.add_paragraph('摘要');d.add_paragraph('正文。');d.add_paragraph('关键词：变量')
            h=d.add_heading('1 绪论',1);h.runs[0].underline=True;h.runs[0].italic=True
            p=d.add_paragraph('保留正文强调。');p.runs[0].italic=True;d.save(source)
            r=G.format_document(source,final);root=F.parse(F.package(final)['word/document.xml'])
            h=next(p for p in root.iter(F.q('p')) if G.text(p).startswith('绪论'))
            self.assertFalse(h.findall('.//w:rPr/w:u',F.NS));self.assertFalse(h.findall('.//w:rPr/w:i',F.NS))
            p=next(p for p in root.iter(F.q('p')) if G.text(p)=='保留正文强调。')
            self.assertTrue(p.findall('.//w:rPr/w:i',F.NS))
            self.assertEqual(F.SPEC['roles']['cn_keywords']['before_pt'],20)
            self.assertEqual(F.SPEC['roles']['en_keywords']['before_pt'],16)

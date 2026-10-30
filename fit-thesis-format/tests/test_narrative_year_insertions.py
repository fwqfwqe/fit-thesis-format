"""Source-bound year insertion, preservation and adversarial evidence tests."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from lxml import etree as E

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import citation_style as C
import format_engine as G
import narrative_citations as N
import narrative_years as Y


class NarrativeYearInsertionTests(unittest.TestCase):
    def make_parts(self, text='Vial的综述讨论转型[5]。', refs=None):
        refs = refs or ['[5] VIAL G. Digital transformation[J]. Journal, 2019.']
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'source.docx'
            doc = Document()
            doc.add_paragraph(text)
            doc.add_paragraph('参考文献')
            for ref in refs:
                doc.add_paragraph(ref)
            doc.save(path)
            return C.package(path)

    def records(self, parts):
        inspection = N.inspect_parts(parts)
        refs = {r['pid']: r for r in inspection['references']}
        ps = C.indexed(C.parse(parts[Y.DOC]))
        result = []
        for candidate in inspection['candidates']:
            if 'missing_year' not in candidate['issues']:
                continue
            ref = refs[candidate['reference_pids'][0]]
            result.append({'pid': candidate['pid'], 'expected_text': C.ptext(ps[candidate['pid']]),
                           'author_start': candidate['author_start'], 'author_end': candidate['author_end'],
                           'expected_author': candidate['expected_text'], 'reference_pid': ref['pid'],
                           'reference_expected_text': ref['expected_text'], 'year': ref['years'][0] if ref['years'] else '2019', 'reviewed': True})
        return result

    def apply(self, parts, records=None, role='body'):
        root = C.parse(parts[Y.DOC])
        ps = C.indexed(root)
        operations, issues = [], []
        Y.apply(parts, root, {ps['p000001']: role}, self.records(parts) if records is None else records,
                G._replace_plain_span, operations, issues)
        return root, operations, issues

    def test_unique_existing_year_inserted_without_prose_or_marker_changes(self):
        parts = self.make_parts()
        root, operations, issues = self.apply(parts)
        self.assertFalse(issues)
        self.assertEqual(C.ptext(C.indexed(root)['p000001']), 'Vial（2019）的综述讨论转型[5]。')
        self.assertEqual(len(operations), 1)
        self.assertTrue(Y.verified(operations[0], parts))
        self.assertEqual(C.ptext(C.indexed(root)['p000003']), C.ptext(C.indexed(C.parse(parts[Y.DOC]))['p000003']))
        self.assertEqual(root.xpath('//w:r[w:t="（2019）"]/w:rPr/w:vertAlign/@w:val', namespaces=C.NS), ['baseline'])

    def test_two_authors_receive_one_year_after_whole_author_span(self):
        parts = self.make_parts('Brynjolfsson和Hitt讨论组织投入[7]。',
                                ['[7] BRYNJOLFSSON E, HITT L M. Beyond Computation[J]. Journal, 2000.'])
        root, operations, issues = self.apply(parts)
        self.assertFalse(issues)
        self.assertEqual(C.ptext(C.indexed(root)['p000001']), 'Brynjolfsson和Hitt（2000）讨论组织投入[7]。')
        self.assertTrue(Y.verified(operations[0], parts))

    def test_publication_prefix_can_precede_untouched_doi_link_and_render_hint(self):
        parts=self.make_parts();root=C.parse(parts[Y.DOC]);ps=C.indexed(root)
        ps['p000001'].find(C.q('r')).insert(0,E.Element(C.q('lastRenderedPageBreak')))
        link=E.SubElement(ps['p000003'],C.q('hyperlink'))
        run=E.SubElement(link,C.q('r'));E.SubElement(run,C.q('t')).text=' DOI／来源'
        original_link=E.tostring(link);parts[Y.DOC]=C.serialize(root)
        result,operations,issues=self.apply(parts)
        self.assertFalse(issues);self.assertTrue(Y.verified(operations[0],parts))
        self.assertEqual(E.tostring(result.find('.//w:hyperlink',C.NS)),original_link)

    def test_two_sentences_in_same_paragraph_keep_original_offsets(self):
        parts = self.make_parts('Vial讨论转型[5]。Vial随后分析企业[5]。')
        root, operations, issues = self.apply(parts)
        self.assertFalse(issues)
        self.assertEqual(C.ptext(C.indexed(root)['p000001']), 'Vial（2019）讨论转型[5]。Vial（2019）随后分析企业[5]。')
        self.assertEqual(len(operations), 1)
        self.assertTrue(Y.verified(operations[0], parts))

    def test_stale_full_text_and_forged_year_are_rejected_without_mutation(self):
        parts = self.make_parts()
        for key, replacement in [('expected_text', 'Vial改写了结论[5]。'), ('year', '2020'),
                                  ('reference_expected_text', '伪造书目'), ('expected_author', 'Someone'),
                                  ('reviewed', False), ('author_end', 3)]:
            with self.subTest(key=key):
                records = self.records(parts)
                records[0][key] = replacement
                root, operations, issues = self.apply(parts, records)
                self.assertTrue(issues)
                self.assertFalse(operations)
                self.assertEqual(C.serialize(root), C.serialize(C.parse(parts[Y.DOC])))

    def test_multiple_bibliography_years_are_not_selected_automatically(self):
        parts = self.make_parts(refs=['[5] VIAL G. Transformation[J]. Journal, 2019, 2020.'])
        _, operations, issues = self.apply(parts)
        self.assertTrue(issues)
        self.assertFalse(operations)

    def test_identifier_digits_are_not_a_publication_year(self):
        parts = self.make_parts(refs=['[5] VIAL G. Transformation[J]. Journal. DOI:10.2019/abc.'])
        _, operations, issues = self.apply(parts)
        self.assertTrue(issues)
        self.assertFalse(operations)

    def test_multiwork_sentence_and_duplicate_markers_are_rejected(self):
        for text in ('Vial讨论转型[5,6]。', 'Vial讨论转型[5][5]。'):
            with self.subTest(text=text):
                parts = self.make_parts(text)
                _, operations, issues = self.apply(parts)
                self.assertTrue(issues)
                self.assertFalse(operations)

    def test_ambiguous_reference_number_is_rejected(self):
        parts = self.make_parts(refs=['[5] VIAL G. First work[J]. Journal, 2019.',
                                      '[5] SMITH G. Other work[J]. Journal, 2020.'])
        _, operations, issues = self.apply(parts)
        self.assertTrue(issues)
        self.assertFalse(operations)

    def test_bookmark_and_code_role_are_protected(self):
        parts = self.make_parts()
        records = self.records(parts)
        _, operations, issues = self.apply(parts, records, role='appendix_code')
        self.assertTrue(issues)
        self.assertFalse(operations)
        root = C.parse(parts[Y.DOC])
        C.indexed(root)['p000001'].append(E.Element(C.q('bookmarkStart'), {C.q('id'): '1', C.q('name'): 'KeepMe'}))
        parts[Y.DOC] = C.serialize(root)
        _, operations, issues = self.apply(parts, records)
        self.assertTrue(issues)
        self.assertFalse(operations)

    def test_verifier_rejects_operation_tampering(self):
        parts = self.make_parts()
        _, operations, _ = self.apply(parts)
        for key, value in [('before', '虚假原文'), ('after', 'Vial（2019）强化结论[5]。'), ('locator', 'p000003')]:
            operation = copy.deepcopy(operations[0])
            operation[key] = value
            self.assertFalse(Y.verified(operation, parts))
        operation = copy.deepcopy(operations[0])
        operation['records'][0]['year'] = '2020'
        self.assertFalse(Y.verified(operation, parts))

    def test_duplicate_records_and_existing_year_are_rejected(self):
        parts = self.make_parts()
        records = self.records(parts)
        _, operations, issues = self.apply(parts, records + records)
        self.assertTrue(issues)
        self.assertFalse(operations)
        formatted = self.make_parts('Vial（2019）的综述讨论转型[5]。')
        records[0]['expected_text'] = 'Vial（2019）的综述讨论转型[5]。'
        _, operations, issues = self.apply(formatted, records)
        self.assertTrue(issues)
        self.assertFalse(operations)


if __name__ == '__main__':
    unittest.main()

import sys,tempfile,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import narrative_citations as N
from format_audit import Audit


class BibliographyPublicationEvidenceTests(unittest.TestCase):
    def test_long_organization_is_one_author_even_when_name_contains_he(self):
        ref='[20] 工业和信息化部办公厅. 关于印发中小企业数字化转型指南的通知[EB/OL]. (2022-11-03)[2026-10-30]. DOI／来源链接：https://example.org/2026'
        result=N._guess_reference(ref)
        self.assertEqual(result['author_kind'],'organization')
        self.assertEqual(len(result['authors']),1)
        self.assertEqual(result['authors'][0]['exact_name'],'工业和信息化部办公厅')
        self.assertEqual(result['years'],['2022'])

    def test_personal_coauthors_and_bad_separators_are_not_waived(self):
        ref='[1] 张三和李四. 研究问题[J]. 期刊,2020.'
        result=N._guess_reference(ref);self.assertEqual(len(result['authors']),2)
        self.assertNotEqual(result['author_kind'],'organization')
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'source.docx';d=Document();d.add_heading('参考文献',1)
            d.add_paragraph(ref);d.add_paragraph('[2] 工业和信息化部办公厅. 通知[EB/OL]. (2022-11-03)[2026-10-30].');d.save(path)
            a=Audit(path);a.infer_roles();a.check_citations()
            warnings=[i for i in a.issues if i['category']=='reference.author_listing_needs_review']
            self.assertEqual(len(warnings),1);self.assertIn('张三和李四',warnings[0]['actual']['author_list'])

    def test_four_digit_journal_pages_are_not_publication_years(self):
        refs=[('[1] HANELT A. Research[J]. Journal of Management Studies,2021,58(5):1159-1197. DOI:10.1111/joms.2020.1234.', ['2021']),
              ('[2] HADLOCK C J. Research[J]. Journal,2010,23(5):1909-1940.', ['2010'])]
        for reference,years in refs:
            with self.subTest(reference=reference):self.assertEqual(N._guess_reference(reference)['years'],years)

    def test_book_publisher_colon_is_not_mistaken_for_page_colon(self):
        self.assertEqual(N._guess_reference('[1] 张三. 书名[M]. 北京: 某出版社,2018.')['years'],['2018'])

    def test_real_multiple_publication_years_remain_ambiguous(self):
        self.assertEqual(N._guess_reference('[1] VIAL G. Work[J]. Journal,2019; revised edition,2021:1159-1197.')['years'],['2019','2021'])

    def test_typed_reference_without_year_does_not_guess_page_range(self):
        self.assertEqual(N._guess_reference('[1] VIAL G. Work[J]. Journal,58(5):1159-1197.')['years'],[])

    def test_unknown_reference_format_keeps_conservative_old_hints(self):
        self.assertEqual(N._guess_reference('[1] VIAL G. Work. Journal,2021:1159-1197.')['years'],['2021','1159','1197'])


if __name__=='__main__':unittest.main()

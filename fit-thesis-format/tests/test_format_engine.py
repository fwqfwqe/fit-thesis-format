import copy
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import thesis_format as F
import format_engine as G
import narrative_citations as N


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def sample(self):
        d=Document();d.add_paragraph('测试论文题目');d.add_paragraph('摘要');d.add_paragraph('摘要正文。');d.add_paragraph('目录')
        d.add_heading('1. 绪论',1);d.add_paragraph('  保留 English words 与数据 0.001[1]。  ');d.add_paragraph('')
        p=d.add_paragraph();m=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='x';mr.append(mt);m.append(mr);p._p.append(m)
        d.add_paragraph('表1-1 机制检验');t=d.add_table(rows=5,cols=4)
        for row,values in zip(t.rows,[['Panel A 回归','(1)','(2)','(3)'],['DID','0.10','0.20','0.30'],['Panel B 效应分解','效应值','区间','占比'],['间接效应','0.03','[0.01,0.05]','10%'],['企业数','12','12','12']]):
            for c,v in zip(row.cells,values):c.text=v
        d.add_heading('参考文献',1);d.add_paragraph('[1] 张三. 文献[J].期刊,2020.');d.add_heading('致谢',1);d.add_paragraph('感谢导师。');d.add_heading('附录',1);d.add_paragraph('附录数据 123。')
        p=self.root/'source.docx';d.save(p);r=G.inspect_document(p);dec={'source_sha256':r['source_sha256'],'paragraphs':{},'metadata':{'title':'测试论文题目','name':'测试用户','student_id':'123','college':'经济管理学院','major':'投资学','grade':'2027','date':'2027年'}}
        for c in r['candidates']:
            if c['text']=='测试论文题目':dec['paragraphs'][c['id']]={'role':'cn_title','expected_text':c['text']}
        return p,dec
    def test_replay_and_second_run(self):
        p,dec=self.sample();a=self.root/'a.docx';b=self.root/'b.docx';source=p.read_bytes()
        report=G.format_document(p,a,dec);G.format_document(a,b)
        self.assertEqual(source,p.read_bytes());self.assertTrue(G.verify_format(p,a,dec)['ok'])
        self.assertEqual(F.package(a),F.package(b))
        self.assertEqual(report['final_mapping']['source_sha256'],F.sha(a.read_bytes()))
    def test_panel_numbers_and_math(self):
        p,dec=self.sample();a=self.root/'a.docx';G.format_document(p,a,dec)
        original=F.parse(F.package(p)['word/document.xml']);final=F.parse(F.package(a)['word/document.xml'])
        self.assertEqual(original.xpath('.//m:t/text()',namespaces=F.NS),final.xpath('.//m:t/text()',namespaces=F.NS))
        table=next(t for t in final.iter(F.q('tbl')) if 'Panel B' in G.text(t))
        self.assertIn('0.03',G.text(table));self.assertIn('[0.01,0.05]',G.text(table))
        edge=table.findall(F.q('tr'))[2].find('w:tc/w:tcPr/w:tcBorders/w:top',F.NS)
        self.assertEqual(edge.get(F.q('val')),'single');self.assertEqual(edge.get(F.q('sz')),'6')
        self.assertLess(G.text(final).index('附录数据'),G.text(final).index('感谢导师'))
    def test_missing_identity_never_copies_example(self):
        p,dec=self.sample();dec['metadata']={'title':'测试论文题目'};a=self.root/'a.docx';report=G.format_document(p,a,dec)
        final=G.text(F.parse(F.package(a)['word/document.xml']))
        for value in ('李四','098605569','2025年11月3日','福建纺织品'):self.assertNotIn(value,final)
        self.assertTrue(any(i['code']=='missing_identity_fields' for i in report['unresolved_issues']))
    def test_hidden_unmapped_body_fails_legacy_preservation(self):
        p,dec=self.sample();parts=F.package(p);root=F.parse(parts['word/document.xml']);before=F.inventory(parts)
        paragraph=root.findall('.//w:p',F.NS)[5];run=paragraph.find(F.q('r'));F.child(F.ensure(run,'rPr',True),'vanish')
        changed=dict(parts);changed['word/document.xml']=F.xml(root)
        self.assertTrue(F.compare_content(before,F.inventory(changed)))
    def test_stale_decisions_and_existing_output_rejected(self):
        p,dec=self.sample();out=self.root/'out.docx';bad=copy.deepcopy(dec);bad['source_sha256']='0'*64
        with self.assertRaises(ValueError):G.format_document(p,out,bad)
        with self.assertRaises(ValueError):G.format_document(p,p,dec)
    def test_long_table_report_is_body_not_caption(self):
        p,dec=self.sample();d=Document(p);d.add_paragraph('表1-1报告了所有企业的结果。'+('正文讨论。'*35));d.save(p)
        r=G.inspect_document(p);item=next(v for v in r['map']['paragraphs'].values() if v['expected_text'].startswith('表1-1报告'))
        self.assertEqual(item['role'],'body')
    def test_complex_citation_keeps_math_and_styles_full_brackets(self):
        p,dec=self.sample();d=Document(p);body=d.paragraphs[5];body.text='变量 '
        math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='x';mr.append(mt);math.append(mr);body._p.append(math);body.add_run(' 使用文献[1]。');d.save(p)
        dec['source_sha256']=F.sha(p.read_bytes());out=self.root/'out.docx';r=G.format_document(p,out,dec)
        root=F.parse(F.package(out)['word/document.xml']);target=next(x for x in root.iter(F.q('p')) if G.text(x).startswith('变量'))
        self.assertEqual(target.xpath('.//m:t/text()',namespaces=F.NS),['x'])
        cited=[x for x in target.findall(F.q('r')) if x.find('w:rPr/w:rStyle',F.NS) is not None and x.find('w:rPr/w:rStyle',F.NS).get(F.q('val'))=='FIT_citation_superscript']
        self.assertEqual(''.join(G.text(x) for x in cited),'[1]');self.assertFalse(any(x['code'].startswith('numeric_citation') for x in r['unresolved_issues']))
    def test_broken_reference_repaired_without_list_number_switch(self):
        p,dec=self.sample();d=Document(p);body=d.paragraphs[5];body.text='原引文';G.field(body._p,' REF _MissingTarget \\r \\h ','[1]');d.save(p)
        dec['source_sha256']=F.sha(p.read_bytes());out=self.root/'out.docx';G.format_document(p,out,dec)
        root=F.parse(F.package(out)['word/document.xml']);codes=root.xpath('.//w:instrText/text()',namespaces=F.NS)
        ref=next(x for x in codes if 'FITv10Reference1Full' in x)
        self.assertNotIn('\\r',ref);self.assertTrue(root.xpath('.//w:bookmarkStart[@w:name="FITv10Reference1Full"]',namespaces=F.NS))
        self.assertIn('原引文[1]',G.text(root))
    def test_author_display_uses_existing_identity_and_decimal_sentence(self):
        p,dec=self.sample();d=Document(p);d.paragraphs[5].text='工业和信息化部（2026）公布规模1.2万亿元[1]。'
        ref=next(x for x in d.paragraphs if x.text.startswith('[1]'));ref.text='[1] 工业和信息化部.行业统计[EB/OL].2026.';d.save(p)
        inv=N.inspect(p);self.assertEqual(len(inv['references'][0]['authors']),1)
        self.assertFalse(any('missing_numeric_marker' in x['issues'] for x in inv['candidates']))
    def test_equation_reference_native_field_preserves_visible_label(self):
        p,dec=self.sample();d=Document(p);d.paragraphs[5].text='参见公式（1-1）。';d.save(p);dec['source_sha256']=F.sha(p.read_bytes())
        out=self.root/'out.docx';G.format_document(p,out,dec);root=F.parse(F.package(out)['word/document.xml'])
        self.assertIn('参见公式（1-1）。',G.text(root));self.assertTrue(any('REF FITv10Equation' in code for code in root.xpath('.//w:instrText/text()',namespaces=F.NS)))
    def test_unsupported_citation_insertion_is_rejected(self):
        p,dec=self.sample();ins={'pid':'p000006','expected_text':'wrong','position':0,'text':'[1]','claim':{'start':0,'end':0,'expected_text':''},'references':[]};dec['citation_insertions']=[ins]
        with self.assertRaises(ValueError):G.format_document(p,self.root/'out.docx',dec)

    def test_layout_repair_source_bound_and_persistent(self):
        p,dec=self.sample();root=F.parse(F.package(p)['word/document.xml']);pid='p000006';target=F.indexed(root)[0][pid]
        dec['layout_repairs']=[{'kind':'paragraph','id':pid,'expected_text':G.text(target),'set':{'page_break_before':True,'keep_with_next':True,'keep_lines':True,'before_pt':2,'after_pt':3,'line_min_pt':24}}]
        a=self.root/'a.docx';b=self.root/'b.docx';report=G.format_document(p,a,dec);G.format_document(a,b)
        self.assertEqual(F.package(a),F.package(b));self.assertTrue(G.verify_format(p,a,dec)['ok'])
        repair=next(o for o in report['operations'] if o['kind']=='layout_repair');self.assertEqual(repair['source_target']['id'],pid)
        final=F.parse(F.package(a)['word/document.xml']);paragraph=next(x for x in final.iter(F.q('p')) if G.text(x).startswith('保留 English'))
        styles=F.parse(F.package(a)['word/styles.xml']);sid=paragraph.find('w:pPr/w:pStyle',F.NS).get(F.q('val'));st=styles.find('w:style[@w:styleId="'+sid+'"]',F.NS)
        self.assertEqual(G._read_repair_style(st)[1],dec['layout_repairs'][0]['set'])

    def test_layout_repair_rejects_stale_invalid_and_unbounded(self):
        p,dec=self.sample();root=F.parse(F.package(p)['word/document.xml']);source=G.text(F.indexed(root)[0]['p000006'])
        for values in ({'before_pt':float('nan')},{'after_pt':145},{'font':'Arial'},{'keep_lines':'yes'},{'line_min_pt':0}):
            bad=copy.deepcopy(dec);bad['layout_repairs']=[{'kind':'paragraph','id':'p000006','expected_text':source,'set':values}]
            with self.assertRaises(ValueError):G.format_document(p,self.root/'bad.docx',bad)
        bad=copy.deepcopy(dec);bad['layout_repairs']=[{'kind':'paragraph','id':'p000006','expected_text':'stale','set':{'keep_lines':True}}]
        with self.assertRaises(ValueError):G.format_document(p,self.root/'bad.docx',bad)

    def test_picture_fit_preserves_crop_media_and_aspect(self):
        from PIL import Image
        pic=self.root/'picture.png';Image.new('RGB',(400,200),'red').save(pic)
        p,dec=self.sample();d=Document(p);shape=d.add_picture(str(pic),width=Inches(12));crop=OxmlElement('a:srcRect');crop.set('l','500');shape._inline.find('.//{'+G.PIC+'}blipFill').insert(1,crop);d.save(p);dec['source_sha256']=F.sha(p.read_bytes())
        a=self.root/'a.docx';b=self.root/'b.docx';report=G.format_document(p,a,dec);G.format_document(a,b)
        before=F.package(p);after=F.package(a);self.assertEqual({n:v for n,v in before.items() if n.startswith('word/media/')},{n:v for n,v in after.items() if n in before and n.startswith('word/media/')})
        root=F.parse(after['word/document.xml']);image=next(x for x in root.iter('{'+G.WP+'}inline') if any(r.get('l')=='500' for r in x.findall('.//{'+G.A+'}srcRect')))
        extent=image.find('{'+G.WP+'}extent');cx=int(extent.get('cx'));cy=int(extent.get('cy'));self.assertLessEqual(abs(cx/cy-2),0.00001);self.assertEqual(image.find('.//{'+G.A+'}srcRect').get('l'),'500')
        self.assertTrue(any(o['kind']=='picture_fit' for o in report['operations']));self.assertEqual(F.package(a),F.package(b))

    def test_table_continuation_repeats_only_header_and_is_idempotent(self):
        p,dec=self.sample();root=F.parse(F.package(p)['word/document.xml']);table=F.indexed(root)[1]['t0001'];source_rows=[G.text(r) for r in table.findall(F.q('tr'))]
        dec['tables']={'t0001':{'continuation_breaks':[3],'expected_sha256':F.sha(F.xml(table))}}
        a=self.root/'a.docx';b=self.root/'b.docx';report=G.format_document(p,a,dec);G.format_document(a,b)
        final=F.parse(F.package(a)['word/document.xml']);tables=[t for t in final.iter(F.q('tbl')) if 'Panel A 回归' in G.text(t)]
        self.assertEqual(len(tables),2);self.assertEqual([G.text(r) for r in tables[0].findall(F.q('tr'))],source_rows[:3]);self.assertEqual([G.text(r) for r in tables[1].findall(F.q('tr'))],[source_rows[0]]+source_rows[3:])
        label=tables[1].getprevious();self.assertEqual(G.text(label),'续前表');self.assertEqual(label.find('w:pPr/w:pStyle',F.NS).get(F.q('val')),'FIT_continuation_label')
        op=next(o for o in report['operations'] if o['kind']=='table_continuation');self.assertEqual(op['added_header_texts'],['Panel A 回归','(1)','(2)','(3)']);self.assertEqual(op['label_count'],1)
        self.assertEqual(F.package(a),F.package(b));self.assertTrue(G.verify_format(p,a,dec)['ok'])

    def test_continuation_rejects_cross_boundary_merge_and_stale_hash(self):
        p,dec=self.sample();d=Document(p);d.tables[0].cell(2,0).merge(d.tables[0].cell(3,0));d.save(p);dec['source_sha256']=F.sha(p.read_bytes());root=F.parse(F.package(p)['word/document.xml']);table=F.indexed(root)[1]['t0001']
        dec['tables']={'t0001':{'continuation_breaks':[3],'expected_sha256':F.sha(F.xml(table))}}
        with self.assertRaisesRegex(ValueError,'Vertical merge'):G.format_document(p,self.root/'out.docx',dec)
        dec['tables']['t0001']['expected_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'exact source'):G.format_document(p,self.root/'out.docx',dec)

    def test_appendix_code_keeps_spaces_fonts_and_indent(self):
        p,dec=self.sample();d=Document(p);values=['    for i in range(3):  ','        print(i)','    mutate(x = x + 1)  ','    regress y x  ']
        for value in values:
            par=d.add_paragraph(value);par.runs[0].font.name='Consolas';par.paragraph_format.left_indent=Inches(0.3)
        d.save(p);dec['source_sha256']=F.sha(p.read_bytes());root=F.parse(F.package(p)['word/document.xml'])
        for pid,par in F.indexed(root)[0].items():
            if G.text(par) in values:dec['paragraphs'][pid]={'role':'appendix_code','expected_text':G.text(par)}
        a=self.root/'a.docx';b=self.root/'b.docx';G.format_document(p,a,dec);G.format_document(a,b);after=F.parse(F.package(a)['word/document.xml'])
        self.assertEqual([G.text(par) for par in after.iter(F.q('p')) if G.text(par) in values],values)
        for par in after.iter(F.q('p')):
            if G.text(par) in values:self.assertEqual(par.find('w:r/w:rPr/w:rFonts',F.NS).get(F.q('ascii')),'Consolas');self.assertEqual(par.find('w:pPr/w:ind',F.NS).get(F.q('left')),'432')
        self.assertEqual(F.package(a),F.package(b))

    def test_appendix_removes_spacing_paragraph_but_keeps_confirmed_blank_code(self):
        p,dec=self.sample();d=Document(p)
        d.add_paragraph('   ');d.add_paragraph('');d.save(p)
        dec['source_sha256']=F.sha(p.read_bytes())
        ps=F.indexed(F.parse(F.package(p)['word/document.xml']))[0]
        spacer=next(k for k,v in ps.items() if G.text(v)=='   ')
        code_id=list(ps)[-1]
        dec['paragraphs'][code_id]={'role':'appendix_code','expected_text':''}
        a=self.root/'a.docx';b=self.root/'b.docx'
        report=G.format_document(p,a,dec);G.format_document(a,b)
        removed={o['locator'] for o in report['operations'] if o['kind']=='remove_empty_paragraph'}
        self.assertIn(spacer,removed);self.assertNotIn(code_id,removed)
        final_ps=F.indexed(F.parse(F.package(a)['word/document.xml']))[0]
        code=[final_ps[k] for k,v in report['final_mapping']['paragraphs'].items() if (v if isinstance(v,str) else v['role'])=='appendix_code']
        self.assertEqual(len(code),1);self.assertEqual(G.text(code[0]),'')
        self.assertTrue(G.verify_format(p,a,dec)['ok']);self.assertEqual(F.package(a),F.package(b))

    def test_heading_style_only_prose_requires_decision(self):
        p,dec=self.sample();d=Document(p);value='这是错误套用了标题样式的一段正文，不能仅凭样式判断章节层级。';d.add_heading(value,1);d.save(p);dec['source_sha256']=F.sha(p.read_bytes())
        inventory=G.inspect_document(p);candidate=next(c for c in inventory['candidates'] if c['text']==value);self.assertEqual(candidate['role'],'chapter')
        a=self.root/'a.docx';report=G.format_document(p,a,dec);self.assertTrue(any(x['code']=='role_confirmation' and x['location']==candidate['id'] for x in report['unresolved_issues']))
        final=F.parse(F.package(a)['word/document.xml']);par=next(p for p in final.iter(F.q('p')) if G.text(p)==value);self.assertIsNone(par.find('w:pPr/w:numPr',F.NS));self.assertFalse(par.findall('.//w:instrText',F.NS))
        d=Document(p);d.add_heading('企业依预算安排投资',1);d.save(p)
        self.assertTrue(any(c['text']=='企业依预算安排投资' for c in G.inspect_document(p)['candidates']))

    def test_adjacent_caption_positions_are_corrected_without_text_change(self):
        from PIL import Image
        p,dec=self.sample();d=Document(p);caption=next(x._p for x in d.paragraphs if x.text.startswith('表1-1'));table=d.tables[0]._tbl;caption.getparent().remove(caption);table.addnext(caption)
        pic=self.root/'p.png';Image.new('RGB',(30,20),'blue').save(pic);d.add_paragraph('图1-1 测试图');d.add_picture(str(pic),width=Inches(2));d.save(p);dec['source_sha256']=F.sha(p.read_bytes())
        a=self.root/'a.docx';b=self.root/'b.docx';report=G.format_document(p,a,dec);G.format_document(a,b);root=F.parse(F.package(a)['word/document.xml'])
        tc=next(x for x in root.iter(F.q('p')) if G.text(x)=='表1-1 机制检验');fc=next(x for x in root.iter(F.q('p')) if G.text(x)=='图1-1 测试图')
        self.assertEqual(tc.getnext().tag,F.q('tbl'));self.assertTrue(fc.getprevious().findall('.//w:drawing',F.NS));self.assertEqual(sum(o['kind']=='caption_position' for o in report['operations']),2);self.assertEqual(F.package(a),F.package(b))

    def test_plain_formula_transcription_integrity_and_tamper_rejection(self):
        from integrity import verify_transition
        p,dec=self.sample();d=Document(p);value='y=x+1 公式（1-2）';d.add_paragraph(value);d.save(p);dec['source_sha256']=F.sha(p.read_bytes());root=F.parse(F.package(p)['word/document.xml']);pid=next(k for k,v in F.indexed(root)[0].items() if G.text(v)==value)
        dec['paragraphs'][pid]={'role':'equation','expected_text':value};a=self.root/'a.docx';b=self.root/'b.docx';report=G.format_document(p,a,dec)
        self.assertTrue(any(o['kind']=='formula_transcription' and o['locator']==pid for o in report['operations']));self.assertTrue(verify_transition(p,a,'format',report)['passed']);G.format_document(a,b);self.assertEqual(F.package(a),F.package(b))
        parts=F.package(a);final=F.parse(parts['word/document.xml']);mt=next(x for x in final.iter('{'+F.M+'}t') if x.text=='y=x+1');mt.text='y=x+2';parts['word/document.xml']=F.xml(final);tampered=self.root/'tampered.docx'
        with zipfile.ZipFile(tampered,'w') as z:
            for name,data in parts.items():z.writestr(name,data)
        self.assertFalse(verify_transition(p,tampered,'format',report)['passed'])

    def test_ambiguous_plain_formula_remains_unresolved(self):
        p,dec=self.sample();d=Document(p);value='y=x^2 公式（1-2）';d.add_paragraph(value);d.save(p);dec['source_sha256']=F.sha(p.read_bytes());root=F.parse(F.package(p)['word/document.xml']);pid=next(k for k,v in F.indexed(root)[0].items() if G.text(v)==value);dec['paragraphs'][pid]={'role':'equation','expected_text':value}
        a=self.root/'a.docx';report=G.format_document(p,a,dec);self.assertTrue(any(i['code']=='formula_transcription_review' for i in report['unresolved_issues']));self.assertIn(value,G.text(F.parse(F.package(a)['word/document.xml'])))

if __name__=='__main__':unittest.main()
